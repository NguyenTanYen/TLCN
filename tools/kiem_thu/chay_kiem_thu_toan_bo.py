"""KIỂM THỬ TOÀN BỘ CHỨC NĂNG – chạy mọi ca kiểm thử tự động, đo độ phủ API và xuất bảng kết quả Excel.

Chạy (từ thư mục HeThong, sau khi đã chạy 1_CAI_DAT_PYTHON.bat):
    .venv\\Scripts\\python tools\\kiem_thu\\chay_kiem_thu_toan_bo.py            (Windows)
    python tools/kiem_thu/chay_kiem_thu_toan_bo.py [--omr]                      (Linux)

Kết quả: docs/kiem_thu/KET_QUA_KIEM_THU.xlsx gồm
  1. Tổng hợp            – số ca đạt/không đạt theo nhóm chức năng, độ phủ API
  2. Ca kiểm thử tự động – từng ca: nhóm, nội dung, kết quả mong đợi, kết quả, thời gian
  3. Độ phủ API          – từng endpoint: nhóm, mô tả, đã gọi thành công chưa, bởi ca nào
  4. Kiểm thử giao diện  – danh sách ca kiểm thử thủ công trên giao diện và Moodle (điền khi trình diễn)
  5. Chấm phiếu (OMR)    – (khi có --omr) độ chính xác bộ chấm phiếu theo điều kiện ảnh và kiểu tô
LƯU Ý: kiểm thử tự động DỰNG LẠI CSDL assessment_db (dữ liệu minh họa); chạy xong hãy chạy lại 4_TAO_DU_LIEU_DEMO_TREN_MOODLE.bat.
"""
from __future__ import annotations

import datetime as dt
import json
import os
import subprocess
import sys
import xml.etree.ElementTree as ET
from collections import defaultdict
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
BACKEND = ROOT / "backend"
OUT_DIR = ROOT / "docs" / "kiem_thu"
sys.path.insert(0, str(HERE)); sys.path.insert(0, str(BACKEND))
from mo_ta_ca_kiem_thu import M, NHOM  # noqa: E402

GIAO_DIEN = [
    # (vai trò, chức năng, các bước, kết quả mong đợi)
    ("GV", "Đăng nhập", "Mở http://localhost:8000, đăng nhập gv.son / Gv@123456", "Vào trang 'Môn học của tôi'; tên GV ở góc trên bên phải"),
    ("GV", "Đăng nhập một lần từ Moodle", "Trên Moodle, vào khóa DBMS330284_01, bấm 'Phân tích CĐR'", "Mở thẳng lớp DBMS330284_01, không hỏi mật khẩu"),
    ("SV", "SV không vào hệ thống", "Đăng nhập hệ thống bằng tài khoản SV 22130001", "Bị từ chối, hướng dẫn xem kết quả trên Moodle"),
    ("GV", "Môn học của tôi", "Xem trang chủ GV", "Môn DBMS330284 với 2 lớp của GV; khung khóa học Moodle chưa có trong hệ thống"),
    ("GV", "Đề cương chi tiết", "Bấm tên môn DBMS330284", "Thẻ Đề cương: thông tin chung, tiên quyết, CTĐT, CLO–PLO, PI, chương, BM6a; không có nút sửa"),
    ("GV", "GV không sửa CLO", "Tìm nút sửa CLO/chương trên đề cương", "Không có; ghi 'Bộ môn quản lý – giảng viên chỉ xem'"),
    ("Bộ môn", "CTĐT 2026", "Đăng nhập admin / Admin@123 → Chương trình đào tạo", "3 CTĐT; CNTT-2026: 7 PLO, 17 PI, 79 học phần; KTDL-2026: 57 học phần"),
    ("Bộ môn", "Ma trận học phần – PI", "Mở KTDL-2026 → thẻ Ma trận học phần – PI; sửa một ô", "Ô lưu được, hiện ngay trên đề cương học phần"),
    ("Bộ môn", "Học phần theo học kỳ", "Thẻ Học phần của CTĐT", "Nhóm theo học kỳ, khối kiến thức, bắt buộc/tự chọn"),
    ("Bộ môn", "Khai báo học phần", "Học phần → DBMS330284 → thẻ Khai báo; thêm chương, CLO, sửa đề cương", "Lưu được; bảng kiểm tra 5 bước đủ"),
    ("GV", "Ngân hàng câu hỏi", "Ngân hàng câu hỏi → lọc CLO2, 'Chưa gán'", "Lọc đúng; p, DI trung bình tô màu"),
    ("GV", "Thêm/sửa/Hủy câu hỏi", "Thêm câu 4 phương án gắn CLO; Hủy câu có lý do", "Lưu đúng; câu đã có kết quả không sửa được"),
    ("GV", "Nhập câu hỏi từ tệp, gán CLO từ tệp", "Nhập câu hỏi từ tệp (Excel/Aiken/XML) → Gán CLO/Bloom từ file", "Xem trước, phát hiện trùng/lỗi; câu được gán"),
    ("GV", "Đồng bộ câu hỏi với Moodle", "Ngân hàng câu hỏi → ⇅ Đồng bộ với Moodle → chọn khóa #41 → Đồng bộ", "Câu mới được đẩy lên; câu sửa trên Moodle kéo về; Moodle giữ phiên bản cũ"),
    ("GV", "Sinh đề theo ma trận", "Lớp DBMS330284_01 → + Tạo bài kiểm tra → nhập 20 câu → Sinh đề", "Ma trận chương × Bloom, phủ đủ CLO, báo ô thiếu"),
    ("GV", "Tạo Quiz trên Moodle", "Bài trực tuyến mới → Tạo Quiz trên Moodle", "Quiz xuất hiện trong khóa học Moodle, đúng câu hỏi"),
    ("GV", "Dùng đề cho lớp khác", "Trong bài KT → 'Dùng đề cho lớp khác' → chọn lớp 02", "Lớp 02 có bài giống hệt; bài trực tuyến tạo Quiz nếu lớp đã gắn Moodle"),
    ("GV", "Đề giấy trên hệ thống: mã đề, xem trước", "Bài giấy mới → In và chấm ngay trên hệ thống → sinh 2 mã đề → Xem trước", "Trang A4 từng mã đề; đổi 1 cột/2 cột; sửa câu ngay trên bản xem trước"),
    ("GV", "Câu 2/3/4 phương án", "Đề có câu Đúng/Sai và câu 3 phương án", "Đề và phiếu chỉ có đúng số ô của từng câu"),
    ("GV", "Phiếu trả lời, xuất tệp", "Thẻ Phiếu trả lời; tải PDF đề, Word, đáp án, trọn bộ .zip", "Tệp mở được, tiếng Việt đúng"),
    ("GV", "Chấm phiếu bằng ảnh", "In phiếu, tô, chụp bằng điện thoại → Chấm phiếu → tải ảnh → soát → Lưu", "Đọc đúng MSSV, mã đề, điểm; phiếu tô 2 ô bị cảnh báo; lưu xong bài 'Đã phân tích'"),
    ("GV", "Đề giấy qua Moodle Offline Quiz", "Bài giấy → Dùng Moodle Offline Quiz → tạo, tải đề/phiếu; quét phiếu lên Moodle; Đồng bộ", "Moodle chấm; hệ thống kéo kết quả theo nhóm đề"),
    ("GV", "Nhập điểm thủ công", "Lớp 02 → + Tạo bài kiểm tra → Nhập điểm → khai báo câu/ý, CLO → nhập điểm", "Kiểm tra giới hạn điểm; Lưu & phân tích → CLO, p, DI, alpha"),
    ("GV", "Nhập điểm từ Excel", "Tải bảng điểm mẫu → điền → Nhập từ Excel; hoặc 'Tạo từ file Excel có sẵn'", "Xem trước lỗi từng dòng; ghi điểm đúng cột"),
    ("GV", "Đồng bộ bài làm Moodle", "Bài Thi cuối kỳ → Đồng bộ & phân tích", "38 bài làm, 2 vắng; nhật ký đồng bộ thành công"),
    ("GV", "Phân tích câu hỏi", "Thẻ Phân tích câu hỏi", "p, DI, phân loại; câu 14 DI âm được đánh dấu"),
    ("GV", "Gộp bài kiểm tra", "Trang môn → Gộp bài KT đánh giá CĐR → chọn 3 bài → Phân tích gộp", "Mức đạt CLO; đóng góp PLO theo từng CTĐT; xuất Excel"),
    ("GV", "Chọn bài làm minh chứng BM6", "Bấm 'Có gộp/Không gộp' ở một bài", "BM6b của lớp tính lại ngay"),
    ("GV", "Phụ đạo", "Lớp 01 → Phụ đạo (câu sai) → bấm 'Câu sai' của 1 SV → Tạo bài phụ đạo", "Danh sách câu sai đúng; bài phụ đạo không gộp minh chứng"),
    ("GV", "Kết quả CĐR (BM6), xuất biểu mẫu", "Lớp 01 → Kết quả CĐR (BM6) → Xuất Excel; BM6c từng bài", "Tệp đúng bố cục biểu mẫu"),
    ("GV", "Thống kê lớp, hồ sơ SV", "Thống kê lớp & từng SV → bấm 1 SV", "Chỉ số, bản đồ nhiệt, phần cần học lại"),
    ("GV", "Công bố kết quả lên Moodle", "Bài Thi cuối kỳ → Công bố kết quả lên Moodle", "SV nhận thông báo; trên Moodle thấy radar CLO, lộ trình"),
    ("SV", "Xem kết quả trên Moodle", "Đăng nhập Moodle bằng SV → khóa học → Kết quả phân tích CĐR", "Điểm, CLO đạt/chưa đạt, lộ trình ôn tập; bài chưa công bố không hiện"),
    ("GV", "Liên kết lớp ⇄ Moodle", "Lớp chưa gắn Moodle → 'Tự nối theo mã lớp'; hoặc bấm 'Phân tích CĐR' trong khóa trùng mã", "Lớp được gắn với khóa học có shortname = mã lớp"),
    ("Bộ môn", "Kế hoạch đo PI (BM3b), Tổng hợp PLO", "Kế hoạch đo PI → thêm/sửa; Tổng hợp PLO → xuất Excel/Word", "BM2, BM3 đúng biểu mẫu"),
    ("Bộ môn", "Kết nối Moodle", "Kết nối Moodle", "CSDL Moodle đọc được, cầu nối đã cài, Web Service OK (7 hàm)"),
]


def run_pytest(out_json: Path, junit: Path) -> int:
    env = {**os.environ, "PYTHONPATH": os.pathsep.join([str(HERE), os.environ.get("PYTHONPATH", "")]), "GHI_API_OUT": str(out_json),
           "PYTHONIOENCODING": "utf-8"}
    cmd = [sys.executable, "-m", "pytest", "-q", "-p", "ghi_api", f"--junitxml={junit}", "tests"]
    print("Đang chạy:", " ".join(cmd), "(khoảng 1–2 phút)…", flush=True)
    return subprocess.run(cmd, cwd=BACKEND, env=env).returncode


def all_routes() -> list[dict]:
    from app.main import app  # noqa: F401
    from app.routers import (analytics, auth, catalog, courses, exams, grading, insights, lms, paper_exam, programs, qsync,
                             question_io, questions, remedial, reports, sync)
    out = []
    for mod in (auth, programs, catalog, courses, question_io, questions, exams, grading, remedial, paper_exam, qsync, sync, lms,
                analytics, insights, reports):
        for r in mod.router.routes:
            doc = (r.endpoint.__doc__ or "").strip().split("\n")[0]
            for m in sorted(r.methods - {"HEAD"}):
                out.append({"method": m, "path": r.path, "group": (mod.router.tags or [mod.__name__])[0], "desc": doc or r.endpoint.__name__})
    out += [{"method": "GET", "path": "/api/health", "group": "Hệ thống", "desc": "Kiểm tra dịch vụ"},
            {"method": "GET", "path": "/api/public-config", "group": "Hệ thống", "desc": "Cấu hình công khai cho giao diện"}]
    return out


def omr_rows() -> list[list]:
    r = subprocess.run([sys.executable, str(ROOT / "tools" / "omr" / "kiem_tra_omr.py")], capture_output=True, text=True,
                       encoding="utf-8", cwd=ROOT)
    rows = []
    for line in r.stdout.splitlines():
        parts = line.split("\t")
        if len(parts) >= 3:
            rows.append(parts)
    return rows


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    hit_json, junit = OUT_DIR / "ghi_api.json", OUT_DIR / "junit.xml"
    code = run_pytest(hit_json, junit)
    # ---------------- kết quả từng ca
    cases = []
    for tc in ET.parse(junit).getroot().iter("testcase"):
        name, f = tc.get("name"), tc.get("classname", "").split(".")[-1]
        bad = tc.find("failure") if tc.find("failure") is not None else tc.find("error")
        skip = tc.find("skipped") is not None
        g, title, what, expect = M.get(name, (len(NHOM) - 1, name, "", ""))
        cases.append({"file": f, "name": name, "group": NHOM[g], "title": title, "what": what, "expect": expect,
                      "result": "BỎ QUA" if skip else ("KHÔNG ĐẠT" if bad is not None else "ĐẠT"),
                      "msg": (bad.get("message") or "")[:300] if bad is not None else "", "time": float(tc.get("time", 0))})
    # ---------------- độ phủ API
    hits = json.load(open(hit_json, encoding="utf-8")) if hit_json.exists() else []
    by = defaultdict(lambda: {"ok": set(), "any": set()})
    for m, p, st, t in hits:
        by[(m, p)]["any"].add(t)
        if st < 400:
            by[(m, p)]["ok"].add(t)
    routes = all_routes()
    for r in routes:
        b = by.get((r["method"], r["path"]), {"ok": set(), "any": set()})
        r["ok"], r["tests"] = bool(b["ok"]), sorted(b["ok"] or b["any"])
        r["err_only"] = not b["ok"] and bool(b["any"])
    n_ok = sum(c["result"] == "ĐẠT" for c in cases)
    n_cov = sum(r["ok"] for r in routes)
    # ---------------- Excel
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    thin = Side(style="thin", color="999999"); B = Border(left=thin, right=thin, top=thin, bottom=thin)
    head = PatternFill("solid", fgColor="1F4E79"); green = PatternFill("solid", fgColor="E2F0D9"); red = PatternFill("solid", fgColor="F8D7DA")

    def table(ws, header, rows, widths, start=1):
        for j, h in enumerate(header, 1):
            c = ws.cell(start, j, h); c.font = Font(bold=True, color="FFFFFF"); c.fill = head; c.border = B
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        for i, row in enumerate(rows, start + 1):
            for j, v in enumerate(row, 1):
                c = ws.cell(i, j, v); c.border = B; c.alignment = Alignment(wrap_text=True, vertical="top")
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.freeze_panes = ws.cell(start + 1, 1)

    ws = wb.active; ws.title = "Tong hop"
    ws["A1"] = "KẾT QUẢ KIỂM THỬ TOÀN BỘ CHỨC NĂNG – HỆ THỐNG KHẢO THÍ & ĐÁNH GIÁ CĐR"; ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Thời điểm chạy: {dt.datetime.now():%d/%m/%Y %H:%M} · Python {sys.version.split()[0]} · CSDL: kiểm thử dựng lại assessment_db"
    ws["A3"] = f"Ca kiểm thử tự động: {n_ok}/{len(cases)} đạt · Độ phủ API: {n_cov}/{len(routes)} endpoint được gọi thành công"
    ws["A3"].font = Font(bold=True, color="006100" if n_ok == len(cases) else "9C0006")
    rows = []
    for g in NHOM:
        cs = [c for c in cases if c["group"] == g]
        if cs:
            rows.append([g, len(cs), sum(c["result"] == "ĐẠT" for c in cs), sum(c["result"] == "KHÔNG ĐẠT" for c in cs),
                         "ĐẠT" if all(c["result"] == "ĐẠT" for c in cs) else "KHÔNG ĐẠT"])
    rows.append(["TỔNG", len(cases), n_ok, len(cases) - n_ok, "ĐẠT" if n_ok == len(cases) else "KHÔNG ĐẠT"])
    table(ws, ["Nhóm chức năng", "Số ca", "Đạt", "Không đạt", "Kết luận"], rows, [55, 10, 10, 12, 14], start=5)
    for r in ws.iter_rows(min_row=6, max_row=5 + len(rows)):
        r[4].fill = green if r[4].value == "ĐẠT" else red
    ws.cell(7 + len(rows), 1, "Kiểm thử bổ sung (ngoài pytest): tools/moodle_e2e/run_real_moodle_e2e.py (Moodle 4.5 thật), "
                              "tools/moodle_e2e/kiem_tra_dong_bo_cau_hoi.py, tools/kiem_tra_do_luong.py (tính lại độc lập các độ đo), "
                              "tools/omr/kiem_tra_omr.py (độ bền chấm phiếu), backend/tests/ui_screens.py (duyệt giao diện).")

    ws = wb.create_sheet("Ca kiem thu tu dong")
    table(ws, ["STT", "Nhóm chức năng", "Ca kiểm thử", "Nội dung kiểm tra", "Kết quả mong đợi", "Kết quả", "Thời gian (s)", "Tệp / hàm", "Ghi chú lỗi"],
          [[i, c["group"], c["title"], c["what"], c["expect"], c["result"], round(c["time"], 2), f"{c['file']}::{c['name']}", c["msg"]]
           for i, c in enumerate(sorted(cases, key=lambda c: (NHOM.index(c["group"]), c["file"])), 1)],
          [6, 30, 34, 50, 40, 12, 10, 40, 30])
    for r in ws.iter_rows(min_row=2):
        r[5].fill = green if r[5].value == "ĐẠT" else red

    ws = wb.create_sheet("Do phu API")
    table(ws, ["STT", "Phương thức", "Đường dẫn", "Nhóm", "Mô tả", "Gọi thành công", "Ca kiểm thử gọi"],
          [[i, r["method"], r["path"], r["group"], r["desc"], "Có" if r["ok"] else ("Chỉ lỗi" if r["err_only"] else "Chưa"),
            ", ".join(r["tests"][:6])] for i, r in enumerate(routes, 1)], [6, 10, 52, 30, 50, 14, 60])
    for r in ws.iter_rows(min_row=2):
        r[5].fill = green if r[5].value == "Có" else red

    ws = wb.create_sheet("Kiem thu giao dien")
    table(ws, ["Mã", "Vai trò", "Chức năng", "Các bước", "Kết quả mong đợi", "Kết quả (Đạt/Không)", "Ghi chú"],
          [[f"GD-{i:02d}", *g, "", ""] for i, g in enumerate(GIAO_DIEN, 1)], [8, 10, 30, 60, 55, 16, 25])

    if "--omr" in sys.argv:
        ws = wb.create_sheet("Cham phieu OMR")
        table(ws, ["Điều kiện ảnh / kiểu tô", "MSSV đúng (ảnh) · Câu đọc đúng (kiểu tô)", "Mã đề đúng (ảnh) · Số câu bị cảnh báo (kiểu tô)", "Câu đọc đúng", "Phiếu lỗi định vị", "Thời gian/phiếu"],
              omr_rows(), [60, 22, 26, 16, 14, 14])

    out = OUT_DIR / "KET_QUA_KIEM_THU.xlsx"
    wb.save(out)
    for f in (hit_json, junit):
        f.unlink(missing_ok=True)
    print(f"\nCa kiểm thử: {n_ok}/{len(cases)} đạt · Độ phủ API: {n_cov}/{len(routes)} endpoint")
    if n_cov < len(routes):
        print("Endpoint chưa được gọi thành công:", [f"{r['method']} {r['path']}" for r in routes if not r["ok"]])
    print("Bảng kết quả:", out)
    return 0 if code == 0 and n_ok == len(cases) else 1


if __name__ == "__main__":
    sys.exit(main())
