"""UC-01 (mở rộng): nhập hàng loạt câu hỏi và gán CLO / Bloom / chương từ file.

Định dạng câu hỏi được hỗ trợ:
- Excel (.xlsx) / CSV theo mẫu: Nội dung | A | B | C | D ... | Đáp án | (Chương | Bloom | CLO – tùy chọn)
- Aiken (.txt hoặc Word .docx): câu hỏi, các dòng "A. ...", dòng "ANSWER: B" – định dạng Moodle cũng dùng
- Moodle XML (.xml): câu hỏi multichoice một đáp án (xuất từ ngân hàng câu hỏi Moodle)
Chương / Bloom / CLO có thể để trống lúc nhập rồi gán sau bằng file gán CLO (Excel).
"""
from __future__ import annotations

import csv
import io
import re
import unicodedata
import xml.etree.ElementTree as ET
from html import unescape

LABELS = "ABCDEFGHIJ"
MAX_OPTIONS = 10


class ImportError_(ValueError):
    """File không đọc được hoặc sai định dạng."""


def norm(s) -> str:
    """Bỏ dấu, chữ thường, chỉ giữ chữ-số (so khớp tiêu đề cột, tên mức Bloom...)."""
    s = unicodedata.normalize("NFD", str(s or "")).replace("đ", "d").replace("Đ", "D")
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    return re.sub(r"[^a-z0-9]", "", s.lower())


def norm_content(s: str) -> str:
    """Chuẩn hóa nội dung để phát hiện câu trùng."""
    return re.sub(r"\s+", " ", str(s or "")).strip().lower()


def cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return "" if s.lower() == "nan" else s


# ------------------------------------------------------------------ đọc bảng (xlsx/csv)
HEADERS = {
    "id": {"id", "idcauhoi", "macau", "macauhoi", "questionid"},
    "content": {"noidung", "cauhoi", "noidungcauhoi", "question", "content", "questiontext", "debai"},
    "answer": {"dapan", "dapandung", "answer", "correct", "key", "phuonganđung", "phuongandung"},
    "chapter": {"chuong", "chapter", "chuongdecuong", "sochuong"},
    "bloom": {"bloom", "mucbloom", "mucdobloom", "capdobloom", "mucnhanthuc", "level"},
    "clo": {"clo", "cdr", "clos", "clotrongso", "cdrmonhoc", "chuandaura"},
}


def _map_header(h: str) -> str | None:
    n = norm(h)
    for key, names in HEADERS.items():
        if n in names:
            return key
    m = re.fullmatch(r"(?:phuongan|option|pa|dapan)?([a-j])", n)
    if m:
        return "opt_" + m.group(1).upper()
    return None


def read_table(data: bytes, filename: str) -> list[dict]:
    """Đọc sheet đầu tiên (hoặc sheet có cột Nội dung/ID) thành danh sách dict {khóa chuẩn: giá trị}."""
    name = filename.lower()
    rows: list[list] = []
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as ex:  # noqa: BLE001
            raise ImportError_(f"Không mở được file Excel: {ex}")
        sheets = wb.worksheets
        # ưu tiên sheet có tiêu đề nhận diện được
        for ws in sheets:
            cand = [list(r) for r in ws.iter_rows(values_only=True)]
            head_idx = _find_header(cand)
            if head_idx is not None:
                rows = cand[head_idx:]
                break
        else:
            raise ImportError_("Không tìm thấy dòng tiêu đề (cần có cột 'Nội dung' hoặc 'ID')")
    elif name.endswith(".csv"):
        text = _decode(data)
        dialect = csv.Sniffer().sniff(text[:2000], delimiters=",;\t") if text.strip() else csv.excel
        rows = [r for r in csv.reader(io.StringIO(text), dialect)]
        head_idx = _find_header(rows)
        if head_idx is None:
            raise ImportError_("Không tìm thấy dòng tiêu đề (cần có cột 'Nội dung' hoặc 'ID')")
        rows = rows[head_idx:]
    else:
        raise ImportError_("Chỉ hỗ trợ .xlsx hoặc .csv")
    header = [_map_header(cell(h)) for h in rows[0]]
    out = []
    for i, r in enumerate(rows[1:], start=2):
        d = {"_row": i}
        for k, v in zip(header, r):
            if k:
                d[k] = cell(v)
        if any(v for k, v in d.items() if k != "_row"):
            out.append(d)
    return out


def _find_header(rows: list[list]) -> int | None:
    for i, r in enumerate(rows[:15]):
        keys = {_map_header(cell(h)) for h in r if cell(h)}
        if "content" in keys or "id" in keys:
            return i
    return None


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1258", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", "replace")


# ------------------------------------------------------------------ đọc câu hỏi theo từng định dạng
def parse_questions(data: bytes, filename: str) -> tuple[str, list[dict]]:
    """Trả về (tên định dạng, danh sách câu thô {row, content, options[], answer, chapter, bloom, clo, errors[]})."""
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm", ".csv")):
        return "Excel/CSV", _from_table(read_table(data, filename))
    if name.endswith(".xml"):
        return "Moodle XML", _from_moodle_xml(data)
    if name.endswith(".docx"):
        from docx import Document
        try:
            doc = Document(io.BytesIO(data))
        except Exception as ex:  # noqa: BLE001
            raise ImportError_(f"Không mở được file Word: {ex}")
        return "Aiken (Word)", _from_aiken([p.text for p in doc.paragraphs])
    if name.endswith((".txt", ".gift", ".aiken")):
        return "Aiken (văn bản)", _from_aiken(_decode(data).splitlines())
    raise ImportError_("Định dạng chưa hỗ trợ. Dùng .xlsx, .csv, .txt/.docx (Aiken) hoặc .xml (Moodle XML)")


def _from_table(rows: list[dict]) -> list[dict]:
    out = []
    for r in rows:
        opts = [r.get("opt_" + L, "") for L in LABELS]
        while opts and not opts[-1]:
            opts.pop()
        q = {"row": r["_row"], "content": r.get("content", ""), "options": opts, "answer": r.get("answer", ""),
             "chapter": r.get("chapter", ""), "bloom": r.get("bloom", ""), "clo": r.get("clo", ""), "errors": []}
        if "" in opts:
            q["errors"].append(f"Phương án {LABELS[opts.index('')]} bị trống giữa các phương án")
        out.append(q)
    return out


OPT_RE = re.compile(r"^\s*([A-Ja-j])\s*[\.\)]\s*(.*\S)\s*$")
ANS_RE = re.compile(r"^\s*(ANSWER|ĐÁP ÁN|DAP AN|Đáp án)\s*[:：]\s*([A-Ja-j])\s*$", re.I)
META_RE = re.compile(r"^\s*(CHƯƠNG|CHUONG|CHAPTER|BLOOM|CLO|CĐR)\s*[:：]\s*(.+?)\s*$", re.I)


def _from_aiken(lines: list[str]) -> list[dict]:
    """Aiken: mỗi câu = nội dung (có thể nhiều dòng) + các dòng 'A. ...' + dòng 'ANSWER: X'.
    Có thể thêm các dòng tùy chọn 'CHƯƠNG: 3', 'BLOOM: Vận dụng', 'CLO: CLO2' ngay sau dòng ANSWER."""
    out, cur, line_no = [], None, 0

    def new(n):
        return {"row": n, "content": "", "options": [], "answer": "", "chapter": "", "bloom": "", "clo": "", "errors": []}

    for line_no, raw in enumerate(lines, 1):
        line = raw.strip()
        if not line:
            continue
        m_meta = META_RE.match(line)
        if m_meta and cur is not None and cur["answer"]:
            key = norm(m_meta.group(1))
            key = {"chuong": "chapter", "chapter": "chapter", "bloom": "bloom"}.get(key, "clo")
            cur[key] = m_meta.group(2)
            continue
        m_ans = ANS_RE.match(line)
        if m_ans:
            if cur is None:
                continue
            cur["answer"] = m_ans.group(2).upper()
            continue
        m_opt = OPT_RE.match(line)
        if m_opt and cur is not None and not cur["answer"] and cur["content"]:
            expected = LABELS[len(cur["options"])] if len(cur["options"]) < MAX_OPTIONS else "?"
            if m_opt.group(1).upper() == expected:
                cur["options"].append(m_opt.group(2))
                continue
        # dòng nội dung: bắt đầu câu mới nếu câu hiện tại đã có đáp án (hoặc chưa có câu)
        if cur is None or cur["answer"]:
            if cur is not None:
                out.append(cur)
            cur = new(line_no)
            line = re.sub(r"^\s*(Câu|Cau|Question|Q)\s*\d+\s*[\.:\)]\s*", "", line, flags=re.I)
            cur["content"] = line
        elif cur["options"]:
            # dòng thường sau khi đã có phương án: nối vào phương án cuối
            cur["options"][-1] += " " + line
        else:
            cur["content"] += "\n" + line
    if cur is not None:
        out.append(cur)
    for q in out:
        if not q["answer"]:
            q["errors"].append("Thiếu dòng 'ANSWER: X'")
    return out


def _strip_html(s: str) -> str:
    s = re.sub(r"<br\s*/?>|</p>", "\n", s or "", flags=re.I)
    s = re.sub(r"<[^>]+>", "", s)
    return re.sub(r"[ \t]+", " ", unescape(s)).strip()


def _from_moodle_xml(data: bytes) -> list[dict]:
    try:
        root = ET.fromstring(data)
    except ET.ParseError as ex:
        raise ImportError_(f"File XML không hợp lệ: {ex}")
    out = []
    for i, q in enumerate(root.iter("question"), 1):
        qtype = q.get("type")
        if qtype == "category":
            continue
        item = {"row": i, "content": _strip_html(q.findtext("questiontext/text") or ""), "options": [], "answer": "",
                "chapter": "", "bloom": "", "clo": "", "errors": []}
        if qtype != "multichoice":
            item["errors"].append(f"Bỏ qua câu loại '{qtype}' (chỉ hỗ trợ trắc nghiệm một đáp án)")
            out.append(item)
            continue
        if (q.findtext("single") or "true").strip().lower() in ("false", "0"):
            item["errors"].append("Câu nhiều đáp án đúng – hệ thống chỉ hỗ trợ một đáp án đúng")
        for j, a in enumerate(q.findall("answer")):
            item["options"].append(_strip_html(a.findtext("text") or ""))
            try:
                frac = float(a.get("fraction") or 0)
            except ValueError:
                frac = 0
            if frac >= 99.9 and j < MAX_OPTIONS:
                item["answer"] = LABELS[j]
        for t in q.findall("tags/tag/text"):   # thẻ Moodle kiểu "CLO2", "Bloom:3", "Chuong:2"
            tag = (t.text or "").strip()
            k, _, v = tag.partition(":")
            nk = norm(k)
            if nk.startswith("clo") and not v:
                item["clo"] = (item["clo"] + ";" if item["clo"] else "") + tag
            elif nk in ("clo", "cdr"):
                item["clo"] = (item["clo"] + ";" if item["clo"] else "") + v
            elif nk == "bloom":
                item["bloom"] = v
            elif nk in ("chuong", "chapter"):
                item["chapter"] = v
        out.append(item)
    return out


# ------------------------------------------------------------------ giải nghĩa Chương / Bloom / CLO
def parse_chapter(v: str, chapters: dict[int, int]) -> tuple[int | None, str | None]:
    """'3', 'Chương 3', '3. Trigger...' → outline_id. Trả (outline_id, lỗi)."""
    if not cell(v):
        return None, None
    m = re.search(r"\d+", str(v))
    if not m:
        return None, f"Chương '{v}' không hợp lệ (ghi số chương, VD: 3)"
    n = int(m.group())
    if n not in chapters:
        return None, (f"Môn học không có Chương {n} (có: {', '.join(map(str, sorted(chapters)))})" if chapters
                      else "Môn học chưa khai báo đề cương (chương) – để trống cột Chương")
    return chapters[n], None


def parse_bloom(v: str, levels: list[dict]) -> tuple[int | None, str | None]:
    """'3', '3 - Vận dụng', 'Vận dụng', 'Apply' → bloom_level_id."""
    s = cell(v)
    if not s:
        return None, None
    m = re.match(r"^\s*([1-6])(?!\d)", s)   # '3', '3 - Vận dụng', '3.'; không nhận '30'
    if m:
        return int(m.group(1)), None
    n = norm(s)
    for b in levels:
        if n in (norm(b["name_vi"]), norm(b["code"])):
            return b["id"], None
    return None, f"Mức Bloom '{s}' không hợp lệ (1–6 hoặc Nhớ, Hiểu, Vận dụng, Phân tích, Đánh giá, Sáng tạo)"


def parse_clos(v: str, clos: dict[str, int]) -> tuple[list[tuple[int, float]] | None, str | None]:
    """'CLO2' | 'CLO1; CLO3' (chia đều) | 'CLO1:0.6; CLO3:0.4' | 'CLO1(0.6), CLO3(0.4)' → [(clo_id, weight)]."""
    s = cell(v)
    if not s:
        return None, None
    s = re.sub(r"([:=(]\s*\d+),(\d+)", r"\1.\2", s)   # trọng số viết dấu phẩy thập phân: CLO1:0,6
    parts = [p.strip() for p in re.split(r"[;,\n]+|\s+và\s+", s) if p.strip()]
    items, weights = [], []
    for p in parts:
        m = re.match(r"^([A-Za-zĐđ]*\s*\d+)\s*(?:[:=\(]\s*(\d+(?:\.\d+)?|\.\d+)\s*\)?)?\s*$", p)
        if not m:
            return None, f"Không hiểu '{p}' (ghi VD: CLO2 hoặc CLO1:0.6; CLO3:0.4)"
        code = m.group(1).replace(" ", "")
        key = norm(code)
        if key not in clos:                       # "2", "CĐR2", "clo 2" → CLO2
            key = "clo" + re.search(r"\d+", code).group()
            code = key.upper()
        if key not in clos:
            return None, (f"Môn học không có {code.upper()} (có: {', '.join(sorted(k.upper() for k in clos))})" if clos
                          else "Môn học chưa khai báo CLO – để trống cột CLO")
        if any(c == clos[key] for c, _ in items):
            return None, f"{code.upper()} bị ghi hai lần"
        items.append((clos[key], None))
        weights.append(float(m.group(2)) if m.group(2) else None)
    if all(w is None for w in weights):
        n = len(items)
        base = round(1 / n, 2)
        ws = [base] * (n - 1) + [round(1 - base * (n - 1), 2)]
    elif any(w is None for w in weights):
        return None, "Ghi trọng số cho tất cả CLO hoặc không ghi cho CLO nào (khi đó chia đều)"
    else:
        ws = weights
        if any(not 0 < w <= 1 for w in ws):
            return None, "Trọng số CLO phải trong khoảng (0; 1]"
        if abs(sum(ws) - 1) > 0.011:
            return None, f"Tổng trọng số CLO = {sum(ws):g}, phải bằng 1"
    ws = [round(w, 2) for w in ws]
    ws[-1] = round(1 - sum(ws[:-1]), 2)          # làm tròn 2 chữ số nhưng tổng vẫn đúng bằng 1
    if ws[-1] <= 0:
        return None, "Trọng số CLO không hợp lệ sau khi làm tròn 2 chữ số"
    return [(c, w) for (c, _), w in zip(items, ws)], None


def clo_text(links: list[tuple[str, float]]) -> str:
    """[('CLO1', 1.0)] → 'CLO1'; nhiều CLO → 'CLO1:0.6; CLO3:0.4' (định dạng đọc lại được)."""
    if not links:
        return ""
    if len(links) == 1:
        return links[0][0]
    return "; ".join(f"{c}:{w:g}" for c, w in links)
