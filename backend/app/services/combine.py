"""Gộp 2, 3… bài kiểm tra (cùng môn, có thể khác lớp HP của GV) để đánh giá CĐR môn học (CLO) và CĐR CTĐT (PI/PLO).

Hai cách gộp:
* merge (mặc định – "cộng điểm theo SV"): điểm CLO của mỗi SV = Σ điểm đạt các câu đo CLO / Σ điểm tối đa của các câu đó
  trên mọi bài được chọn; SV đạt CLO nếu tỷ lệ ≥ ngưỡng đạt (BM6a). Tỷ lệ đạt CLO = số SV đạt / số SV có dữ liệu CLO.
* cumulative ("cộng dồn lượt" – BM6d): mỗi bài là một lần đánh giá; số lượt đánh giá và số lượt đạt cộng dồn qua các bài.
PLO/PI: cộng dồn (số đánh giá, số đạt) của các CLO đóng góp cho PLO (ma trận CLO–PLO) – cùng quy tắc với BM3c/BM2.
"""
from __future__ import annotations

from collections import defaultdict

from sqlalchemy import bindparam, text
from sqlalchemy.orm import Session

from ..config import settings
from .results import clo_thresholds


def combined(db: Session, exam_ids: list[int], mode: str = "merge") -> dict:
    q = text("""SELECT e.id, e.exam_title, e.assessment_type, e.exam_type, e.status, e.max_score, e.class_section_id,
                       cs.section_code, cs.course_id, cs.semester_id, s.name AS semester,
                       (SELECT SUM(points) FROM exam_questions WHERE exam_id=e.id) AS raw_max
                FROM exams e JOIN class_sections cs ON cs.id=e.class_section_id JOIN semesters s ON s.id=cs.semester_id
                WHERE e.id IN :ids ORDER BY e.id""").bindparams(bindparam("ids", expanding=True))
    exams = [dict(r) for r in db.execute(q, {"ids": exam_ids}).mappings()]
    if len(exams) != len(set(exam_ids)):
        raise ValueError("Có bài kiểm tra không tồn tại")
    if len({e["course_id"] for e in exams}) > 1:
        raise ValueError("Chỉ gộp các bài kiểm tra của cùng một môn học")
    if any(e["status"] != "Analyzed" for e in exams):
        bad = ", ".join(e["exam_title"] for e in exams if e["status"] != "Analyzed")
        raise ValueError(f"Bài chưa phân tích (cần đồng bộ/nhập điểm và phân tích trước): {bad}")
    course_id = exams[0]["course_id"]
    thr, tgt, _ = clo_thresholds(db, exams[0]["class_section_id"])
    clos = [dict(r) for r in db.execute(text("SELECT id, clo_code, description FROM clos WHERE course_id=:c ORDER BY clo_code"),
                                        {"c": course_id}).mappings()]
    q = text("""SELECT a.exam_id, a.student_id, s.student_code, s.full_name, r.clo_id, r.score_earned, r.score_max, r.is_achieved
                FROM attempt_clo_results r JOIN exam_attempts a ON a.id=r.attempt_id JOIN students s ON s.id=a.student_id
                WHERE a.exam_id IN :ids AND a.status='finished'""").bindparams(bindparam("ids", expanding=True))
    rows = [dict(r) for r in db.execute(q, {"ids": exam_ids}).mappings()]
    q = text("""SELECT a.exam_id, a.student_id, a.total_score, a.status FROM exam_attempts a WHERE a.exam_id IN :ids""").bindparams(
        bindparam("ids", expanding=True))
    totals = defaultdict(dict)
    for r in db.execute(q, {"ids": exam_ids}).mappings():
        e = next(x for x in exams if x["id"] == r["exam_id"])
        totals[r["student_id"]][r["exam_id"]] = None if r["status"] == "absent" else \
            round(float(r["total_score"]) * float(e["max_score"]) / float(e["raw_max"] or 1), 2)

    students: dict[int, dict] = {}
    agg = defaultdict(lambda: [0.0, 0.0])           # (sid, clo) -> earned, max
    per_exam = defaultdict(lambda: [0, 0])           # (exam, clo) -> n, ok   (cumulative)
    for r in rows:
        students.setdefault(r["student_id"], {"student_id": r["student_id"], "student_code": r["student_code"], "full_name": r["full_name"]})
        a = agg[(r["student_id"], r["clo_id"])]
        a[0] += float(r["score_earned"]); a[1] += float(r["score_max"])
        pe = per_exam[(r["exam_id"], r["clo_id"])]
        pe[0] += 1; pe[1] += int(bool(r["is_achieved"]))
    clo_stats = {}
    for c in clos:
        cid = c["id"]
        th = thr.get(cid, settings.default_pass_threshold)
        if mode == "cumulative":
            n = sum(v[0] for (e, k), v in per_exam.items() if k == cid)
            ok = sum(v[1] for (e, k), v in per_exam.items() if k == cid)
        else:
            vals = [(e, m) for (sid, k), (e, m) in agg.items() if k == cid and m > 0]
            n = len(vals)
            ok = sum(1 for e, m in vals if e + 1e-9 >= m * th / 100)
        pct = round(100 * ok / n, 2) if n else None
        t = tgt.get(cid, settings.default_clo_target)
        clo_stats[cid] = {"clo_id": cid, "clo_code": c["clo_code"], "description": c["description"], "threshold": th, "target": t,
                          "n_evaluated": n, "n_achieved": ok, "achieved_pct": pct, "is_achieved": None if pct is None else pct + 1e-9 >= t,
                          "by_exam": [{"exam_id": e["id"], "n": per_exam[(e["id"], cid)][0], "ok": per_exam[(e["id"], cid)][1]} for e in exams]}
    out_students = []
    for sid, s in sorted(students.items(), key=lambda x: x[1]["student_code"]):
        cl = {}
        for c in clos:
            e, m = agg.get((sid, c["id"]), (0, 0))
            if m > 0:
                pct = round(100 * e / m, 2)
                cl[c["clo_code"]] = {"pct": pct, "ok": e + 1e-9 >= m * clo_stats[c["id"]]["threshold"] / 100,
                                     "earned": round(e, 2), "max": round(m, 2)}
        out_students.append({**s, "clos": cl, "scores": {str(k): v for k, v in totals.get(sid, {}).items()},
                             "n_ok": sum(1 for v in cl.values() if v["ok"]), "n_clo": len(cl)})
    # PLO / PI theo các CTĐT có môn học
    maps = [dict(r) for r in db.execute(text("""SELECT m.clo_id, m.level, p.id AS plo_id, p.plo_code, p.description, p.target_pct,
                                                     pr.id AS program_id, pr.code AS program_code, pr.name AS program_name
                                                FROM clo_plo_mapping m JOIN plos p ON p.id=m.plo_id JOIN programs pr ON pr.id=p.program_id
                                                JOIN clos c ON c.id=m.clo_id WHERE c.course_id=:c"""), {"c": course_id}).mappings()]
    pis = [dict(r) for r in db.execute(text("""SELECT pi.id, pi.pi_code, pi.description, pi.plo_id, pc.level FROM pi_courses pc
                                               JOIN performance_indicators pi ON pi.id=pc.pi_id WHERE pc.course_id=:c"""), {"c": course_id}).mappings()]
    programs = {}
    for m in maps:
        st = clo_stats.get(m["clo_id"])
        if not st or not st["n_evaluated"]:
            continue
        pr = programs.setdefault(m["program_id"], {"program_id": m["program_id"], "code": m["program_code"], "name": m["program_name"], "plos": {}})
        p = pr["plos"].setdefault(m["plo_id"], {"plo_id": m["plo_id"], "plo_code": m["plo_code"], "description": m["description"],
                                                 "target": float(m["target_pct"]), "clos": [], "n_evaluated": 0, "n_achieved": 0})
        p["clos"].append(f"{st['clo_code']} ({m['level']})")
        p["n_evaluated"] += st["n_evaluated"]; p["n_achieved"] += st["n_achieved"]
    for pr in programs.values():
        lst = []
        for p in sorted(pr["plos"].values(), key=lambda x: (len(x["plo_code"]), x["plo_code"])):
            p["achieved_pct"] = round(100 * p["n_achieved"] / p["n_evaluated"], 2)
            p["is_achieved"] = p["achieved_pct"] + 1e-9 >= p["target"]
            p["pis"] = [{"pi_code": x["pi_code"], "description": x["description"], "level": x["level"]} for x in pis if x["plo_id"] == p["plo_id"]]
            lst.append(p)
        pr["plos"] = lst
    return {"mode": mode, "course_id": course_id, "exams": [{**e, "max_score": float(e["max_score"]), "raw_max": float(e["raw_max"] or 0)} for e in exams],
            "clos": list(clo_stats.values()), "students": out_students, "programs": list(programs.values())}


def combined_xlsx(data: dict, title: str) -> bytes:
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active; ws.title = "TongHop_CLO"
    bold = Font(bold=True); fill = PatternFill("solid", fgColor="DCE6F1")
    ws.append([title]); ws["A1"].font = Font(bold=True, size=12)
    ws.append(["Các bài gộp: " + "; ".join(f"{e['exam_title']} ({e['section_code']})" for e in data["exams"])])
    ws.append(["Cách gộp: " + ("cộng điểm các câu đo CLO của từng SV trên mọi bài" if data["mode"] == "merge" else "cộng dồn lượt đánh giá qua các bài (BM6d)")])
    ws.append([])
    ws.append(["CLO", "Nội dung", "Ngưỡng đạt (%)", "Số SV đánh giá", "Số SV đạt", "Tỷ lệ đạt (%)", "Chỉ tiêu (%)", "Kết luận"])
    for c in ws[5]:
        c.font = bold; c.fill = fill
    for c in data["clos"]:
        ws.append([c["clo_code"], c["description"], c["threshold"], c["n_evaluated"], c["n_achieved"], c["achieved_pct"], c["target"],
                   "" if c["is_achieved"] is None else ("Đạt" if c["is_achieved"] else "Không đạt")])
    for pr in data["programs"]:
        ws.append([]); ws.append([f"CĐR CTĐT – {pr['code']} {pr['name']}"]); ws.cell(ws.max_row, 1).font = bold
        ws.append(["PLO", "Nội dung", "CLO đóng góp", "Số lượt đánh giá", "Số lượt đạt", "Tỷ lệ đạt (%)", "Chỉ tiêu (%)", "Kết luận", "PI của học phần"])
        for c in ws[ws.max_row]:
            c.font = bold; c.fill = fill
        for p in pr["plos"]:
            ws.append([p["plo_code"], p["description"], ", ".join(p["clos"]), p["n_evaluated"], p["n_achieved"], p["achieved_pct"], p["target"],
                       "Đạt" if p["is_achieved"] else "Không đạt", "; ".join(f"{x['pi_code']} ({x['level'] or ''})" for x in p["pis"])])
    for i, w in enumerate([10, 60, 18, 14, 12, 13, 12, 12, 30], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=6):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    s = wb.create_sheet("TungSinhVien")
    codes = [c["clo_code"] for c in data["clos"]]
    s.append(["MSSV", "Họ và tên"] + [f"{e['exam_title']} ({e['section_code']})" for e in data["exams"]] + [f"{c} (%)" for c in codes] + ["Số CLO đạt"])
    for c in s[1]:
        c.font = bold; c.fill = fill; c.alignment = Alignment(wrap_text=True)
    for st in data["students"]:
        s.append([st["student_code"], st["full_name"]] + [st["scores"].get(str(e["id"])) for e in data["exams"]]
                 + [st["clos"].get(c, {}).get("pct") for c in codes] + [f"{st['n_ok']}/{st['n_clo']}"])
    s.column_dimensions["A"].width = 13; s.column_dimensions["B"].width = 26
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()
