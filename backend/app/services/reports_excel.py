"""Xuất báo cáo đo lường CĐR theo đúng bố cục các biểu mẫu HCMUTE đang dùng.

* build_bm6(class_section_id)  -> tệp "Biểu mẫu 6": mục lục, BM6a (kế hoạch + nhận xét), BM6b (kết quả, công thức
  liên kết tới BM6c/6d), BM6c/6d (minh chứng từng CLO – mỗi bài KT một khối, COUNTIF theo ngưỡng)
* build_bm6_exam(exam_id)      -> BM6c riêng của MỘT bài kiểm tra: trang tổng hợp các CLO bài đo + BM6c từng CLO (kèm MSSV, họ tên)
* build_bm3(program_id, year)  -> BM2a (kế hoạch đo CĐR CTĐT), BM2b (báo cáo tổng kết), BM3b/BM3c cho từng CĐR đo trong năm
* build_assignments(semester)  -> bảng "Phân công đánh giá PIs" của học kỳ

Mỗi biểu mẫu gồm hai bước: hàm *_data đọc CSDL thành dict, hàm render_* dựng tệp từ dict (kiểm thử được không cần CSDL).
"""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import text
from sqlalchemy.orm import Session

from ..config import settings

FONT = "Times New Roman"
BODY, TITLE = 13, 14                      # cỡ chữ của biểu mẫu gốc: nội dung 13, tiêu đề 14
THIN = Side(style="thin", color="000000")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="center")
LEFT_TOP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
DATE_LINE = "TP. Hồ Chí Minh, ngày … tháng … năm 20…"   # để trống cho người ký tự điền, như biểu mẫu
SCHOOL = "TRƯỜNG ĐẠI HỌC CÔNG NGHỆ KỸ THUẬT\nTHÀNH PHỐ HỒ CHÍ MINH"
ALL_SUPERVISORS = "Tất cả thầy/cô có hướng dẫn"
TYPE_NAME = {"process": "Quá trình", "final": "Cuối kỳ"}
EV_NAME = {"process": "Quá trình", "final": "Cuối kỳ", "any": "Quá trình + Cuối kỳ", None: ""}

# Hướng dẫn in nghiêng dưới mỗi mục nhận xét của BM6a / BM3b (nguyên văn biểu mẫu) – dùng khi mục chưa có nội dung
GUIDE = {
    "Tổng hợp dữ liệu": "Mô tả kỹ cách tính số liệu tổng hợp từ dữ liệu của từng CĐR và cung cấp con số tính tổng. Sử dụng thêm "
                        "các đồ thị, công thức nếu cần thiết nhưng nhớ trích dẫn cụ thể số lượng sinh viên đã được đánh giá "
                        "trong khi đánh giá từng CĐR",
    "Đánh giá kết quả": "Dựa trên số liệu đã tổng hợp, phân tích đánh giá để mô tả mức độ mà CĐR này đã đạt được. Sử dụng biểu "
                        "đồ/ đồ thị với các mô tả chi tiết cho các con số.",
    "Những hành động": "Mô tả ngắn gọn các hành động mà chương trình đã thực hiện dẫn đến kết quả hiện thời. Chỉ ra các hành "
                       "động dự tính sẽ cải tiến trong tương lai (nếu cần thiết).",
    "Kết quả của các": "Mô tả ngắn gọn kết quả của bất kỳ thay đổi nào đã thực hiện (dù là có hiệu quả hay không có hiệu quả) "
                       "khi thực hiện đánh giá lại CĐR này.",
    "Công cụ đánh giá": "Dữ liệu của các đánh giá và các kết quả được văn bản hóa và lưu trữ như thế nào? Đính kèm bản copy của "
                        "các minh chứng cùng với bản này. Đính kèm mẫu một vài bài làm của SV với các mức khác nhau (kém, TB, "
                        "khá, …). Có thể tách ra thành bản phụ lục riêng",
}
ASSIGN_NOTE = "- Đánh giá PIs tất cả các CĐR của môn học\n- Các môn học làm đề tài cuối kỳ các thầy/cô cần lưu minh chứng rubric"


def _f(bold=False, size=BODY, italic=False):
    return Font(name=FONT, size=size, bold=bold, italic=italic)


def _put(ws, ref: str, value, *, bold=False, italic=False, size=BODY, align=WRAP, merge: str | None = None, box=False):
    if merge:
        ws.merge_cells(merge)
    c = ws[ref]
    c.value = value
    c.font = _f(bold, size, italic); c.alignment = align
    if box:
        c.border = BOX
    return c


def _box_range(ws, rng: str) -> None:
    for row in ws[rng]:
        for c in row:
            c.border = BOX


def _header(ws, unit: str, last_col: int, *, left_cols: int | None = None, doc_no: str | None = None) -> None:
    """Khối tiêu đề của mọi biểu mẫu: trường + đơn vị bên trái, quốc hiệu + ngày tháng bên phải."""
    lc = get_column_letter(last_col)
    split = left_cols or max(3, last_col // 2)
    lm, rs = get_column_letter(split), get_column_letter(split + 1)
    _put(ws, "A1", SCHOOL, align=CENTER, merge=f"A1:{lm}1")
    _put(ws, "A2", unit.upper(), bold=True, align=CENTER, merge=f"A2:{lm}2")
    if doc_no:
        _put(ws, "A3", doc_no, align=CENTER, merge=f"A3:{lm}3")
    _put(ws, f"{rs}1", "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nĐộc lập – Tự do – Hạnh phúc", bold=True, align=CENTER, merge=f"{rs}1:{lc}1")
    _put(ws, f"{rs}3", DATE_LINE, italic=True, align=CENTER, merge=f"{rs}3:{lc}3")
    ws.row_dimensions[1].height = 36


def _lines(ws, row: int, last_col: int, lines: list[str], bold_prefix=("Kết luận",)) -> int:
    lc = get_column_letter(last_col)
    for line in lines:
        _put(ws, f"A{row}", line, bold=line.startswith(bold_prefix), merge=f"A{row}:{lc}{row}")
        row += 1
    return row


def _head_cells(ws, row: int, headers: list[str]) -> None:
    for j, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=j, value=h)
        c.font = _f(True); c.alignment = CENTER; c.border = BOX


def _row(ws, row: int, values: list, *, wrap_cols=(2,), pct_cols=(), bold=False) -> None:
    for j, v in enumerate(values, 1):
        c = ws.cell(row=row, column=j, value=v)
        c.font = _f(bold); c.border = BOX
        c.alignment = WRAP if j in wrap_cols else CENTER
        if j in pct_cols:   # chỉ tiêu (số tròn) 0%, tỷ lệ tính ra (công thức / số lẻ) 0.00%
            frac = isinstance(v, str) and v.startswith("=") or isinstance(v, float) and abs(v * 100 - round(v * 100)) > 1e-9
            c.number_format = "0.00%" if frac else "0%"


def _narratives(ws, row: int, last_col: int, items: list[tuple[str, str | None]], unit_word: str = "CĐR") -> int:
    """Các mục nhận xét: tiêu đề in đậm + nội dung (chưa có nội dung thì in nghiêng lời hướng dẫn của biểu mẫu)."""
    lc = get_column_letter(last_col)
    for title, body in items:
        _put(ws, f"A{row}", title, bold=True, merge=f"A{row}:{lc}{row}")
        row += 1
        guide = next((g for k, g in GUIDE.items() if title.startswith(k)), "").replace("từng CĐR", f"từng {unit_word}")
        _put(ws, f"A{row}", body or guide, italic=not body, align=LEFT_TOP, merge=f"A{row}:{lc}{row}")
        ws.row_dimensions[row].height = max(48, 17 * (1 + len(body or guide) // 140))
        row += 2
    return row


def _signer(ws, row: int, first_col: int, last_col: int, title: str) -> int:
    rng = f"{get_column_letter(first_col)}{row}:{get_column_letter(last_col)}{row}"
    _put(ws, f"{get_column_letter(first_col)}{row}", f"{title}\n(Ký, ghi rõ Họ và Tên)", bold=True, align=CENTER, merge=rng)
    ws.row_dimensions[row].height = 36
    return row + 1


def _widths(ws, widths: list[float]) -> None:
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def _print_setup(wb, landscape=True) -> None:
    """In A4 (ngang như biểu mẫu), vừa 1 trang bề ngang."""
    for ws in wb.worksheets:
        ws.page_setup.orientation = "landscape" if landscape else "portrait"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = ws.page_margins.right = 0.4


def _pct(x) -> float | str:
    return "" if x is None else round(float(x) / 100, 4)


def _ref(sheet: str, cell: str) -> str:
    return "='" + sheet.replace("'", "''") + "'!" + cell


def _save(wb) -> bytes:
    bio = io.BytesIO(); wb.save(bio)
    return bio.getvalue()


# =====================================================================================  BM6 – CĐR môn học


def bm6_data(db: Session, cs_id: int) -> dict:
    cs = dict(db.execute(text("""SELECT cs.id, cs.section_code, co.id AS course_id, co.course_code, co.course_name, co.clo_target_pct,
                                        s.id AS semester_id, s.name AS sem_name, s.academic_year, l.full_name AS lecturer, l.department
                                 FROM class_sections cs JOIN courses co ON co.id = cs.course_id
                                 JOIN semesters s ON s.id = cs.semester_id JOIN lecturers l ON l.id = cs.lecturer_id
                                 WHERE cs.id = :cs"""), {"cs": cs_id}).mappings().one())
    clos = [dict(c) for c in db.execute(text("""
        SELECT c.id, c.clo_code, c.description, p.assessments_text, p.evidence_type, p.method, p.cycle,
               COALESCE(p.pass_threshold_pct, :thr0) AS thr, COALESCE(p.target_pct, :t) AS target
        FROM clos c LEFT JOIN clo_assessment_plans p ON p.clo_id = c.id AND p.semester_id = :s
        WHERE c.course_id = :c ORDER BY c.clo_code"""),
        {"c": cs["course_id"], "s": cs["semester_id"], "t": cs["clo_target_pct"], "thr0": settings.default_pass_threshold}).mappings()]
    res = {x["clo_id"]: dict(x) for x in db.execute(text("SELECT * FROM clo_results WHERE class_section_id = :cs"), {"cs": cs_id}).mappings()}
    for c in clos:
        c["result"] = res.get(c["id"])
        ev = c["evidence_type"] or "any"
        c["tests"] = [t for e in _analyzed_exams(db, cs_id) if ev == "any" or e["assessment_type"] == ev
                      for t in [_exam_test(db, e, c["id"])] if t]
    return {"cs": cs, "clos": clos}


def _analyzed_exams(db: Session, cs_id: int, exam_id: int | None = None):
    return db.execute(text("""SELECT id, exam_title, assessment_type, max_score,
                                     (SELECT SUM(points) FROM exam_questions WHERE exam_id = e.id) AS raw_max
                              FROM exams e WHERE class_section_id = :cs AND status = 'Analyzed'"""
                           + (" AND id = :e" if exam_id else "") + " ORDER BY id"), {"cs": cs_id, "e": exam_id}).mappings().all()


def exam_label(e) -> str:
    t = TYPE_NAME.get(e["assessment_type"])
    return e["exam_title"] + (f" ({t})" if t and t.lower() not in e["exam_title"].lower() else "")


def _exam_test(db: Session, e, clo_id: int) -> dict | None:
    """Điểm phần câu hỏi của một CLO trong một bài KT cho từng SV đã làm bài (đúng tập dữ liệu tính BM6b)."""
    pts = db.execute(text("""SELECT r.score_earned, r.score_max, r.is_achieved, a.total_score,
                                    COALESCE(st.student_code, '') AS code, COALESCE(st.full_name, '') AS name
                             FROM attempt_clo_results r JOIN exam_attempts a ON a.id = r.attempt_id
                             LEFT JOIN students st ON st.id = a.student_id
                             WHERE a.exam_id = :e AND r.clo_id = :c AND a.status = 'finished'
                             ORDER BY st.student_code, a.id"""), {"e": e["id"], "c": clo_id}).mappings().all()
    if not pts:
        return None
    emax, raw_max = float(e["max_score"]), float(e["raw_max"] or 0) or 1
    return {"title": exam_label(e), "type": e["assessment_type"], "emax": emax, "qmax": float(pts[0]["score_max"]),
            "n_ok": sum(1 for p in pts if p["is_achieved"]),
            "rows": [{"q": float(p["score_earned"]), "qmax": float(p["score_max"]),
                      "s": round(float(p["total_score"] or 0) * emax / raw_max, 2), "code": p["code"], "name": p["name"]} for p in pts]}


def render_bm6(d: dict) -> bytes:
    cs, clos = d["cs"], d["clos"]
    unit = cs["department"] or "Bộ môn"
    course_line = f"Tên môn học: {cs['course_name']} ({cs['course_code']}) – Lớp HP: {cs['section_code']} – {cs['sem_name']} – GV: {cs['lecturer']}"
    target_line = f"Chỉ tiêu đạt CĐR môn học: {float(cs['clo_target_pct']):.0f}%"
    wb = Workbook()

    # ---------------- Mục lục "Biểu mẫu 6"
    ws = wb.active; ws.title = "BM6"
    _put(ws, "A1", "Biểu mẫu 6", bold=True, align=Alignment(vertical="center"))
    _put(ws, "A2", f"DỮ LIỆU ĐO LƯỜNG CHI TIẾT CĐR MÔN HỌC {cs['course_name'].upper()}", bold=True, size=TITLE, align=CENTER, merge="A2:K2")
    _head_cells(ws, 4, ["STT", "Biểu mẫu", "Nội dung biểu mẫu"]); ws.merge_cells("C4:K4"); _box_range(ws, "C4:K4")
    for i, (code, desc) in enumerate([
            ("Biểu mẫu 6a", "Bảng kế hoạch kiểm tra, đánh giá mức độ đạt cho từng chuẩn đầu ra môn học"),
            ("Biểu mẫu 6b", "Kết quả tổng hợp của từng chuẩn đầu ra môn học"),
            ("Biểu mẫu 6c", "Bảng minh chứng đo lường từ kết quả kiểm tra, đánh giá \n(Nếu CĐR chỉ cần dùng kết quả đánh giá của 1 bài kiểm tra)"),
            ("Biểu mẫu 6d", "Bảng minh chứng đo lường từ kết quả kiểm tra, đánh giá \n(Trường hợp CĐR nào đó cần lấy mẫu từ nhiều bài kiểm tra)")], 5):
        _row(ws, i, [i - 4, code, desc], wrap_cols=(3,))
        ws.merge_cells(f"C{i}:K{i}"); _box_range(ws, f"C{i}:K{i}"); ws.row_dimensions[i].height = 48
    _widths(ws, [14.6, 15.6] + [9.1] * 9)

    # ---------------- BM6a – kế hoạch + 5 mục nhận xét + chữ ký (bố cục sheet BM6a của biểu mẫu)
    ws = wb.create_sheet("BM6a_KH-KQ CĐR môn học")
    _header(ws, unit, 8, left_cols=4)
    _put(ws, "A5", "BẢNG KẾ HOẠCH KIỂM TRA, ĐÁNH GIÁ MỨC ĐỘ ĐẠT CHO TỪNG CHUẨN ĐẦU RA MÔN HỌC", bold=True, size=TITLE, align=CENTER, merge="A5:H5")
    _lines(ws, 6, 8, [course_line, target_line])
    _head_cells(ws, 9, ["STT", "Nội dung CĐR môn học", "Các bài KT có CĐR xuất hiện", "Bài KT sẽ lấy minh chứng",
                        "Phương pháp kiểm tra, đánh giá", "Chu kỳ lấy minh chứng", "Thời gian lấy minh chứng", "Chỉ tiêu mong muốn"])
    r = 10
    for i, c in enumerate(clos, 1):
        _row(ws, r, [i, f"{c['clo_code']}: {c['description']}", c["assessments_text"] or "", EV_NAME[c["evidence_type"]],
                     c["method"] or "", c["cycle"] or "", cs["sem_name"], _pct(c["target"])], pct_cols=(8,))
        r += 1
    _row(ws, r, ["CHỈ TIÊU MONG MUỐN ĐẠT CĐR MÔN HỌC", "", "", "", "", "", "", _pct(cs["clo_target_pct"])], pct_cols=(8,), bold=True)
    ws.merge_cells(f"A{r}:G{r}"); ws[f"A{r}"].alignment = CENTER
    measured = [c for c in clos if c["result"]]
    tot_ok = sum(c["result"]["n_achieved"] for c in measured); tot_n = sum(c["result"]["n_evaluated"] for c in measured)
    notes = "\n".join(f"{c['clo_code']}: {c['result']['analysis']}" for c in measured if c["result"]["analysis"])
    impr = "\n".join(f"{c['clo_code']}: {c['result']['improvement']}" for c in measured if c["result"]["improvement"])
    summary = None
    if tot_n:
        summary = (f"Tổng hợp từ {tot_n} lượt SV được đánh giá trên các bài KT lấy minh chứng của {len(measured)} CĐR; {tot_ok} lượt đạt "
                   f"({100 * tot_ok / tot_n:.2f}%). Mỗi CĐR: SV đạt khi điểm các câu hỏi thuộc CĐR ≥ ngưỡng % điểm tối đa của các câu đó; "
                   "tỷ lệ của CĐR = tổng số SV đạt / tổng số SV đánh giá qua các bài KT (chi tiết ở các sheet BM6c/BM6d).")
    r = _narratives(ws, r + 2, 8, [
        ("Tổng hợp dữ liệu đã đánh giá cho CĐR này", summary),
        ("Đánh giá kết quả của số liệu tổng hợp", notes or None),
        ("Những hành động cải tiến", impr or None),
        ("Kết quả của các cải tiến đã thực hiện", None),
        ("Công cụ đánh giá", "Bài KT trắc nghiệm (Moodle Quiz / bài giấy chấm bằng Moodle Offline Quiz) – câu hỏi và đáp án; "
                             "ma trận câu hỏi – CĐR và kết quả từng câu của từng SV được lưu trong hệ thống." if tot_n else None),
    ])
    _signer(ws, r + 1, 4, 8, "TRƯỞNG ĐƠN VỊ")
    _widths(ws, [7.7, 42, 16.4, 14, 16, 12, 13, 11])

    # ---------------- BM6c / BM6d – dựng trước để BM6b dùng công thức liên kết như biểu mẫu
    ws6b = wb.create_sheet("BM6b_KH-KQ CĐR môn học")
    links = {}
    for c in clos:
        name = f"BM6{'d' if len(c['tests']) > 1 else 'c'}_{c['clo_code']}"[:31]
        links[c["id"]] = (name, _render_evidence(wb.create_sheet(name), c, cs, unit))

    # ---------------- BM6b – kết quả tổng hợp (tiêu đề 2 tầng, số liệu = công thức trỏ sang sheet minh chứng)
    ws = ws6b
    _header(ws, unit, 8, left_cols=3)
    _put(ws, "A5", "KẾT QUẢ TỔNG HỢP CỦA TỪNG CHUẨN ĐẦU RA MÔN HỌC", bold=True, size=TITLE, align=CENTER, merge="A5:H5")
    _lines(ws, 6, 8, [course_line, target_line])
    first, last = 12, 11 + len(clos)
    total = last + 1
    _put(ws, "A8", f'="Kết luận: "&H{total}', bold=True, merge="A8:H8")
    _head_cells(ws, 10, ["STT", "Nội dung CĐR \nmôn học", "Công cụ kiểm tra, đánh giá", "Kết quả tổng hợp của từng CĐR", "", "",
                         "Chỉ tiêu mong muốn ", "Kết quả\nđạt được"])
    _head_cells(ws, 11, ["", "", "", "Tổng số SV đạt yêu cầu", "Tổng số SV đánh giá", "Tỷ lệ % đạt yêu cầu", "", ""])
    for col in "ABCGH":
        ws.merge_cells(f"{col}10:{col}11")
    ws.merge_cells("D10:F10")
    for i, c in enumerate(clos):
        rr = first + i
        sheet, cells = links[c["id"]]
        ok = _ref(sheet, cells["ok"]) if cells else ""
        n = _ref(sheet, cells["n"]) if cells else ""
        tools = "; ".join(t["title"] for t in c["tests"]) or "Câu hỏi và đáp án"
        _row(ws, rr, [i + 1, f"{c['clo_code']}: {c['description']}", tools, ok, n,
                      f'=IF(N(E{rr})>0,D{rr}/E{rr},"")', _pct((c["result"] or {}).get("target_pct", c["target"])),
                      f'=IF(N(E{rr})=0,"Chưa đo",IF(F{rr}+1E-9>=G{rr},"Đạt","Không đạt"))'], pct_cols=(6, 7))
    _row(ws, total, ["KẾT QUẢ ĐẠT ĐƯỢC CĐR MÔN HỌC", "", "", f"=SUM(D{first}:D{last})", f"=SUM(E{first}:E{last})",
                     f'=IF(E{total}>0,D{total}/E{total},"")', _pct(cs["clo_target_pct"]),
                     f'=IF(E{total}=0,"Chưa đo",IF(F{total}+1E-9>=G{total},"Đạt","Không đạt"))'], pct_cols=(6, 7), bold=True)
    ws.merge_cells(f"A{total}:C{total}"); ws[f"A{total}"].alignment = CENTER
    _signer(ws, total + 3, 5, 8, "TRƯỞNG ĐƠN VỊ")
    _widths(ws, [7.7, 42, 18, 15.7, 15.7, 15.7, 12, 13])
    _print_setup(wb)
    return _save(wb)


def _render_evidence(ws, c: dict, cs: dict, unit: str, *, with_names=False, extra: list[str] | None = None,
                     concl_suffix: str = "") -> dict | None:
    """BM6c (1 bài KT) / BM6d (nhiều bài KT): mỗi bài một khối cột TT | Điểm câu hỏi | Điểm bài thi/KT, cách nhau 1 cột.

    with_names: thêm cột MSSV, Họ và tên (tệp BM6c xuất riêng cho một bài KT). Khi điểm tối đa phần câu hỏi của CLO khác nhau
    giữa các SV (đề rút ngẫu nhiên / nhiều mã đề), khối có thêm cột "Điểm tối đa câu hỏi" và số SV đạt tính theo từng dòng.
    Trả về ô 'tổng số SV đạt' và 'tổng số SV đánh giá' để BM6b liên kết."""
    tests = c["tests"]
    blocks = []
    for t in tests:
        varying = len({r["qmax"] for r in t["rows"]}) > 1
        cols = ["TT"] + (["MSSV", "Họ và tên"] if with_names else []) + ["Điểm câu hỏi "] + (["Điểm tối đa câu hỏi"] if varying else []) + ["Điểm bài thi/KT"]
        blocks.append((t, varying, cols))
    width = sum(len(cols) + 1 for _, _, cols in blocks) - 1
    last_col = max(11, width)
    _header(ws, unit, last_col, left_cols=5)
    lc = get_column_letter(last_col)
    _put(ws, "A5", "BẢNG MINH CHỨNG ĐO LƯỜNG TỪ KẾT QUẢ KIỂM TRA, ĐÁNH GIÁ", bold=True, size=TITLE, align=CENTER, merge=f"A5:{lc}5")
    thr = float(c["thr"])
    concl_row = _lines(ws, 7, last_col, [
        f"CĐR: {c['clo_code']} – {c['description']}",
        f"Chỉ tiêu: {float(c['target']):.0f}% SV đạt được {thr:.0f}% điểm tối đa của câu hỏi kiểm tra có liên quan",
        f"Môn học lấy mẫu: {cs['course_name']} ({cs['course_code']})",
        "Bài kiểm tra lấy mẫu: " + ("; ".join(t["title"] for t in tests) or "(chưa có bài KT đã phân tích)"),
        f"Năm học, HK lấy mẫu: {cs['sem_name']}", *(extra or [])])
    W = {"TT": 6.6, "MSSV": 12, "Họ và tên": 26, "Điểm câu hỏi ": 18, "Điểm tối đa câu hỏi": 14, "Điểm bài thi/KT": 14}
    if not blocks:
        _widths(ws, [6.6, 18, 14, 6.5] * 3)
        _put(ws, f"A{concl_row}", "Kết luận: Chưa đo", bold=True, merge=f"A{concl_row}:{lc}{concl_row}")
        return None
    top = concl_row + 1                                           # dòng tiêu đề khối (biểu mẫu: dòng 13)
    end = top + 2 + max(len(t["rows"]) for t in tests)            # dòng dữ liệu cuối chung của mọi khối (như BM6d)
    cnt, rate = end + 1, end + 2
    n_cells, ok_cells, col = [], [], 1
    for k, (t, varying, cols) in enumerate(blocks):
        L = lambda j, col=col: get_column_letter(col + j)
        jq, je = cols.index("Điểm câu hỏi "), len(cols) - 1
        jm = cols.index("Điểm tối đa câu hỏi") if varying else None
        _put(ws, f"{L(0)}{top}", f"Bài KT {k + 1}: {t['title']}", bold=True, align=Alignment(vertical="center"))  # tràn sang ô trống bên phải
        for j, h in enumerate(cols):
            _put(ws, f"{L(j)}{top + 1}", h, bold=True, align=CENTER, box=True)
            ws[f"{L(j)}{top + 2}"].border = BOX
            ws.column_dimensions[L(j)].width = W[h]
        ws.column_dimensions[L(je + 1)].width = 6.5
        if not varying:
            _put(ws, f"{L(jq)}{top + 2}", f"Điểm tối đa: {t['qmax']:g}", align=CENTER, box=True)
        _put(ws, f"{L(je)}{top + 2}", f"Điểm tối đa: {t['emax']:g}", align=CENTER, box=True)
        for i, r in enumerate(t["rows"], 1):
            vals = {0: i, jq: r["q"], je: r["s"]}
            if with_names:
                vals[1], vals[2] = r["code"], r["name"]
            if varying:
                vals[jm] = r["qmax"]
            for j, v in vals.items():
                cell = ws.cell(row=top + 2 + i, column=col + j, value=v); cell.font = _f(); cell.border = BOX
                cell.alignment = WRAP if with_names and j == 2 else CENTER
        rng = f"{L(jq)}{top + 3}:{L(jq)}{end}"
        n_ref = f"COUNT({L(0)}{top + 3}:{L(0)}{end})"
        if varying:   # so điểm của từng SV với điểm tối đa của chính SV đó
            f_cnt = f"=SUMPRODUCT(({rng}<>\"\")*({rng}+1E-9>={L(jm)}{top + 3}:{L(jm)}{end}*{thr:g}/100))"
        else:
            f_cnt = f'=COUNTIF({rng},">={round(t["qmax"] * thr / 100, 4):g}")'
        _put(ws, f"{L(jq)}{cnt}", f_cnt, bold=True, italic=True, align=CENTER, box=True)
        _put(ws, f"{L(jq)}{rate}", f"=IF({n_ref}>0,{L(jq)}{cnt}/{n_ref},0)", bold=True, italic=True, align=CENTER, box=True).number_format = "0.00%"
        for i, (lab, f) in enumerate([(f"Tổng số SV thực hiện bài KT{k + 1}", f"={n_ref}"),
                                      (f"Tổng số SV đạt bài KT{k + 1}", f"={L(jq)}{cnt}"),
                                      (f"Tỷ lệ % đạt KT{k + 1}", f"={L(jq)}{rate}")]):
            rr = end + 3 + i
            _put(ws, f"{L(0)}{rr}", lab, merge=f"{L(0)}{rr}:{L(je - 1)}{rr}")
            cell = _put(ws, f"{L(je)}{rr}", f, bold=True, align=CENTER)
            if i == 2:
                cell.number_format = "0.00%"
        n_cells.append(f"{L(je)}{end + 3}"); ok_cells.append(f"{L(je)}{end + 4}")
        col += len(cols) + 1
    b = end + 7
    for i, (v, bold) in enumerate([("Tổng kết cho CĐR", True), ("Tổng số SV tham gia đánh giá:", False), ("=" + "+".join(n_cells), True),
                                   (f"Tổng số SV đạt của {c['clo_code']}:", False), ("=" + "+".join(ok_cells), True),
                                   ("Tỷ lệ %:", False), (f"=IF(A{b + 2}>0,A{b + 4}/A{b + 2},0)", True)]):
        cell = _put(ws, f"A{b + i}", v, bold=bold, merge=f"A{b + i}:C{b + i}", align=Alignment(horizontal="left", vertical="center"))
        if i == 6:
            cell.number_format = "0.00%"
    _put(ws, f"A{concl_row}", f'="Kết luận: "&IF(A{b + 6}+1E-9>={float(c["target"]) / 100:g},"Đạt","Không đạt")'
         + (f'&"{concl_suffix}"' if concl_suffix else ""), bold=True, merge=f"A{concl_row}:{lc}{concl_row}")
    _signer(ws, b + 8, 7, min(last_col, 10), "TRƯỞNG ĐƠN VỊ")
    return {"n": f"A{b + 2}", "ok": f"A{b + 4}"}


def build_bm6(db: Session, cs_id: int) -> bytes:
    return render_bm6(bm6_data(db, cs_id))


def bm6_exam_data(db: Session, exam_id: int) -> dict:
    """Dữ liệu BM6c của MỘT bài KT: mọi CLO bài đo (kể cả CLO mà kế hoạch BM6a không lấy minh chứng từ loại bài này)."""
    e = db.execute(text("SELECT id, class_section_id FROM exams WHERE id = :e"), {"e": exam_id}).mappings().one()
    d = bm6_data(db, e["class_section_id"])
    ex = _analyzed_exams(db, e["class_section_id"], exam_id)
    d["exam"] = dict(ex[0]) if ex else None
    d["n_sv"] = db.execute(text("SELECT COUNT(*) FROM exam_attempts WHERE exam_id = :e AND status = 'finished'"), {"e": exam_id}).scalar()
    for c in d["clos"]:
        t = _exam_test(db, ex[0], c["id"]) if ex else None
        ev = c["evidence_type"] or "any"
        c["tests"] = [t] if t else []
        c["used"] = ev == "any" or (ex and ev == ex[0]["assessment_type"])
    return d


def render_bm6_exam(d: dict) -> bytes:
    """Tệp BM6c riêng cho một bài KT: sheet tổng hợp (kết quả từng CLO theo riêng bài này) + một sheet BM6c cho mỗi CLO bài đo."""
    cs, e = d["cs"], d["exam"]
    unit = cs["department"] or "Bộ môn"
    measured = [c for c in d["clos"] if c["tests"]]
    wb = Workbook(); ws = wb.active; ws.title = "TongHop_BaiKT"
    _header(ws, unit, 9, left_cols=4)
    _put(ws, "A5", "KẾT QUẢ ĐO LƯỜNG CHUẨN ĐẦU RA MÔN HỌC THEO BÀI KIỂM TRA", bold=True, size=TITLE, align=CENTER, merge="A5:I5")
    r = _lines(ws, 6, 9, [
        f"Tên môn học: {cs['course_name']} ({cs['course_code']}) – Lớp HP: {cs['section_code']} – {cs['sem_name']} – GV: {cs['lecturer']}",
        f"Bài kiểm tra: {exam_label(e)} – Điểm tối đa: {float(e['max_score']):g} – Số SV làm bài: {d['n_sv']}",
        f"Số CĐR môn học được đo trong bài: {len(measured)}/{len(d['clos'])}"])
    r += 1
    _head_cells(ws, r, ["STT", "Nội dung CĐR môn học", "Số SV đạt", "Số SV đánh giá", "Tỷ lệ % đạt", "Ngưỡng đạt (% điểm câu hỏi CĐR)",
                        "Chỉ tiêu mong muốn", "Kết quả (theo bài KT này)", "Dùng cho BM6b của lớp HP"])
    for i, c in enumerate(measured, 1):
        t, rr = c["tests"][0], r + i
        _row(ws, rr, [i, f"{c['clo_code']}: {c['description']}", t["n_ok"], len(t["rows"]), f'=IF(D{rr}>0,C{rr}/D{rr},"")',
                      _pct(c["thr"]), _pct(c["target"]), f'=IF(D{rr}=0,"Chưa đo",IF(E{rr}+1E-9>=G{rr},"Đạt","Không đạt"))',
                      "Có" if c["used"] else f"Không (minh chứng: {EV_NAME[c['evidence_type'] or 'any']})"], pct_cols=(5, 6, 7), wrap_cols=(2, 9))
        ws[f"E{rr}"].number_format = "0.00%"
    r += len(measured) + 2
    _put(ws, f"A{r}", "Ghi chú: kết quả CĐR chính thức của lớp HP (BM6b) cộng dồn các bài KT thuộc loại minh chứng trong kế hoạch BM6a – "
                      "xuất bộ BM6 của lớp ở trang Kết quả CĐR (BM6).", italic=True, merge=f"A{r}:I{r}")
    ws.row_dimensions[r].height = 36
    _widths(ws, [7, 44, 10, 11, 11, 14, 12, 14, 22])
    for c in measured:
        ev = c["evidence_type"] or "any"
        _render_evidence(wb.create_sheet(f"BM6c_{c['clo_code']}"[:31]), c, cs, unit, with_names=True, concl_suffix=" (theo bài KT này)",
                         extra=[f"Bài KT lấy minh chứng theo kế hoạch BM6a: {EV_NAME[ev]}" + ("" if c["used"] else " – bài này không tính vào BM6b")])
    _print_setup(wb)
    return _save(wb)


def build_bm6_exam(db: Session, exam_id: int) -> bytes:
    return render_bm6_exam(bm6_exam_data(db, exam_id))


# =====================================================================================  BM2 / BM3 – CĐR chương trình


def prev_next_year(year: str) -> tuple[str, str]:
    a, b = (int(x) for x in year.split("-"))
    return f"{a - 1}-{b - 1}", f"{a + 1}-{b + 1}"


def bm2_data(db: Session, program_id: int, year: str) -> dict:
    """Dữ liệu chung cho BM2 (Excel + Word) và BM3b/BM3c."""
    prog = dict(db.execute(text("SELECT * FROM programs WHERE id = :p"), {"p": program_id}).mappings().one())
    plos = [dict(p) for p in db.execute(text("SELECT * FROM plos WHERE program_id = :p ORDER BY group_no, id"), {"p": program_id}).mappings()]
    ids = [p["id"] for p in plos] or [0]
    rounds, results = {}, {}
    for r in db.execute(text(f"SELECT * FROM plo_measurement_plans WHERE plo_id IN ({','.join(map(str, ids))}) ORDER BY round_no")).mappings():
        rounds.setdefault(r["plo_id"], {})[int(r["round_no"])] = (r["academic_year"], float(r["target_pct"]))
    for r in db.execute(text(f"SELECT * FROM plo_results WHERE plo_id IN ({','.join(map(str, ids))})")).mappings():
        results.setdefault(r["plo_id"], {})[r["academic_year"]] = dict(r)
    for p in plos:
        p["rounds"] = rounds.get(p["id"], {})
        p["results"] = results.get(p["id"], {})
        p["pis"] = [dict(x) for x in db.execute(text("""
            SELECT pl.id, pi.pi_code, pi.description, co.course_name, s.name AS sem, pl.method, pl.cycle, pl.target_pct,
                   l.full_name AS lecturer,
                   (SELECT GROUP_CONCAT(c2.course_name ORDER BY c2.course_name SEPARATOR ', ') FROM pi_courses pc
                     JOIN courses c2 ON c2.id = pc.course_id WHERE pc.pi_id = pi.id) AS courses_text,
                   r.n_evaluated, r.n_achieved, r.is_achieved
            FROM performance_indicators pi
            LEFT JOIN pi_assessment_plans pl ON pl.pi_id = pi.id
                 AND pl.semester_id IN (SELECT id FROM semesters WHERE academic_year = :y)
            LEFT JOIN courses co ON co.id = pl.course_id
            LEFT JOIN semesters s ON s.id = pl.semester_id
            LEFT JOIN lecturers l ON l.id = pl.lecturer_id
            LEFT JOIN pi_results r ON r.plan_id = pl.id
            WHERE pi.plo_id = :plo ORDER BY pi.id, s.term"""), {"plo": p["id"], "y": year}).mappings()]
        p["planned"] = any(x["id"] for x in p["pis"])
    return {"prog": prog, "plos": plos, "year": year}


def bm2_summary(d: dict) -> dict:
    """Số liệu và nhận xét tự sinh của BM2b (dùng chung cho Excel và Word)."""
    year = d["year"]; prev, nxt = prev_next_year(year)
    target = float(d["prog"]["target_pct"])
    measured = [p for p in d["plos"] if year in p["results"]]
    ok = sum(p["results"][year]["n_achieved"] for p in measured); n = sum(p["results"][year]["n_evaluated"] for p in measured)
    # Bảng 1: kết quả lần đo gần nhất của mỗi CĐR tới năm báo cáo (giữ kết quả các năm trước – hướng dẫn của BM2b)
    latest = []
    for p in d["plos"]:
        ys = sorted(y for y in p["results"] if y <= year)
        if ys:
            latest.append(p["results"][ys[-1]])
    c_ok = sum(x["n_achieved"] for x in latest); c_n = sum(x["n_evaluated"] for x in latest)
    pct = lambda a, b: round(100 * a / b, 2) if b else None
    verdict = lambda a, b: ("Đạt" if 100 * a / b + 1e-9 >= target else "Không đạt") if b else ""
    weak = [p for p in measured if not p["results"][year]["is_achieved"]]
    weak_pi = [f"{p['plo_code']}-{x['pi_code']}" for p in measured for x in p["pis"] if x["n_evaluated"] and not x["is_achieved"]]
    lowest = sorted(measured, key=lambda p: float(p["results"][year]["achieved_pct"]))[:2]
    analyses = [f"CĐR {p['plo_code']}: {p['results'][year]['analysis']}" for p in measured if p["results"][year].get("analysis")]
    current = (f"Năm học {year} đo {len(measured)} CĐR: {ok}/{n} lượt SV đạt ({pct(ok, n) or 0:.2f}%), chỉ tiêu CTĐT {target:.0f}% – "
               f"{verdict(ok, n) or 'chưa có kết quả'}. CĐR không đạt: {', '.join(p['plo_code'] for p in weak) or 'không có'}"
               f"; PI không đạt: {', '.join(weak_pi) or 'không có'}." + ("\n" + "\n".join(analyses) if analyses else "")) if n else None
    prev_actions = [f"CĐR {p['plo_code']}: {p['results'][prev]['improvement_actions']}" for p in d["plos"]
                    if prev in p["results"] and p["results"][prev].get("improvement_actions")]
    prev_results = [f"CĐR {p['plo_code']}: {p['results'][year]['improvement_results']}" for p in measured
                    if p["results"][year].get("improvement_results")]
    proposals = [f"CĐR {p['plo_code']}: {p['results'][year]['improvement_actions']}" for p in measured
                 if p["results"][year].get("improvement_actions")]
    focus = ("Cần tập trung cải tiến cho – CĐR/PI không đạt: " + (", ".join([p["plo_code"] for p in weak] + weak_pi) or "không có")
             + "; hai CĐR có tỷ lệ đạt thấp nhất: "
             + (", ".join(f"{p['plo_code']} ({float(p['results'][year]['achieved_pct']):.1f}%)" for p in lowest) or "—") + ".")
    return {"prev": prev, "next": nxt, "target": target, "measured": measured, "ok": ok, "n": n, "pct": pct(ok, n),
            "verdict": verdict(ok, n), "cum_ok": c_ok, "cum_n": c_n, "cum_pct": pct(c_ok, c_n), "cum_verdict": verdict(c_ok, c_n),
            "n_pi_cols": max([3] + [len(p["pis"]) for p in measured]),
            "narratives": [
                (f"Nhận xét chung về kết quả đạt được trong năm học {year} (năm học hiện tại)", current),
                (f"Nhận xét chung về việc triển khai các hoạt động cải tiến đã đề ra trong năm học {prev} (năm học trước)",
                 "\n".join(prev_actions + prev_results) or None),
                (f"Đề xuất giải pháp cải tiến trong năm học {nxt} (năm học tiếp theo)", "\n".join([focus] + proposals) if n else None)]}


def _plo_head(ws, row: int, first_label: str) -> None:
    """Tiêu đề 2 tầng: CĐR (2 cột: nhóm, mã) | Nội dung | ..."""
    ws.merge_cells(f"A{row}:B{row + 1}"); ws.merge_cells(f"C{row}:C{row + 1}")
    _put(ws, f"A{row}", first_label, bold=True, align=CENTER, box=True)
    _box_range(ws, f"A{row}:C{row + 1}")


def render_bm3(d: dict) -> bytes:
    prog, plos, year = d["prog"], d["plos"], d["year"]
    s = bm2_summary(d)
    unit = prog["department"] or "Bộ môn"
    abbr = "".join(w[0] for w in unit.replace("Bộ môn", "").split() if w[:1].isalpha()).upper() or "BM"
    doc_no = f"Số:…../KH-{abbr}"
    vv = f"V/v triển khai đo lường mức độ đạt được chuẩn đầu ra CTĐT {prog['name'].upper()} năm học {year}"
    wb = Workbook()

    # ---------------- BM2a – kế hoạch đo lường CĐR CTĐT
    ws = wb.active; ws.title = "BM2a_KeHoach"
    _header(ws, unit, 7, left_cols=3, doc_no=doc_no)
    _put(ws, "G4", "BM2a", italic=True, align=Alignment(horizontal="right"))
    _put(ws, "A5", "KẾ HOẠCH", bold=True, size=TITLE, align=CENTER, merge="A5:G5")
    _put(ws, "A6", vv, bold=True, size=TITLE, align=CENTER, merge="A6:G6"); ws.row_dimensions[6].height = 36
    _put(ws, "A8", f"Kế hoạch đo lường mức độ đạt được chuẩn đầu ra CTĐT / Trình độ {prog['level'].lower()}", bold=True, merge="A8:G8")
    _plo_head(ws, 9, "CĐR")
    _put(ws, "C9", f"Nội dung của CĐR CTĐT\n{prog['code']}", bold=True, align=CENTER, box=True)
    for col, lab in (("D", "KH đo CĐR lần 1"), ("F", "KH đo CĐR lần 2")):
        nxt = chr(ord(col) + 1)
        _put(ws, f"{col}9", lab, bold=True, align=CENTER, box=True, merge=f"{col}9:{nxt}9"); ws[f"{nxt}9"].border = BOX
        _put(ws, f"{col}10", "Chỉ tiêu", bold=True, align=CENTER, box=True)
        _put(ws, f"{nxt}10", "Năm học đo lường", bold=True, align=CENTER, box=True)
    r = _plo_rows(ws, 11, plos, lambda p: [_pct(p["rounds"].get(1, (None, None))[1]), p["rounds"].get(1, ("", 0))[0],
                                          _pct(p["rounds"].get(2, (None, None))[1]), p["rounds"].get(2, ("", 0))[0]], pct_cols=(4, 6))
    planned = [p for p in plos if p["planned"]]
    r = _lines(ws, r + 1, 7, ["Kế hoạch đo lường chi tiết mức độ đạt được từng chuẩn đầu ra: "
                              + (", ".join(f"CĐR {p['plo_code']}" for p in planned) + " – xem các sheet BM3b tương ứng."
                                 if planned else f"chưa có kế hoạch đo PI trong năm học {year}.")], bold_prefix=("Kế hoạch",))
    _signer(ws, r + 1, 5, 7, "TRƯỞNG BỘ MÔN")
    _widths(ws, [6, 7, 62, 11, 16, 11, 16])

    # ---------------- BM2b – báo cáo tổng kết
    ws = wb.create_sheet("BM2b_TongKet")
    n_pi = s["n_pi_cols"]
    last_col = 7 + n_pi + 1
    _header(ws, unit, last_col, left_cols=3)
    _put(ws, f"{get_column_letter(last_col)}4", "BM2b", italic=True, align=Alignment(horizontal="right"))
    lc = get_column_letter(last_col)
    _put(ws, "A5", "BÁO CÁO TỔNG KẾT", bold=True, size=TITLE, align=CENTER, merge=f"A5:{lc}5")
    _put(ws, "A6", vv, bold=True, size=TITLE, align=CENTER, merge=f"A6:{lc}6"); ws.row_dimensions[6].height = 36
    _put(ws, "A8", "Bảng tổng hợp kết quả đo lường mức độ đạt được chuẩn đầu ra CTĐT", bold=True, merge=f"A8:{lc}8")
    _plo_head(ws, 9, "CĐR")
    _put(ws, "C9", f"Nội dung của CĐR CTĐT\n{prog['code']}", bold=True, align=CENTER, box=True)
    for col, lab in (("D", "KQ đo CĐR lần 1"), ("G", "KQ đo CĐR lần 2")):
        c2 = chr(ord(col) + 2)
        _put(ws, f"{col}9", lab, bold=True, align=CENTER, merge=f"{col}9:{c2}9"); _box_range(ws, f"{col}9:{c2}9")
        for j, h in enumerate(["Tỷ lệ % đạt", "Kết quả", "Năm học đo lường"]):
            _put(ws, f"{chr(ord(col) + j)}10", h, bold=True, align=CENTER, box=True)

    def round_cells(p):
        out = []
        for k in (1, 2):
            y = p["rounds"].get(k, ("", 0))[0]
            x = p["results"].get(y) if y and y <= year else None
            out += [_pct(x["achieved_pct"]) if x else "", ("Đạt" if x["is_achieved"] else "Không đạt") if x else "", y]
        return out
    r = _plo_rows(ws, 11, plos, round_cells, pct_cols=(4, 7))
    _row(ws, r, ["", "KẾT QUẢ ĐẠT ĐƯỢC CĐR CTĐT", "", _pct(s["cum_pct"]), s["cum_verdict"], "", "", "", ""], pct_cols=(4,), bold=True)
    ws.merge_cells(f"B{r}:C{r}")
    r += 2
    _put(ws, f"A{r}", f"Bảng tổng hợp chi tiết kết quả đạt được của từng CĐR đo lường trong năm học {year}", bold=True, merge=f"A{r}:{lc}{r}")
    r += 1
    _plo_head(ws, r, "CĐR #")
    _put(ws, f"C{r}", f"Nội dung của CĐR CTĐT\n{prog['code']}", bold=True, align=CENTER, box=True)
    _put(ws, f"D{r}", "Kết quả tổng hợp của từng CĐR", bold=True, align=CENTER, merge=f"D{r}:F{r}"); _box_range(ws, f"D{r}:F{r}")
    _put(ws, f"H{r}", "Kết quả thực hiện", bold=True, align=CENTER, merge=f"H{r}:{lc}{r}"); _box_range(ws, f"H{r}:{lc}{r}")
    for j, h in enumerate(["Tổng số SV đã đạt", "Tổng số SV đã khảo sát", "Tỷ lệ % đã đạt"]):
        _put(ws, f"{chr(ord('D') + j)}{r + 1}", h, bold=True, align=CENTER, box=True)
    for j in range(n_pi):
        _put(ws, f"{get_column_letter(8 + j)}{r + 1}", f"PI {j + 1}", bold=True, align=CENTER, box=True)
    _put(ws, f"{lc}{r + 1}", "CĐR", bold=True, align=CENTER, box=True)
    first = r + 2

    def detail_cells(p, rr):
        x = p["results"][year]
        res = [(("Đạt" if q["is_achieved"] else "Không đạt") if q["n_evaluated"] else "") for q in p["pis"][:n_pi]]
        res += [""] * (n_pi - len(res))                      # cột PI k = PI thứ k của CĐR (PI chưa đo để trống)
        return [x["n_achieved"], x["n_evaluated"], f'=IF(E{rr}>0,D{rr}/E{rr},"")', ""] + res + ["Đạt" if x["is_achieved"] else "Không đạt"]
    r = _plo_rows(ws, first, s["measured"], detail_cells, pct_cols=(6,), with_row=True)
    if s["measured"]:
        _row(ws, r, ["", "KẾT QUẢ ĐẠT ĐƯỢC CĐR CTĐT", "", f"=SUM(D{first}:D{r - 1})", f"=SUM(E{first}:E{r - 1})",
                     f'=IF(E{r}>0,D{r}/E{r},"")', "", s["verdict"]] + [""] * n_pi, pct_cols=(6,), bold=True)
        ws.merge_cells(f"B{r}:C{r}")
        r += 1
    else:
        r = _lines(ws, r, last_col, [f"(Chưa có CĐR nào được đo trong năm học {year}.)"])
    r = _lines(ws, r + 1, last_col, ["Kết quả đo lường mức độ đạt được của từng CĐR theo PI: xem các sheet BM3c tương ứng."],
               bold_prefix=())
    r = _narratives(ws, r + 1, last_col, s["narratives"])
    r = _signer(ws, r, last_col - 3, last_col, "TRƯỞNG BỘ MÔN")
    r = _narratives(ws, r + 1, last_col, [("Nhận xét và góp ý của Trưởng đơn vị", None)])
    _signer(ws, r, last_col - 3, last_col, "TRƯỞNG ĐƠN VỊ")
    _widths(ws, [6, 7, 52] + [11] * (last_col - 3))

    # ---------------- BM3b / BM3c cho từng CĐR có kế hoạch đo PI trong năm
    for p in plos:
        if not p["planned"]:
            continue
        x = p["results"].get(year)
        head = [f"Tên CTĐT: {prog['name']}", f"CĐR {p['plo_code']} - {p['description']}", f"Chỉ tiêu đạt CĐR: {float(p['target_pct']):.0f} %"]
        ws = wb.create_sheet(f"BM3b_CDR{p['plo_code']}"[:31])
        _header(ws, unit, 9, left_cols=4)
        _put(ws, "A5", "BẢNG KẾ HOẠCH KIỂM TRA, ĐÁNH GIÁ MỨC ĐỘ ĐẠT CHO TỪNG CHUẨN ĐẦU RA CTĐT", bold=True, size=TITLE, align=CENTER, merge="A5:I5")
        _lines(ws, 6, 9, head)
        _head_cells(ws, 10, ["TT", "Performance indicator (PI) cho CĐR này", "Các môn học có PI xuất hiện", "Môn học sẽ lấy minh chứng",
                             "Phương pháp kiểm tra, đánh giá", "Chu kỳ lấy minh chứng", "Thời gian lấy minh chứng", "Chỉ tiêu mong muốn",
                             "GV \nphụ trách"])
        r = 11
        for q in p["pis"]:
            _row(ws, r, [q["pi_code"], q["description"], q["courses_text"] or "", q["course_name"] or "", q["method"] or "",
                         q["cycle"] or "", q["sem"] or "", _pct(q["target_pct"]), q["lecturer"] or ""], wrap_cols=(2, 3), pct_cols=(8,))
            r += 1
        _row(ws, r, [f"CĐR {p['plo_code']}", "CHỈ TIÊU MONG MUỐN CỦA CHUẨN ĐẦU RA", "", "", "", "", "", _pct(p["target_pct"]), ""],
             pct_cols=(8,), bold=True)
        ws.merge_cells(f"B{r}:G{r}")
        r = _narratives(ws, r + 2, 9, [
            ("Tổng hợp dữ liệu đã đánh giá cho CĐR này (bao gồm tất cả các PIs)", x["data_summary"] if x else None),
            ("Đánh giá kết quả của số liệu tổng hợp", x["analysis"] if x else None),
            ("Những hành động cải tiến", x["improvement_actions"] if x else None),
            ("Kết quả của các cải tiến đã thực hiện", x["improvement_results"] if x else None),
            ("Công cụ đánh giá", x["evidence_tools"] if x else None)], unit_word="PI")
        _signer(ws, r + 1, 5, 8, "TRƯỞNG ĐƠN VỊ")
        _widths(ws, [9.9, 40, 22, 18, 19, 12.7, 12.4, 12, 16])

        ws = wb.create_sheet(f"BM3c_CDR{p['plo_code']}"[:31])
        _header(ws, unit, 8, left_cols=3)
        _put(ws, "A5", "KẾT QUẢ TỔNG HỢP CỦA TỪNG PI CHO CHUẨN ĐẦU RA", bold=True, size=TITLE, align=CENTER, merge="A5:H5")
        _lines(ws, 6, 8, head)
        _head_cells(ws, 10, ["TT", "Performance indicator (PI)\ncho CĐR này ", "Phương pháp/\nCông cụ đánh giá", "SL SV đạt yêu cầu",
                             "Tổng số SV đánh giá", "Tỷ lệ % đạt yêu cầu", "Chỉ tiêu mong muốn ", "Kết quả\nđạt được"])
        first = r = 11
        for q in p["pis"]:                                     # đủ mọi PI của CĐR, PI chưa đo để trống số liệu
            _row(ws, r, [q["pi_code"], q["description"], q["method"] or "", q["n_achieved"] if q["n_evaluated"] else "",
                         q["n_evaluated"] or "", f'=IF(N(E{r})>0,D{r}/E{r},"")', _pct(q["target_pct"]),
                         f'=IF(N(E{r})=0,"",IF(F{r}+1E-9>=G{r},"Đạt","Không đạt"))'], wrap_cols=(2, 3), pct_cols=(6, 7))
            r += 1
        _row(ws, r, [f"CĐR {p['plo_code']}", "KẾT QUẢ ĐẠT ĐƯỢC CỦA CĐR", "", f"=SUM(D{first}:D{r - 1})", f"=SUM(E{first}:E{r - 1})",
                     f'=IF(E{r}>0,D{r}/E{r},"")', _pct(x["target_pct"] if x else p["target_pct"]),
                     f'=IF(E{r}=0,"Chưa đo",IF(F{r}+1E-9>=G{r},"Đạt","Không đạt"))'], pct_cols=(6, 7), bold=True)
        ws.merge_cells(f"B{r}:C{r}")
        _signer(ws, r + 4, 4, 8, "TRƯỞNG ĐƠN VỊ")
        _widths(ws, [13.7, 40, 24, 13.6, 11.9, 12, 12.7, 12.9])
    _print_setup(wb)
    return _save(wb)


def _plo_rows(ws, row: int, plos: list[dict], cells, *, pct_cols=(), with_row=False) -> int:
    """Các dòng CĐR: cột A = nhóm CĐR (gộp ô theo nhóm), cột B = mã, cột C = nội dung, sau đó là cells(p)."""
    start, prev = row, None
    for p in plos:
        vals = cells(p, row) if with_row else cells(p)
        _row(ws, row, [p["group_no"], p["plo_code"], p["description"]] + vals, wrap_cols=(3,), pct_cols=pct_cols)
        if p["group_no"] != prev:
            if prev is not None and row - start > 1:
                ws.merge_cells(f"A{start}:A{row - 1}")
            start, prev = row, p["group_no"]
        row += 1
    if prev is not None and row - start > 1:
        ws.merge_cells(f"A{start}:A{row - 1}")
    return row


def build_bm3(db: Session, program_id: int, academic_year: str) -> bytes:
    return render_bm3(bm2_data(db, program_id, academic_year))


# =====================================================================================  Phân công đánh giá PIs


def assignments_data(db: Session, semester_id: int) -> dict:
    sem = dict(db.execute(text("SELECT * FROM semesters WHERE id = :s"), {"s": semester_id}).mappings().one())
    prog = db.execute(text("SELECT name FROM programs ORDER BY id LIMIT 1")).scalar()
    rows = [dict(r) for r in db.execute(text("""
        SELECT c.course_code, c.course_name, IF(a.all_supervisors = 1, :all, l.full_name) AS lecturer, a.all_supervisors, a.note
        FROM assessment_assignments a JOIN courses c ON c.id = a.course_id LEFT JOIN lecturers l ON l.id = a.lecturer_id
        WHERE a.semester_id = :s ORDER BY c.course_code, a.all_supervisors DESC, l.full_name"""),
        {"s": semester_id, "all": ALL_SUPERVISORS}).mappings()]
    return {"sem": sem, "program": prog, "rows": rows}


def render_assignments(d: dict) -> bytes:
    sem = d["sem"]
    wb = Workbook(); ws = wb.active; ws.title = "PhanCong"
    _put(ws, "A1", f"PHÂN CÔNG ĐÁNH GIÁ PIs HỌC KỲ {sem['term']} NĂM HỌC {sem['academic_year']}", size=TITLE,
         align=CENTER, merge="A1:E1")
    _put(ws, "A2", f"Ngành {d['program'] or ''}", size=TITLE, align=CENTER, merge="A2:E2")
    _put(ws, "A3", "Chú ý:", bold=True, align=LEFT_TOP)
    _put(ws, "B3", ASSIGN_NOTE, align=LEFT_TOP, merge="B3:E3"); ws.row_dimensions[3].height = 36
    _head_cells(ws, 5, ["STT", "MÃ MH", "TÊN MH", "GV ĐÁNH GIÁ", "GHI CHÚ"])
    groups: list[list[dict]] = []
    for x in d["rows"]:                                      # một môn có nhiều GV: các dòng liền nhau, gộp ô STT/mã/tên
        if groups and groups[-1][0]["course_code"] == x["course_code"]:
            groups[-1].append(x)
        else:
            groups.append([x])
    r = 6
    for i, g in enumerate(groups, 1):
        for k, x in enumerate(g):
            # môn giao cho "Tất cả thầy/cô có hướng dẫn" (TLCN, KLTN): ghi rõ ở cột GV đánh giá như biểu mẫu
            gv, note = x["lecturer"], x["note"]
            if x["all_supervisors"] and (note or "").lower().startswith("tất cả thầy/cô"):
                note = None                                   # ghi chú chỉ nhắc lại "Tất cả thầy/cô có HD"
            _row(ws, r + k, [i if k == 0 else "", x["course_code"] if k == 0 else "", x["course_name"] if k == 0 else "",
                             gv or "", note or ""], wrap_cols=(3, 4, 5))
        if len(g) > 1:
            for col in "ABC":
                ws.merge_cells(f"{col}{r}:{col}{r + len(g) - 1}")
        r += len(g)
    _widths(ws, [8, 15.1, 32, 26, 24])
    _print_setup(wb, landscape=False)
    return _save(wb)


def build_assignments(db: Session, semester_id: int) -> bytes:
    return render_assignments(assignments_data(db, semester_id))
