"""Xuất đề cương chi tiết học phần ra tệp Word theo bố cục mẫu đề cương của trường (A4 đứng, Times New Roman 13):
quốc hiệu – tiêu đề – 11 mục như trang Đề cương chi tiết – chữ ký Trưởng khoa / Trưởng bộ môn / Người biên soạn.

Dữ liệu lấy đúng từ GET /api/courses/{cid}/syllabus (routers.programs.syllabus) nên Word và trang web luôn khớp nhau.
Góp ý 23/10: "In" đề cương phải ra văn bản theo mẫu, không in nguyên trang web.
"""
from __future__ import annotations

import io
from datetime import datetime

from docx import Document
from docx.shared import Cm, Pt

from .reports_excel import SCHOOL
from .reports_word import CENTER, FONT, JUSTIFY, LEFT, _cell, _p, _table, qn

BLOOM_VI = {1: "Nhớ", 2: "Hiểu", 3: "Vận dụng", 4: "Phân tích", 5: "Đánh giá", 6: "Sáng tạo"}
EVIDENCE_VI = {"process": "Quá trình", "final": "Cuối kỳ", "any": "Mọi bài KT"}


def _h(doc, text: str) -> None:
    _p(doc, text, bold=True, align=LEFT, space_after=3).paragraph_format.space_before = Pt(8)


def _rows(doc, head: list[str], widths: list[float], data: list[list], aligns: list | None = None, size=11):
    """Bảng có dòng tiêu đề đậm; aligns: căn lề từng cột (mặc định giữa, cột chữ dài căn trái)."""
    t = _table(doc, len(data) + 1, widths)
    for j, h in enumerate(head):
        _cell(t.cell(0, j), h, bold=True, size=size)
    for i, r in enumerate(data, 1):
        for j, v in enumerate(r):
            _cell(t.cell(i, j), "" if v is None else v, align=(aligns[j] if aligns else CENTER), size=size)
    return t


def _text_or_dash(doc, v) -> None:
    _p(doc, v.strip() if v and v.strip() else "—", align=JUSTIFY)


def render_syllabus_docx(s: dict, faculty: str | None = None) -> bytes:
    d = s.get("doc") or {}
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(2), Cm(2), Cm(2)
    st = doc.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(13)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    # ---------------- Quốc hiệu, tên trường / khoa
    t = doc.add_table(rows=1, cols=2); t.autofit = False
    t.rows[0].cells[0].width, t.rows[0].cells[1].width = Cm(7.9), Cm(8.5)
    _cell(t.cell(0, 0), SCHOOL + "\n" + (faculty or "KHOA CÔNG NGHỆ THÔNG TIN").upper(), bold=True, size=10.5)
    _cell(t.cell(0, 1), "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nĐộc lập – Tự do – Hạnh phúc", bold=True, size=10.5)
    for run in t.cell(0, 0).paragraphs[0].runs:              # tên trường chữ thường, tên khoa in đậm như mẫu văn bản
        if run.text in SCHOOL.split("\n"):
            run.bold = False
    _p(doc, "")
    _p(doc, "ĐỀ CƯƠNG CHI TIẾT HỌC PHẦN", bold=True, size=15, align=CENTER, space_after=2)
    progs = s.get("programs") or []
    if progs:
        _p(doc, "Chương trình đào tạo: " + "; ".join(f"{p['name']} ({p['code']})" for p in progs), italic=True, align=CENTER, space_after=2)
    if d:
        status = "" if d.get("status") == "approved" else f" – {d.get('status_label', '')}, chưa có hiệu lực"
        _p(doc, f"Phiên bản {d.get('version_no')}{status}", italic=True, size=12, align=CENTER)

    # ---------------- 1. Thông tin chung
    _h(doc, "1. Thông tin chung")
    cred = f"{s['credits']}"
    if s.get("credits_theory") is not None or s.get("credits_practice") is not None:
        cred += f" (Lý thuyết: {s.get('credits_theory') or 0}; Thực hành/Thí nghiệm: {s.get('credits_practice') or 0})"
    pre = ", ".join(f"{p['course_code']}" + (f" – {p['course_name']}" if p.get("course_name") else "") for p in s.get("prerequisite_list") or []) or "Không"
    info = [
        ("Tên học phần (tiếng Việt)", s["course_name"]),
        ("Tên học phần (tiếng Anh)", d.get("english_name") or "—"),
        ("Mã học phần", s["course_code"]),
        ("Số tín chỉ", cred),
        ("Cấu trúc tín chỉ", s.get("course_structure") or "—"),
        ("Số giờ tự học", f"{d['self_study_hours']} giờ" if d.get("self_study_hours") is not None else "—"),
        ("Học phần trước / tiên quyết", pre),
        ("Thuộc CTĐT", "\n".join(f"{p['code']}: học kỳ {p.get('semester_no') or '—'}"
                                 + (f", {p['knowledge_block']}" if p.get("knowledge_block") else "")
                                 + (", tự chọn" if p.get("is_elective") else ", bắt buộc") for p in progs) or "—"),
        ("Chỉ tiêu đạt CĐR môn học", f"{s.get('clo_target_pct', 75):g}% sinh viên đạt"),
        ("Áp dụng từ học kỳ", d.get("effective_semester") or "—"),
    ]
    t = _table(doc, len(info), [5.5, 10.5])
    for i, (k, v) in enumerate(info):
        _cell(t.cell(i, 0), k, bold=True, align=LEFT, size=12); _cell(t.cell(i, 1), v, align=LEFT, size=12)

    # ---------------- 2–3
    _h(doc, "2. Mô tả vắn tắt học phần"); _text_or_dash(doc, s.get("description"))
    _h(doc, "3. Mục tiêu học phần"); _text_or_dash(doc, d.get("objectives"))

    # ---------------- 4. CLO
    _h(doc, "4. Chuẩn đầu ra học phần (CLO)")
    clos = s.get("clos") or []
    if clos:
        _rows(doc, ["CLO", "Mô tả", "Mức Bloom", "Đáp ứng PI › PLO (mức)"], [1.6, 8.0, 2.2, 4.2], [[
            c["clo_code"], c["description"], c.get("bloom") or BLOOM_VI.get(c.get("bloom_level_id"), "—"),
            "\n".join(f"{p['plo_code']} › {p['pi_code']} ({p['level']})" for p in c.get("pis") or [])
            or "\n".join(f"{p['plo_code']} ({p['level']})" for p in c.get("plos") or []) or "—"] for c in clos],
            [CENTER, LEFT, CENTER, LEFT])
    else:
        _text_or_dash(doc, None)

    # ---------------- 5. Ma trận học phần – PI
    _h(doc, "5. Học phần đáp ứng chỉ số PI của chương trình đào tạo")
    code_of = {c["id"]: c["clo_code"] for c in clos}
    any_pi = False
    for p in progs:
        if not p.get("pis"):
            continue
        any_pi = True
        _p(doc, f"{p['code']} – {p['name']}", italic=True, align=LEFT, size=12, space_after=2)
        _rows(doc, ["PLO", "PI", "Nội dung PI", "Mức", "CLO đáp ứng"], [1.6, 1.6, 8.2, 1.4, 3.2], [[
            pi["plo_code"], pi["pi_code"], pi["description"], pi.get("level") or "",
            ", ".join(f"{code_of.get(m['clo_id'], '?')} ({m['level']})" for m in pi.get("clos") or []) or "Chưa có CLO"]
            for pi in p["pis"]], [CENTER, CENTER, LEFT, CENTER, CENTER])
    if not any_pi:
        _text_or_dash(doc, None)

    # ---------------- 6. Nội dung chi tiết
    _h(doc, "6. Nội dung chi tiết học phần")
    outs = s.get("outlines") or []
    if outs:
        data = [[f"Chương {o['chapter_number']}. {o['chapter_name']}",
                 (o.get("content") or "") + (f"\nHoạt động dạy và học: {o['teaching_activities']}" if o.get("teaching_activities") else ""),
                 o.get("theory_hours") if o.get("theory_hours") is not None else "",
                 o.get("practice_hours") if o.get("practice_hours") is not None else "",
                 ", ".join(x["clo_code"] for x in o.get("clos") or [])] for o in outs]
        data.append(["Tổng số tiết", "", sum(o.get("theory_hours") or 0 for o in outs), sum(o.get("practice_hours") or 0 for o in outs), ""])
        t = _rows(doc, ["Chương", "Nội dung", "LT", "TH", "CLO"], [3.8, 8.0, 1.2, 1.2, 1.8], data, [LEFT, LEFT, CENTER, CENTER, CENTER])
        last = len(data)
        _cell(t.cell(last, 0).merge(t.cell(last, 1)), "Tổng số tiết", bold=True, size=11)
    else:
        _text_or_dash(doc, None)

    # ---------------- 7
    _h(doc, "7. Phương pháp giảng dạy"); _text_or_dash(doc, d.get("teaching_methods"))

    # ---------------- 8. Đánh giá học phần
    _h(doc, "8. Đánh giá học phần")
    ass = d.get("assessments") or []
    if ass:
        data = [[a["component_code"], a["name"], a.get("method") or "", a.get("timing") or "", f"{a['weight_pct']:g}%",
                 ", ".join(x["clo_code"] for x in a.get("clos") or [])] for a in ass]
        data.append(["", "Tổng", "", "", f"{sum(a['weight_pct'] for a in ass):g}%", ""])
        _rows(doc, ["Mã", "Thành phần đánh giá", "Hình thức", "Thời điểm", "Tỷ trọng", "CLO"], [1.3, 4.6, 3.6, 2.2, 1.9, 2.4], data,
              [CENTER, LEFT, LEFT, CENTER, CENTER, CENTER])
    else:
        _text_or_dash(doc, None)

    # ---------------- 9. Kế hoạch BM6a
    _h(doc, "9. Kế hoạch kiểm tra – đánh giá chuẩn đầu ra học phần")
    plans = [c for c in clos if c.get("plan")]
    if plans:
        _rows(doc, ["CLO", "Học kỳ", "Bài KT có CLO", "Minh chứng", "Phương pháp", "Ngưỡng đạt", "Chỉ tiêu"],
              [1.4, 2.0, 3.6, 2.0, 3.2, 1.9, 1.9], [[
                  c["clo_code"], c["plan"].get("semester"), c["plan"].get("assessments_text") or "",
                  EVIDENCE_VI.get(c["plan"].get("evidence_type"), ""), c["plan"].get("method") or "",
                  f"≥ {float(c['plan']['pass_threshold_pct']):g}%", f"{float(c['plan']['target_pct']):g}% SV"] for c in plans],
              [CENTER, CENTER, LEFT, CENTER, LEFT, CENTER, CENTER], size=10.5)
    else:
        _p(doc, "Giảng viên phụ trách lập kế hoạch đánh giá CLO (BM6a) cho từng lớp học phần.", italic=True)

    # ---------------- 10. Tài liệu
    _h(doc, "10. Tài liệu học tập")
    refs = d.get("references") or []
    for typ, label in (("main", "Giáo trình chính"), ("extra", "Tài liệu tham khảo")):
        items = [r["citation"] for r in refs if r["ref_type"] == typ]
        _p(doc, label + ":", bold=True, italic=True, align=LEFT, size=12, space_after=1)
        for i, c in enumerate(items, 1):
            _p(doc, f"[{i}] {c}", align=JUSTIFY, space_after=1)
        if not items:
            _p(doc, "—")

    # ---------------- 11
    _h(doc, "11. Quy định của học phần"); _text_or_dash(doc, d.get("course_policy"))

    # ---------------- Ngày ban hành, chữ ký
    when = datetime.fromisoformat(d["approved_at"]) if d.get("approved_at") else None
    date_line = (f"TP. Hồ Chí Minh, ngày {when.day:02d} tháng {when.month:02d} năm {when.year}" if when
                 else "TP. Hồ Chí Minh, ngày …… tháng …… năm ……")
    _p(doc, "")
    _p(doc, date_line, italic=True, align=CENTER).paragraph_format.left_indent = Cm(7)
    t = doc.add_table(rows=1, cols=3); t.autofit = False
    for j, (w, title) in enumerate(zip((5.3, 5.3, 5.4), ("TRƯỞNG KHOA", "TRƯỞNG BỘ MÔN", "NGƯỜI BIÊN SOẠN"))):
        t.rows[0].cells[j].width = Cm(w)
        _cell(t.cell(0, j), title + "\n(Ký và ghi rõ họ tên)", bold=True, size=12)
        t.cell(0, j).paragraphs[0].runs[-1].bold = False
        t.cell(0, j).paragraphs[0].runs[-1].italic = True

    buf = io.BytesIO(); doc.save(buf)
    return buf.getvalue()
