"""Đề thi giấy do HỆ THỐNG sinh và chấm (không cần Moodle) – góp ý 15/10, tham khảo cách làm của TNMaker:

* Sinh nhiều mã đề: xáo trộn thứ tự câu và thứ tự phương án (mỗi câu 2–6 phương án tùy câu) → exam_versions,
  exam_version_questions, exam_version_options (Reverse Shuffle khi chấm).
* Xuất đề in (PDF, Word) một cột hoặc chia 2 cột trên một mặt giấy để tiết kiệm giấy; xuất đáp án từng mã đề.
* Phiếu trả lời theo mẫu phiếu dùng chung nhiều mã đề của trường, thêm 4 ô định vị ở 4 góc, khung tô MSSV và ô tô mã đề
  để chấm bằng ảnh; cùng một bộ tọa độ (sheet_layout) dùng để VẼ phiếu và để ĐỌC phiếu (services/omr.py).
"""
from __future__ import annotations

import html
import io
import math
import random
import re
from dataclasses import dataclass, field
from pathlib import Path

from sqlalchemy import text
from sqlalchemy.orm import Session

from ..models import Exam, ExamVersion, ExamVersionOption, ExamVersionQuestion

LABELS = "ABCDEFGHIJ"
FONT_DIR = Path(__file__).resolve().parents[1] / "assets" / "fonts"
MM = 72 / 25.4            # 1 mm = 2.8346 pt
PAGE_W, PAGE_H = 210.0, 297.0


class PaperError(ValueError):
    pass


# ------------------------------------------------------------------ phông chữ (Times New Roman nếu máy có, nếu không DejaVu Serif)
_FONTS: dict | None = None


def fonts() -> dict:
    global _FONTS
    if _FONTS:
        return _FONTS
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib.fonts import addMapping
    win = Path("C:/Windows/Fonts")
    cands = [("TNR", {"": win / "times.ttf", "B": win / "timesbd.ttf", "I": win / "timesi.ttf", "BI": win / "timesbi.ttf"}),
             ("DJS", {"": FONT_DIR / "DejaVuSerif.ttf", "B": FONT_DIR / "DejaVuSerif-Bold.ttf",
                      "I": FONT_DIR / "DejaVuSerif-Italic.ttf", "BI": FONT_DIR / "DejaVuSerif-BoldItalic.ttf"})]
    for name, files in cands:
        if all(p.exists() for p in files.values()):
            for k, p in files.items():
                pdfmetrics.registerFont(TTFont(name + k, str(p)))
            addMapping(name, 0, 0, name); addMapping(name, 1, 0, name + "B"); addMapping(name, 0, 1, name + "I"); addMapping(name, 1, 1, name + "BI")
            # chữ không chân: Arial như mẫu phiếu của trường nếu máy có, nếu không DejaVu Sans
            sans, sfiles = ("ARIAL", (win / "arial.ttf", win / "arialbd.ttf")) if (win / "arial.ttf").exists() and (win / "arialbd.ttf").exists() \
                else ("DJSANS", (FONT_DIR / "DejaVuSans.ttf", FONT_DIR / "DejaVuSans-Bold.ttf"))
            pdfmetrics.registerFont(TTFont(sans, str(sfiles[0])))
            pdfmetrics.registerFont(TTFont(sans + "B", str(sfiles[1])))
            pdfmetrics.registerFont(TTFont("DJSYM", str(FONT_DIR / "DejaVuSans.ttf")))     # ký hiệu ✓ ✗ (Arial không có)
            _FONTS = {"serif": name, "serifB": name + "B", "serifI": name + "I", "sans": sans, "sansB": sans + "B", "sym": "DJSYM"}
            return _FONTS
    raise PaperError("Không tìm thấy phông chữ tiếng Việt để xuất PDF")


def plain(s: str) -> str:
    """Nội dung câu hỏi (có thể là HTML từ Moodle) -> văn bản thuần."""
    s = re.sub(r"<\s*br\s*/?>|</p>", "\n", s or "", flags=re.I)
    s = re.sub(r"<[^>]+>", "", s)
    return html.unescape(s).strip()


def para(s: str) -> str:
    """Văn bản thuần -> markup an toàn cho Paragraph của ReportLab."""
    return html.escape(plain(s)).replace("\n", "<br/>")


# ------------------------------------------------------------------ mã đề
def generate_versions(db: Session, exam: Exam, n_versions: int, shuffle_questions: bool = True, shuffle_options: bool = True,
                      seed: int | None = None, first_code: int = 1) -> list[dict]:
    if exam.exam_type != "paper":
        raise PaperError("Chỉ bài thi giấy mới sinh mã đề")
    if exam.moodle_offlinequiz_id:
        raise PaperError("Bài đã tạo đề giấy trên Moodle (Offline Quiz) – mã đề do Moodle quản lý")
    if db.execute(text("SELECT COUNT(*) FROM exam_attempts WHERE exam_id=:e"), {"e": exam.id}).scalar():
        raise PaperError("Bài đã có kết quả chấm – không sinh lại mã đề được")
    qs = [eq.question for eq in exam.questions]
    if not qs:
        raise PaperError("Bài kiểm tra chưa có câu hỏi")
    if any(q.question_type != "multichoice" for q in qs):
        raise PaperError("Đề giấy chấm phiếu chỉ gồm câu trắc nghiệm")
    if any(not (2 <= len(q.options) <= 6) for q in qs):
        raise PaperError("Mỗi câu trên phiếu trả lời có từ 2 đến 6 phương án")
    rng = random.Random(seed if seed is not None else exam.id * 7919)
    exam.versions.clear(); db.flush()
    out = []
    for k in range(n_versions):
        code = f"{first_code + k:02d}"
        order = list(qs)
        if shuffle_questions and k > 0:
            rng.shuffle(order)
        v = ExamVersion(exam_id=exam.id, version_code=code)
        db.add(v); db.flush()
        for pos, q in enumerate(order, 1):
            vq = ExamVersionQuestion(version_id=v.id, question_id=q.id, position=pos)
            opts = sorted(q.options, key=lambda o: o.position)
            if shuffle_options and k > 0:
                opts = list(opts); rng.shuffle(opts)
            db.add(vq); db.flush()
            for i, o in enumerate(opts):
                db.add(ExamVersionOption(version_id=v.id, question_id=q.id, displayed_label=LABELS[i], option_id=o.id))
        out.append({"id": v.id, "version_code": code})
    db.flush()
    return out


def version_view(db: Session, version_id: int) -> list[dict]:
    """Câu hỏi của một mã đề theo thứ tự in, phương án theo nhãn in (A, B, …) và đáp án đúng."""
    rows = db.execute(text("""SELECT vq.position, q.id AS question_id, q.content, eq.points, vo.displayed_label, o.content AS opt, o.is_correct, o.id AS option_id
                              FROM exam_version_questions vq JOIN question_bank q ON q.id=vq.question_id
                              JOIN exam_versions v ON v.id=vq.version_id
                              JOIN exam_questions eq ON eq.exam_id=v.exam_id AND eq.question_id=q.id
                              JOIN exam_version_options vo ON vo.version_id=vq.version_id AND vo.question_id=vq.question_id
                              JOIN question_options o ON o.id=vo.option_id
                              WHERE vq.version_id=:v ORDER BY vq.position, vo.displayed_label"""), {"v": version_id}).mappings()
    out: dict[int, dict] = {}
    for r in rows:
        q = out.setdefault(r["position"], {"position": r["position"], "question_id": r["question_id"], "content": plain(r["content"]),
                                           "points": float(r["points"]), "options": [], "answer": None})
        q["options"].append({"label": r["displayed_label"], "content": plain(r["opt"]), "option_id": r["option_id"]})
        if r["is_correct"]:
            q["answer"] = r["displayed_label"]
    return [out[k] for k in sorted(out)]


def versions_data(db: Session, exam: Exam) -> list[dict]:
    return [{"id": v.id, "version_code": v.version_code, "questions": version_view(db, v.id)} for v in exam.versions]


# ------------------------------------------------------------------ bố cục phiếu trả lời (đơn vị mm, gốc tọa độ góc trên-trái)
# Theo mẫu phiếu trả lời dùng chung nhiều mã đề của trường (góp ý 24/10): tiêu đề – môn – họ tên, lớp, ngày thi – mã đề –
# I. TRẮC NGHIỆM LỰA CHỌN (bảng "Câu | PHƯƠNG ÁN TRẢ LỜI", hàng sọc, mọi hàng cùng số phương án) – II. CÂU ĐIỀN / TỰ LUẬN –
# điểm, chữ ký. Phần thêm để máy đọc: 4 ô định vị ở góc, khung tô MSSV, mã đề dạng ô tô, vạch nhận chiều ở mép trên.
@dataclass
class SheetOptions:
    title: str = "PHIẾU TRẢ LỜI"
    school: str = "TRƯỜNG ĐẠI HỌC CÔNG NGHỆ KỸ THUẬT TP. HỒ CHÍ MINH"
    subtitle: str = ""                      # VD: LẬP TRÌNH PYTHON (IPPA233277)
    exam_title: str = ""
    footer: str = ""                        # chân trang, VD: IPPA233277 – Lớp IPPA233277_01
    n_questions: int = 40
    options_per_question: list[int] = field(default_factory=list)   # số phương án của từng câu (2–7); rỗng = 4
    student_digits: int = 8
    version_digits: int = 2
    version_codes: list[str] = field(default_factory=list)
    essay_lines: int = 0                    # số hàng ở phần II (câu điền / tự luận)
    note: str = ("Tô kín ô chữ cái của phương án chọn (bút chì 2B hoặc bút mực đen); chỉ chọn trong các phương án có ở đề. "
                 "Muốn đổi đáp án: tẩy sạch hoặc gạch chéo ô cũ rồi tô ô mới.")


FID = 7.0                                  # cạnh ô định vị (mm)
FID_POS = [(10.0, 10.0), (PAGE_W - 10.0 - FID, 10.0), (PAGE_W - 10.0 - FID, PAGE_H - 10.0 - FID), (10.0, PAGE_H - 10.0 - FID)]   # TL, TR, BR, BL
ORIENT_BAR = (PAGE_W / 2 - 12, 10.5, 24.0, 3.0)   # vạch đen giữa mép trên (x, y, rộng, cao) – nhận ra phiếu chụp lộn ngược
R_BUBBLE = 2.15
X0, X1 = 14.0, PAGE_W - 14.0               # lề trái/phải của nội dung
ID_X0 = 146.0                              # khung MSSV ở góc phải phần đầu phiếu
ROW_H, II_ROW_H = 6.2, 8.0
CONTENT_END = PAGE_H - 26.0                # nội dung kết thúc trên dòng "Điểm" và 2 ô định vị dưới


def fid_centers() -> list[tuple[float, float]]:
    return [(x + FID / 2, y + FID / 2) for x, y in FID_POS]


def sheet_layout(o: SheetOptions) -> dict:
    n = o.n_questions
    if not 1 <= n <= 120:
        raise PaperError("Phiếu trả lời hỗ trợ 1–120 câu")
    opts = (o.options_per_question or [4] * n)[:n] + [4] * max(0, n - len(o.options_per_question or []))
    max_opt = max(opts)
    if max_opt > 7:
        raise PaperError("Phiếu trả lời hỗ trợ tối đa 7 phương án mỗi câu")
    # khung MSSV: hàng ô viết số + lưới tô 0–9 mỗi cột
    dx_digit, dy_digit = min(5.2, (X1 - ID_X0 - 4) / max(o.student_digits, 1)), 4.4
    sx = ID_X0 + 2 + dx_digit / 2
    top = 40.0
    student = [{"col": c, "value": v, "x": sx + c * dx_digit, "y": top + 12.5 + v * dy_digit}
               for c in range(o.student_digits) for v in range(10)]
    id_bottom = top + 12.5 + 9 * dy_digit + 3.5
    # mã đề: một hàng ô tô, mỗi ô là một mã đề (thay cho "Mã đề (khoanh)" của mẫu)
    codes = o.version_codes or []
    vstep = min(9.0, 40.0 / max(len(codes) - 1, 1))
    version = [{"col": 0, "value": k, "code": code, "x": X0 + 30.0 + k * vstep, "y": 58.6} for k, code in enumerate(codes)]
    # I. Trắc nghiệm lựa chọn: 1–3 bảng cột, mọi hàng cùng số ô (= số phương án nhiều nhất)
    ncols = 1 if n <= 10 else 2 if n <= 24 else 3
    gap = 5.0
    col_w = (X1 - X0 - gap * (ncols - 1)) / ncols
    rows_per_col = math.ceil(n / ncols)
    a_top = id_bottom + 9                       # đầu bảng (dưới tiêu đề "I. ...")
    ii_h = (6 + 9 + 7 + o.essay_lines * II_ROW_H) if o.essay_lines else 0      # khoảng cách + tiêu đề + đầu bảng + các hàng
    bottom_limit = CONTENT_END - ii_h
    dy = min(ROW_H, (bottom_limit - a_top - 5) / max(rows_per_col, 1))
    first_row = a_top + 5 + dy / 2              # tâm hàng đầu (ngay dưới dòng tiêu đề bảng)
    if dy < 4.6:
        raise PaperError("Quá nhiều câu cho một phiếu A4 – giảm số câu, số phương án hoặc số câu ở phần II")
    dx_opt = min(9.5, (col_w - 16) / max_opt)
    off = 13 + ((col_w - 13) - dx_opt * max_opt) / 2      # các ô nằm giữa phần "PHƯƠNG ÁN TRẢ LỜI" như mẫu
    answers = []
    for i in range(n):
        c, r = divmod(i, rows_per_col)
        x0 = X0 + c * (col_w + gap)
        y = first_row + r * dy
        for j in range(max_opt):
            answers.append({"q": i + 1, "label": LABELS[j], "x": x0 + off + dx_opt / 2 + j * dx_opt, "y": y})
    a_bottom = first_row + (rows_per_col - 1) * dy + dy / 2
    return {"options": opts, "max_opt": max_opt, "student": student, "version": version, "answers": answers, "ncols": ncols,
            "rows_per_col": rows_per_col, "col_w": col_w, "gap": gap, "dy": dy, "dx_opt": dx_opt, "top": top, "id_bottom": id_bottom,
            "a_top": a_top, "a_bottom": a_bottom, "ii_top": a_bottom + 6, "sx": sx, "dx_digit": dx_digit, "dy_digit": dy_digit,
            "r": R_BUBBLE, "fids": fid_centers()}


def sheet_options_for(db: Session, exam: Exam, **kw) -> SheetOptions:
    cs = exam.class_section
    vs = versions_data(db, exam) if exam.versions else []
    n = len(exam.questions)
    per = []
    for pos in range(1, n + 1):
        cnt = [len(v["questions"][pos - 1]["options"]) for v in vs if len(v["questions"]) >= pos]
        per.append(max(cnt) if cnt else len(sorted(exam.questions, key=lambda q: q.order_index)[pos - 1].question.options))
    o = SheetOptions(subtitle=f"{cs.course.course_name.upper()} ({cs.course.course_code})", exam_title=exam.exam_title,
                     footer=f"{cs.course.course_code} – Lớp {cs.section_code}",
                     n_questions=n, options_per_question=per, version_codes=[v["version_code"] for v in vs])
    for k, v in kw.items():
        if v is not None and hasattr(o, k):
            setattr(o, k, v)
    if o.version_codes and max(len(c) for c in o.version_codes) > o.version_digits:
        o.version_digits = max(len(c) for c in o.version_codes)
    return o


# ------------------------------------------------------------------ vẽ phiếu trả lời
def _y(v: float) -> float:            # mm từ trên xuống -> pt từ dưới lên
    return (PAGE_H - v) * MM


def _dots(c, font, size, x0: float, x1: float, y: float) -> None:
    """Dòng chấm để điền tay từ x0 tới x1 (mm)."""
    w = c.stringWidth(".", font, size) / MM
    c.setFont(font, size); c.drawString(x0 * MM, _y(y), "." * max(1, int((x1 - x0) / w)))


def draw_sheet(c, o: SheetOptions, L: dict, prefill: dict | None = None) -> None:
    """Vẽ một phiếu lên canvas ReportLab. prefill (chỉ dùng cho dữ liệu thử): {'student': '2213...', 'version': '01', 'answers': {1:'A'}}."""
    from reportlab.lib.colors import Color, black
    F = fonts()
    S, SB = F["sans"], F["sansB"]
    ink = Color(0.25, 0.25, 0.25)              # chữ trong ô tô (đậm như mẫu nhưng vẫn nhạt hơn vết tô)
    rule = Color(0.55, 0.55, 0.55)
    zebra = Color(0.945, 0.945, 0.945)
    r = L["r"]
    c.setFillColor(black)
    for x, y in FID_POS:
        c.rect(x * MM, _y(y + FID), FID * MM, FID * MM, stroke=0, fill=1)
    bx, by, bw_, bh = ORIENT_BAR
    c.rect(bx * MM, _y(by + bh), bw_ * MM, bh * MM, stroke=0, fill=1)

    def bubble(x, y, label, size=7.6):
        c.setStrokeColor(ink); c.setLineWidth(0.55)
        c.circle(x * MM, _y(y), (r - 0.2) * MM, stroke=1, fill=0)
        c.setFillColor(ink); c.setFont(S, size); c.drawCentredString(x * MM, _y(y + size * 0.13), label)
        c.setFillColor(black)

    # ---- đầu phiếu
    c.setFont(SB, 16); c.drawCentredString(PAGE_W / 2 * MM, _y(23), o.title)
    if o.subtitle:
        c.setFont(SB, 11); c.drawCentredString(PAGE_W / 2 * MM, _y(29.5), o.subtitle[:80])
    vtxt = ("Dùng chung cho " + ", ".join(f"Đề {v}" for v in o.version_codes[:-1]) + f" và Đề {o.version_codes[-1]}"
            if len(o.version_codes) > 1 else f"Đề {o.version_codes[0]}" if o.version_codes else "")
    line = " • ".join(x for x in (o.exam_title, vtxt, f"{o.n_questions} câu") if x)
    c.setFont(S, 9.5); c.drawCentredString(PAGE_W / 2 * MM, _y(35), line[:120])
    c.setFont(S, 10)
    c.drawString(X0 * MM, _y(44), "Họ và tên:"); _dots(c, S, 10, X0 + 19, ID_X0 - 6, 44)
    c.drawString(X0 * MM, _y(51.5), "Lớp:"); _dots(c, S, 10, X0 + 9, X0 + 52, 51.5)
    c.drawString((X0 + 56) * MM, _y(51.5), "Ngày thi: ......... / ......... / ...........")
    c.setFont(SB, 10); c.drawString(X0 * MM, _y(59.6), "Mã đề (tô):")
    for b in L["version"]:
        bubble(b["x"], b["y"], b["code"], 7.2)
    if not L["version"]:
        _dots(c, S, 10, X0 + 22, X0 + 60, 59.6)
    # khung MSSV (viết số vào hàng ô rồi tô từng cột)
    top, dxd, dyd = L["top"], L["dx_digit"], L["dy_digit"]
    c.setStrokeColor(black); c.setLineWidth(0.6)
    c.rect(ID_X0 * MM, _y(L["id_bottom"]), (X1 - ID_X0) * MM, (L["id_bottom"] - top) * MM, stroke=1, fill=0)
    c.setFont(SB, 9); c.drawString((ID_X0 + 2) * MM, _y(top + 4), "MSSV:")
    c.setFont(S, 7); c.drawRightString((X1 - 2) * MM, _y(top + 4), "viết số rồi tô từng cột")
    for k in range(o.student_digits):
        c.setLineWidth(0.5); c.rect((L["sx"] - dxd / 2 + k * dxd + 0.3) * MM, _y(top + 10.2), (dxd - 0.6) * MM, 4.6 * MM, stroke=1, fill=0)
    for b in L["student"]:
        bubble(b["x"], b["y"], str(b["value"]), 6.6)
    # hướng dẫn (như dòng chữ nhỏ của mẫu)
    from reportlab.platypus import Paragraph
    from reportlab.lib.styles import ParagraphStyle
    st = ParagraphStyle("n", fontName=S, fontSize=8.6, leading=10.4)
    items = [o.note, "Ghi MSSV vào khung bên phải rồi tô chữ số tương ứng ở mỗi cột; tô ô mã đề của đề đang làm.",
             "Không gấp phiếu, không viết vào 4 ô vuông đen ở góc và vạch đen ở mép trên."]
    if o.essay_lines:
        items.append("Các câu điền / tự luận: viết đáp án ở phần II.")
    yy = 65.0
    for t in items:
        p = Paragraph(html.escape(t), st)
        _w, h = p.wrap((ID_X0 - 6 - X0) * MM, 60 * MM)
        p.drawOn(c, X0 * MM, _y(yy) - h)
        yy += h / MM + 0.8
    c.setFont(S, 8.6); c.drawString(X0 * MM, _y(yy + 4), "Tô đúng:")
    c.circle((X0 + 17) * MM, _y(yy + 3), (r - 0.2) * MM, stroke=0, fill=1)
    c.drawString((X0 + 24) * MM, _y(yy + 4), "Tô sai:")
    for k, mark in enumerate(("✓", "✗", "•", "")):          # ô cuối: khoanh vòng ngoài chữ cái – máy không đọc được
        cx = X0 + 38 + k * 7
        c.setStrokeColor(ink); c.setLineWidth(0.55); c.circle(cx * MM, _y(yy + 3), (r - 0.2) * MM, stroke=1, fill=0)
        c.setFont(F["sym"], 7.5); c.drawCentredString(cx * MM, _y(yy + 4), mark)
    c.setLineWidth(0.9); c.circle((X0 + 59) * MM, _y(yy + 3), (r + 0.9) * MM, stroke=1, fill=0)
    # ---- I. Trắc nghiệm lựa chọn
    a_top, dy, col_w, gap = L["a_top"], L["dy"], L["col_w"], L["gap"]
    c.setFont(SB, 10.5); c.drawString(X0 * MM, _y(a_top - 2), "I. TRẮC NGHIỆM LỰA CHỌN")
    by_q = {}
    for b in L["answers"]:
        by_q.setdefault(b["q"], []).append(b)
    for col in range(L["ncols"]):
        x0 = X0 + col * (col_w + gap)
        qs = [q for q in by_q if (q - 1) // L["rows_per_col"] == col]
        if not qs:
            continue
        y_end = by_q[qs[-1]][0]["y"] + dy / 2
        c.setFillColor(zebra); c.rect(x0 * MM, _y(a_top + 5), col_w * MM, 5 * MM, stroke=0, fill=1)
        for k, q in enumerate(qs):                 # hàng sọc như mẫu
            if k % 2:
                yq = by_q[q][0]["y"]
                c.rect(x0 * MM, _y(yq + dy / 2), col_w * MM, dy * MM, stroke=0, fill=1)
        c.setFillColor(black); c.setStrokeColor(rule); c.setLineWidth(0.6)
        c.line(x0 * MM, _y(a_top), (x0 + col_w) * MM, _y(a_top))
        c.line(x0 * MM, _y(a_top + 5), (x0 + col_w) * MM, _y(a_top + 5))
        c.line(x0 * MM, _y(y_end), (x0 + col_w) * MM, _y(y_end))
        c.setFont(SB, 8); c.drawString((x0 + 1.5) * MM, _y(a_top + 3.6), "Câu")
        c.drawCentredString((x0 + 13 + (col_w - 13) / 2) * MM, _y(a_top + 3.6), "PHƯƠNG ÁN TRẢ LỜI")
        for q in qs:
            bs = by_q[q]
            c.setFillColor(black); c.setFont(SB, 9); c.drawString((x0 + 2.5) * MM, _y(bs[0]["y"] + 1.1), f"{q:02d}")
            for b in bs:
                bubble(b["x"], b["y"], b["label"])
    # ---- II. Câu điền / tự luận
    if o.essay_lines:
        y0 = L["ii_top"]
        c.setFillColor(black); c.setFont(SB, 10.5); c.drawString(X0 * MM, _y(y0 + 2), "II. CÂU ĐIỀN / TỰ LUẬN")
        c.setFont(S, 8.6); c.drawString(X0 * MM, _y(y0 + 6.5), "Ghi số câu ở cột đúng mã đề đã tô; ghi đầy đủ nội dung trả lời vào cùng hàng.")
        codes = o.version_codes or ["…"]
        wc = 20.0
        t0, th = y0 + 9, 7.0
        xs = [X0 + k * wc for k in range(len(codes) + 1)]
        c.setStrokeColor(rule); c.setLineWidth(0.6)
        c.setFillColor(zebra); c.rect(X0 * MM, _y(t0 + th), (X1 - X0) * MM, th * MM, stroke=0, fill=1); c.setFillColor(black)
        bottom = t0 + th + o.essay_lines * II_ROW_H
        c.rect(X0 * MM, _y(bottom), (X1 - X0) * MM, (bottom - t0) * MM, stroke=1, fill=0)
        for x in xs[1:]:
            c.line(x * MM, _y(t0), x * MM, _y(bottom))
        c.setFont(SB, 8.2)
        for k, code in enumerate(codes):
            c.drawCentredString((xs[k] + wc / 2) * MM, _y(t0 + 4.6), f"Câu - Đề {code}")
        c.drawCentredString(((xs[-1] + X1) / 2) * MM, _y(t0 + 4.6), "NỘI DUNG TRẢ LỜI")
        for k in range(o.essay_lines):
            yr = t0 + th + k * II_ROW_H
            c.line(X0 * MM, _y(yr), X1 * MM, _y(yr))
            c.setFillColor(rule)
            _dots(c, S, 8, xs[-1] + 2, X1 - 3, yr + 3.4); _dots(c, S, 8, xs[-1] + 2, X1 - 3, yr + 7.2)
            c.setFillColor(black)
    yb = CONTENT_END + 5.5                  # dòng điểm/chữ ký cố định, phía trên 2 ô định vị dưới
    c.setFont(S, 9.5)
    c.drawString(X0 * MM, _y(yb), "Điểm (giảng viên ghi): ....................")
    c.drawString((X0 + 66) * MM, _y(yb), "Sinh viên ký: " + "." * 50)
    c.setFillColor(Color(0.3, 0.3, 0.3)); c.setFont(S, 7.8)
    foot = " • ".join(x for x in (o.footer or o.subtitle, "Phiếu trả lời dùng chung " + "/".join(o.version_codes) if o.version_codes else "Phiếu trả lời") if x)
    c.drawRightString((PAGE_W - 21) * MM, _y(PAGE_H - 13), foot[:110])
    c.setFillColor(black)
    if prefill:
        c.setFillColor(Color(0.1, 0.1, 0.1))
        s = prefill.get("student", "")
        for b in L["student"]:
            if b["col"] < len(s) and s[b["col"]] == str(b["value"]):
                c.circle(b["x"] * MM, _y(b["y"]), (r - 0.2) * MM, stroke=0, fill=1)
        for b in L["version"]:
            if b["code"] == prefill.get("version"):
                c.circle(b["x"] * MM, _y(b["y"]), (r - 0.2) * MM, stroke=0, fill=1)
        for b in L["answers"]:
            if b["label"] in (prefill.get("answers", {}).get(b["q"]) or ""):
                c.circle(b["x"] * MM, _y(b["y"]), (r - 0.2) * MM, stroke=0, fill=1)


def sheet_pdf(o: SheetOptions, copies: int = 1, prefills: list[dict] | None = None) -> bytes:
    from reportlab.pdfgen.canvas import Canvas
    L = sheet_layout(o)
    buf = io.BytesIO()
    c = Canvas(buf, pagesize=(PAGE_W * MM, PAGE_H * MM))
    c.setTitle(o.title)
    for k in range(len(prefills) if prefills else copies):
        draw_sheet(c, o, L, prefills[k] if prefills else None)
        c.showPage()
    c.save()
    return buf.getvalue()


def sheet_png(o: SheetOptions, zoom: float = 1.4) -> bytes:
    import fitz
    doc = fitz.open(stream=sheet_pdf(o), filetype="pdf")
    return doc[0].get_pixmap(matrix=fitz.Matrix(zoom, zoom)).tobytes("png")


# ------------------------------------------------------------------ đề in (PDF/Word) 1 hoặc 2 cột
@dataclass
class PaperOptions:
    columns: int = 1                      # 1 = để nguyên; 2 = chia đôi mặt giấy
    font_size: float = 11.5
    school: str = "TRƯỜNG ĐẠI HỌC CÔNG NGHỆ KỸ THUẬT TP. HỒ CHÍ MINH"
    faculty: str = "KHOA CÔNG NGHỆ THÔNG TIN"
    exam_name: str = ""                   # VD ĐỀ THI CUỐI KỲ HỌC KỲ I NĂM HỌC 2025-2026
    duration: str = ""                    # VD 60 phút
    notes: str = "Sinh viên không được sử dụng tài liệu. Làm bài trên phiếu trả lời."
    show_points: bool = False
    option_layout: str = "auto"           # auto | 1 (mỗi dòng 1 phương án) | 2 | 4


def _option_cols(opts: list[dict], mode: str, width_pt: float, font: str, size: float) -> int:
    """Số cột phương án: đặt 4 (hoặc 2) phương án trên một dòng khi đủ chỗ – đo bằng độ rộng chữ thật của phông."""
    from reportlab.pdfbase.pdfmetrics import stringWidth
    if mode in ("1", "2", "4"):
        return min(int(mode), len(opts))
    longest = max(stringWidth(f"{o['label']}. {o['content']}", font, size) for o in opts) + 8
    n = len(opts)
    if n <= 4 and longest * n <= width_pt:
        return n
    if longest * 2 <= width_pt:
        return 2 if n > 2 else n
    return 1


def exam_pdf(exam: Exam, version: dict, p: PaperOptions, semester_name: str = "") -> bytes:
    from reportlab.lib.colors import black
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import BaseDocTemplate, Frame, KeepTogether, PageTemplate, Paragraph, Spacer, Table, TableStyle
    F = fonts()
    cs = exam.class_section
    W, H = PAGE_W * MM, PAGE_H * MM
    margin = 15 * MM
    gap = 7 * MM
    head_h = (40 if p.notes else 35) * MM
    ncol = 2 if p.columns == 2 else 1
    fw = (W - 2 * margin - (gap if ncol == 2 else 0)) / ncol
    fs = p.font_size if ncol == 1 else min(p.font_size, 10.5)
    st_q = ParagraphStyle("q", fontName=F["serif"], fontSize=fs, leading=fs * 1.28, spaceBefore=fs * 0.45)
    st_o = ParagraphStyle("o", fontName=F["serif"], fontSize=fs, leading=fs * 1.22)
    st_h = ParagraphStyle("h", fontName=F["serifB"], fontSize=10.5, leading=13, alignment=1)
    st_hl = ParagraphStyle("hl", fontName=F["serif"], fontSize=10, leading=13, alignment=1)

    def header(c, doc):
        c.saveState()
        c.setFont(F["serif"], 8.5)
        c.drawRightString(W - margin, 9 * MM, f"Mã đề {version['version_code']} – Trang {doc.page}")
        c.drawString(margin, 9 * MM, f"{cs.course.course_code} – {exam.exam_title}"[:90])
        if ncol == 2:
            top = H - margin - (head_h if doc.page == 1 else 0)
            c.setLineWidth(0.4); c.line(W / 2, top, W / 2, 15 * MM)
        if doc.page == 1:
            t = Table([[Paragraph(f"<b>{html.escape(p.school)}</b><br/>{html.escape(p.faculty)}", st_hl),
                        Paragraph(f"<b>{html.escape(p.exam_name or exam.exam_title.upper())}</b><br/>"
                                  f"Môn: {html.escape(cs.course.course_name)} – Mã môn: {cs.course.course_code}<br/>"
                                  + (f"Thời gian: {html.escape(p.duration)} – " if p.duration else "")
                                  + f"{len(version['questions'])} câu – {html.escape(semester_name)}", st_hl)]],
                      colWidths=[(W - 2 * margin) * 0.42, (W - 2 * margin) * 0.58])
            t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
            w, h = t.wrap(W - 2 * margin, 40 * MM)
            t.drawOn(c, margin, H - margin - h)
            y = H - margin - h - 6 * MM
            c.setFont(F["serifB"], 12); c.drawString(margin, y, f"MÃ ĐỀ: {version['version_code']}")
            c.setFont(F["serif"], 10.5)
            c.drawString(margin + 38 * MM, y, "Họ tên SV: " + "." * 48 + "  MSSV: " + "." * 22)
            if p.notes:
                c.setFont(F["serifI"], 9.5); c.drawString(margin, y - 6 * MM, p.notes[:150])
            c.setLineWidth(0.6); c.line(margin, H - margin - head_h + 2 * MM, W - margin, H - margin - head_h + 2 * MM)
        c.restoreState()

    first = [Frame(margin + i * (fw + gap), 15 * MM, fw, H - margin - head_h - 15 * MM, leftPadding=0, rightPadding=0, topPadding=2, id=f"f{i}")
             for i in range(ncol)]
    later = [Frame(margin + i * (fw + gap), 15 * MM, fw, H - 2 * margin, leftPadding=0, rightPadding=0, topPadding=0, id=f"g{i}")
             for i in range(ncol)]
    buf = io.BytesIO()
    doc = BaseDocTemplate(buf, pagesize=(W, H), title=f"{exam.exam_title} – mã đề {version['version_code']}")
    doc.addPageTemplates([PageTemplate("first", first, onPage=header, autoNextPageTemplate="later"),
                          PageTemplate("later", later, onPage=header)])
    story = []
    for q in version["questions"]:
        pts = f" <i>({q['points']:g} điểm)</i>" if p.show_points else ""
        block = [Paragraph(f"<b>Câu {q['position']}.</b>{pts} {para(q['content'])}", st_q)]
        k = _option_cols(q["options"], p.option_layout, fw - 4 * MM, F["serif"], fs)
        cells = [Paragraph(f"<b>{o['label']}.</b> {para(o['content'])}", st_o) for o in q["options"]]
        rows = [cells[i:i + k] + [""] * (k - len(cells[i:i + k])) for i in range(0, len(cells), k)]
        t = Table(rows, colWidths=[(fw - 4 * MM) / k] * k)
        t.setStyle(TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 2), ("RIGHTPADDING", (0, 0), (-1, -1), 2),
                               ("TOPPADDING", (0, 0), (-1, -1), 0.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 0.5), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        t.hAlign = "LEFT"
        block += [Spacer(1, 1), t]
        story.append(KeepTogether(block))
    story += [Spacer(1, 8), Paragraph("<b>--- HẾT ---</b>", ParagraphStyle("end", parent=st_q, alignment=1))]
    doc.build(story)
    _ = black
    return buf.getvalue()


def exam_docx(exam: Exam, versions: list[dict], p: PaperOptions, semester_name: str = "") -> bytes:
    """Đề thi dạng Word (sửa được trước khi in); chia 2 cột bằng section 2 cột của Word."""
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt
    cs = exam.class_section
    d = Document()
    stl = d.styles["Normal"]; stl.font.name = "Times New Roman"; stl.font.size = Pt(p.font_size if p.columns == 1 else min(p.font_size, 11))
    stl.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    sec = d.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Mm(15))

    def cols(section, n):
        el = section._sectPr.find(qn("w:cols"))
        if el is None:
            el = OxmlElement("w:cols"); section._sectPr.append(el)
        el.set(qn("w:num"), str(n)); el.set(qn("w:space"), "400")
        if n == 2:
            el.set(qn("w:sep"), "1")
    for vi, v in enumerate(versions):
        if vi:
            sec = d.add_section(WD_SECTION.NEW_PAGE); cols(sec, 1)
        t = d.add_table(rows=1, cols=2)
        a, b = t.rows[0].cells
        a.paragraphs[0].add_run(p.school).bold = True
        a.add_paragraph(p.faculty)
        b.paragraphs[0].add_run(p.exam_name or exam.exam_title.upper()).bold = True
        b.add_paragraph(f"Môn: {cs.course.course_name} – Mã môn: {cs.course.course_code}")
        b.add_paragraph((f"Thời gian: {p.duration} – " if p.duration else "") + f"{len(v['questions'])} câu – {semester_name}")
        for cell in (a, b):
            for par in cell.paragraphs:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = d.add_paragraph().add_run(f"MÃ ĐỀ: {v['version_code']}"); r.bold = True
        d.add_paragraph("Họ tên SV: ................................................  MSSV: ......................")
        if p.notes:
            d.add_paragraph().add_run(p.notes).italic = True
        sec = d.add_section(WD_SECTION.CONTINUOUS)
        cols(sec, 2 if p.columns == 2 else 1)
        for q in v["questions"]:
            par = d.add_paragraph()
            par.add_run(f"Câu {q['position']}. ").bold = True
            if p.show_points:
                par.add_run(f"({q['points']:g} điểm) ").italic = True
            par.add_run(q["content"])
            par.paragraph_format.space_before = Pt(4); par.paragraph_format.space_after = Pt(1)
            par.paragraph_format.keep_with_next = True
            for o in q["options"]:
                op = d.add_paragraph()
                op.add_run(f"{o['label']}. ").bold = True
                op.add_run(o["content"])
                op.paragraph_format.left_indent = Mm(5); op.paragraph_format.space_after = Pt(0)
        e = d.add_paragraph(); e.alignment = WD_ALIGN_PARAGRAPH.CENTER; e.add_run("--- HẾT ---").bold = True
    buf = io.BytesIO(); d.save(buf)
    return buf.getvalue()


def answer_key_pdf(exam: Exam, versions: list[dict]) -> bytes:
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from reportlab.lib.colors import Color, black
    F = fonts()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=(PAGE_W * MM, PAGE_H * MM), leftMargin=15 * MM, rightMargin=15 * MM, topMargin=15 * MM, bottomMargin=15 * MM,
                            title=f"Đáp án – {exam.exam_title}")
    h = ParagraphStyle("h", fontName=F["serifB"], fontSize=13, leading=16, alignment=1)
    s = ParagraphStyle("s", fontName=F["serif"], fontSize=10, leading=13, alignment=1)
    story = [Paragraph(f"ĐÁP ÁN – {html.escape(exam.exam_title)}", h),
             Paragraph(f"{html.escape(exam.class_section.course.course_name)} – lớp {exam.class_section.section_code}", s), Spacer(1, 8)]
    for v in versions:
        story.append(Paragraph(f"Mã đề {v['version_code']}", ParagraphStyle("v", fontName=F["serifB"], fontSize=11.5, leading=15, spaceBefore=6)))
        cells = [f"{q['position']}.{q['answer'] or '?'}" for q in v["questions"]]
        per = 10
        rows = [cells[i:i + per] + [""] * (per - len(cells[i:i + per])) for i in range(0, len(cells), per)]
        t = Table(rows, colWidths=[18 * MM] * per)
        t.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, -1), F["sans"]), ("FONTSIZE", (0, 0), (-1, -1), 9.5),
                               ("GRID", (0, 0), (-1, -1), 0.4, Color(0.6, 0.6, 0.6)), ("ALIGN", (0, 0), (-1, -1), "CENTER")]))
        story.append(t)
    doc.build(story)
    _ = black
    return buf.getvalue()
