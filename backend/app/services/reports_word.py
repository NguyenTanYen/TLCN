"""Xuất BM2 (kế hoạch BM2a + báo cáo tổng kết BM2b) ra tệp Word theo bố cục biểu mẫu HCMUTE (A4 ngang, Times New Roman 13).

Dùng chung dữ liệu với bản Excel: reports_excel.bm2_data / bm2_summary.
"""
from __future__ import annotations

import io

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from sqlalchemy.orm import Session

from .reports_excel import bm2_data, bm2_summary

FONT = "Times New Roman"
CENTER, LEFT, RIGHT, JUSTIFY = WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.JUSTIFY


def _fmt_run(run, bold=False, italic=False, size=13):
    run.bold, run.italic = bold, italic
    run.font.size = Pt(size); run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def _p(doc_or_cell, text: str = "", *, bold=False, italic=False, size=13, align=JUSTIFY, space_after=4):
    p = doc_or_cell.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after); p.paragraph_format.space_before = Pt(0)
    for k, line in enumerate(text.split("\n")):
        if k:
            p.add_run().add_break()
        _fmt_run(p.add_run(line), bold, italic, size)
    return p


def _cell(cell, text, *, bold=False, italic=False, align=CENTER, size=12):
    cell.text = ""
    p = cell.paragraphs[0]; p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    for k, line in enumerate(str("" if text is None else text).split("\n")):
        if k:
            p.add_run().add_break()
        _fmt_run(p.add_run(line), bold, italic, size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _table(doc, n_rows: int, widths: list[float]):
    t = doc.add_table(rows=n_rows, cols=len(widths))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for row in t.rows:
        for j, w in enumerate(widths):
            row.cells[j].width = Cm(w)
    return t


def _merge(t, r1, c1, r2, c2, text=None, **kw):
    cell = t.cell(r1, c1).merge(t.cell(r2, c2))
    if text is not None:
        _cell(cell, text, **kw)
    return cell


def _pct(v) -> str:
    return "" if v is None else f"{float(v):.0f}%"


def _ok(flag) -> str:
    return "Đạt" if flag else "Không đạt"


def _plo_block(t, first_row: int, plos: list[dict]) -> None:
    """Cột 0 = nhóm CĐR (gộp ô theo nhóm), cột 1 = mã, cột 2 = nội dung."""
    start, prev = first_row, None
    for i, p in enumerate(plos):
        r = first_row + i
        _cell(t.cell(r, 1), p["plo_code"]); _cell(t.cell(r, 2), p["description"], align=LEFT)
        if p["group_no"] != prev:
            if prev is not None:
                _merge(t, start, 0, r - 1, 0, prev)
            start, prev = r, p["group_no"]
    if prev is not None:
        _merge(t, start, 0, first_row + len(plos) - 1, 0, prev)


def _two_level_head(t, prog_code: str, first_label: str) -> None:
    _merge(t, 0, 0, 1, 1, first_label, bold=True)
    _merge(t, 0, 2, 1, 2, f"Nội dung của CĐR CTĐT\n{prog_code}", bold=True)


def _signature(doc, title: str, date_line: str | None = None) -> None:
    t = doc.add_table(rows=1, cols=2); t.autofit = False
    t.rows[0].cells[0].width, t.rows[0].cells[1].width = Cm(15), Cm(10.5)
    _cell(t.cell(0, 1), (date_line + "\n" if date_line else "") + title, bold=True)
    _p(doc, "\n\n")


def render_bm2_docx(d: dict) -> bytes:
    prog, plos, year = d["prog"], d["plos"], d["year"]
    s = bm2_summary(d)
    unit = prog["department"] or "Bộ môn"
    abbr = "".join(w[0] for w in unit.replace("Bộ môn", "").split() if w[:1].isalpha()).upper() or "BM"
    vv = f"V/v triển khai đo lường mức độ đạt được chuẩn đầu ra CTĐT {prog['name'].upper()} năm học {year}"

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, m, Cm(2))
    st = doc.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(13)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    # ---------------- Tiêu đề văn bản
    t = doc.add_table(rows=1, cols=2); t.autofit = False
    t.rows[0].cells[0].width, t.rows[0].cells[1].width = Cm(11), Cm(14.7)
    _cell(t.cell(0, 0), f"TRƯỜNG ĐẠI HỌC SƯ PHẠM KỸ THUẬT\nTHÀNH PHỐ HỒ CHÍ MINH\n{unit.upper()}\nSố:…../KH-{abbr}", size=13)
    for k, para in enumerate(t.cell(0, 0).paragraphs[0].runs):
        para.bold = k in (4,)        # dòng tên đơn vị in đậm
    _cell(t.cell(0, 1), "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nĐộc lập – Tự do – Hạnh phúc\nTP. Hồ Chí Minh, ngày … tháng … năm 20…", size=13)
    runs = t.cell(0, 1).paragraphs[0].runs
    runs[0].bold = runs[2].bold = True; runs[4].italic = True

    # ================= BM2a – KẾ HOẠCH
    _p(doc, "BM2a", italic=True, size=11, align=RIGHT, space_after=0)
    _p(doc, "KẾ HOẠCH", bold=True, size=14, align=CENTER, space_after=0)
    _p(doc, vv, bold=True, size=14, align=CENTER, space_after=8)
    _p(doc, f"Kế hoạch đo lường mức độ đạt được chuẩn đầu ra CTĐT / Trình độ {prog['level'].lower()}", bold=True)
    t = _table(doc, 2 + len(plos), [1.2, 1.4, 13.1, 2.2, 3.2, 2.2, 3.2])
    _two_level_head(t, prog["code"], "CĐR")
    for c, lab in ((3, "KH đo CĐR lần 1"), (5, "KH đo CĐR lần 2")):
        _merge(t, 0, c, 0, c + 1, lab, bold=True)
        _cell(t.cell(1, c), "Chỉ tiêu", bold=True); _cell(t.cell(1, c + 1), "Năm học đo lường", bold=True)
    _plo_block(t, 2, plos)
    for i, p in enumerate(plos):
        for k, c in ((1, 3), (2, 5)):
            y, tg = p["rounds"].get(k, ("", None))
            _cell(t.cell(2 + i, c), _pct(tg)); _cell(t.cell(2 + i, c + 1), y)

    _p(doc, "")
    _p(doc, "Kế hoạch đo lường chi tiết mức độ đạt được từng chuẩn đầu ra", bold=True)
    planned = [p for p in plos if p["planned"]]
    if not planned:
        _p(doc, f"(Chưa có kế hoạch đo PI trong năm học {year}.)", italic=True)
    for p in planned:
        pis = p["pis"]
        t = _table(doc, 4 + len(pis), [1.6, 6.4, 3.8, 3.2, 3.2, 1.9, 2.2, 1.6, 2.6])
        _merge(t, 0, 0, 0, 8, f"CĐR {p['plo_code']} - {p['description']}", bold=True, align=LEFT)
        _merge(t, 1, 0, 1, 8, f"Chỉ tiêu đạt CĐR: {float(p['target_pct']):.0f} %", align=LEFT)
        for j, h in enumerate(["TT", "Performance indicator (PI) cho CĐR này", "Các môn học có PI xuất hiện", "Môn học sẽ lấy minh chứng",
                               "Phương pháp kiểm tra, đánh giá", "Chu kỳ lấy minh chứng", "Thời gian lấy minh chứng", "Chỉ tiêu mong muốn",
                               "GV phụ trách"]):
            _cell(t.cell(2, j), h, bold=True)
        for i, q in enumerate(pis):
            vals = [q["pi_code"], q["description"], q["courses_text"] or "", q["course_name"] or "", q["method"] or "", q["cycle"] or "",
                    q["sem"] or "", _pct(q["target_pct"]), q["lecturer"] or ""]
            for j, v in enumerate(vals):
                _cell(t.cell(3 + i, j), v, align=LEFT if j in (1, 2) else CENTER)
        last = 3 + len(pis)
        _cell(t.cell(last, 0), f"CĐR {p['plo_code']}", bold=True)
        _merge(t, last, 1, last, 6, "CHỈ TIÊU MONG MUỐN CỦA CHUẨN ĐẦU RA", bold=True)
        _cell(t.cell(last, 7), _pct(p["target_pct"]), bold=True)
        _p(doc, "")
    _signature(doc, "TRƯỞNG BỘ MÔN")

    # ================= BM2b – BÁO CÁO TỔNG KẾT
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    _p(doc, "BM2b", italic=True, size=11, align=RIGHT, space_after=0)
    _p(doc, "BÁO CÁO TỔNG KẾT", bold=True, size=14, align=CENTER, space_after=0)
    _p(doc, vv, bold=True, size=14, align=CENTER, space_after=8)
    _p(doc, "Bảng tổng hợp kết quả đo lường mức độ đạt được chuẩn đầu ra CTĐT", bold=True)
    t = _table(doc, 3 + len(plos), [1.2, 1.4, 10.5, 1.9, 2.2, 2.6, 1.9, 2.2, 2.6])
    _two_level_head(t, prog["code"], "CĐR")
    for c, lab in ((3, "KQ đo CĐR lần 1"), (6, "KQ đo CĐR lần 2")):
        _merge(t, 0, c, 0, c + 2, lab, bold=True)
        for j, h in enumerate(["Tỷ lệ % đạt", "Kết quả", "Năm học đo lường"]):
            _cell(t.cell(1, c + j), h, bold=True)
    _plo_block(t, 2, plos)
    for i, p in enumerate(plos):
        for k, c in ((1, 3), (2, 6)):
            y = p["rounds"].get(k, ("", None))[0]
            x = p["results"].get(y) if y and y <= year else None
            _cell(t.cell(2 + i, c), _pct(x["achieved_pct"]) if x else "")
            _cell(t.cell(2 + i, c + 1), _ok(x["is_achieved"]) if x else "")
            _cell(t.cell(2 + i, c + 2), y)
    last = 2 + len(plos)
    _merge(t, last, 1, last, 2, "KẾT QUẢ ĐẠT ĐƯỢC CĐR CTĐT", bold=True)
    _cell(t.cell(last, 3), _pct(s["cum_pct"]), bold=True); _cell(t.cell(last, 4), s["cum_verdict"], bold=True)

    _p(doc, "")
    _p(doc, f"Bảng tổng hợp chi tiết kết quả đạt được của từng CĐR đo lường trong năm học {year}", bold=True)
    n_pi, measured = s["n_pi_cols"], s["measured"]
    if measured:
        widths = [1.2, 1.4, 9.0, 2.0, 2.2, 1.9, 0.5] + [1.8] * n_pi + [1.8]
        t = _table(doc, 3 + len(measured), widths)
        _two_level_head(t, prog["code"], "CĐR #")
        _merge(t, 0, 3, 0, 5, "Kết quả tổng hợp của từng CĐR", bold=True)
        _merge(t, 0, 6, 1, 6)
        _merge(t, 0, 7, 0, 7 + n_pi, "Kết quả thực hiện", bold=True)
        for j, h in enumerate(["Tổng số SV đã đạt", "Tổng số SV đã khảo sát", "Tỷ lệ % đã đạt"]):
            _cell(t.cell(1, 3 + j), h, bold=True)
        for j in range(n_pi):
            _cell(t.cell(1, 7 + j), f"PI {j + 1}", bold=True)
        _cell(t.cell(1, 7 + n_pi), "CĐR", bold=True)
        _plo_block(t, 2, measured)
        for i, p in enumerate(measured):
            x, r = p["results"][year], 2 + i
            _cell(t.cell(r, 3), x["n_achieved"]); _cell(t.cell(r, 4), x["n_evaluated"]); _cell(t.cell(r, 5), _pct(x["achieved_pct"]))
            for j, q in enumerate(p["pis"][:n_pi]):          # cột PI k = PI thứ k của CĐR (PI chưa đo để trống)
                _cell(t.cell(r, 7 + j), _ok(q["is_achieved"]) if q["n_evaluated"] else "")
            _cell(t.cell(r, 7 + n_pi), _ok(x["is_achieved"]))
        last = 2 + len(measured)
        _merge(t, last, 1, last, 2, "KẾT QUẢ ĐẠT ĐƯỢC CĐR CTĐT", bold=True)
        _cell(t.cell(last, 3), s["ok"], bold=True); _cell(t.cell(last, 4), s["n"], bold=True)
        _cell(t.cell(last, 5), _pct(s["pct"]), bold=True); _cell(t.cell(last, 7), s["verdict"], bold=True)
    else:
        _p(doc, f"(Chưa có CĐR nào được đo trong năm học {year}.)", italic=True)

    _p(doc, "")
    _p(doc, f"Kết quả đo lường mức độ đạt của các CĐR được đo lường trong năm học {year}", bold=True)
    for p in measured:
        x = p["results"][year]
        pis = p["pis"]
        _p(doc, f"Kết quả đo lường mức độ đạt được CĐR {p['plo_code']}", bold=True, space_after=2)
        t = _table(doc, 2 + len(pis), [1.8, 9.5, 4.4, 2.2, 2.2, 2.0, 2.0, 2.2])
        for j, h in enumerate(["TT", "Performance indicator (PI) cho CĐR này", "Phương pháp/Công cụ đánh giá", "SL SV đạt yêu cầu",
                               "Tổng số SV đánh giá", "Tỷ lệ % đạt yêu cầu", "Chỉ tiêu mong muốn", "Kết quả đạt được"]):
            _cell(t.cell(0, j), h, bold=True)
        for i, q in enumerate(pis):
            done = bool(q["n_evaluated"])
            vals = [q["pi_code"], q["description"], q["method"] or "", q["n_achieved"] if done else "", q["n_evaluated"] if done else "",
                    _pct(100 * q["n_achieved"] / q["n_evaluated"]) if done else "", _pct(q["target_pct"]), _ok(q["is_achieved"]) if done else ""]
            for j, v in enumerate(vals):
                _cell(t.cell(1 + i, j), v, align=LEFT if j in (1, 2) else CENTER)
        last = 1 + len(pis)
        _cell(t.cell(last, 0), f"CĐR {p['plo_code']}", bold=True)
        _merge(t, last, 1, last, 2, "KẾT QUẢ ĐẠT ĐƯỢC CỦA CĐR", bold=True)
        for j, v in enumerate([x["n_achieved"], x["n_evaluated"], _pct(x["achieved_pct"]), _pct(x["target_pct"]), _ok(x["is_achieved"])]):
            _cell(t.cell(last, 3 + j), v, bold=True)
        _p(doc, "")

    for title, body in s["narratives"]:
        _p(doc, title, bold=True, space_after=2)
        _p(doc, body or "…………………………………………………………………………………………………………", italic=not body)
    _p(doc, "Hướng dẫn thực hiện: CTĐT cần phải phân tích và lập kế hoạch cải tiến cho những chuẩn đầu ra và PI “không đạt” "
            "và hai chuẩn đầu ra có tỷ lệ đạt thấp nhất.", italic=True, size=12)
    _signature(doc, "TRƯỞNG BỘ MÔN")
    _p(doc, "Nhận xét và góp ý của Trưởng đơn vị", bold=True)
    _p(doc, "…………………………………………………………………………………………………………\n"
            "…………………………………………………………………………………………………………")
    _signature(doc, "TRƯỞNG ĐƠN VỊ", "Ngày ..... tháng ..... năm 20..")

    bio = io.BytesIO(); doc.save(bio)
    return bio.getvalue()


def build_bm2_docx(db: Session, program_id: int, academic_year: str) -> bytes:
    return render_bm2_docx(bm2_data(db, program_id, academic_year))
