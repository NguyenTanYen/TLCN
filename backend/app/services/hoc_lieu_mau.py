"""Học liệu minh họa cho học phần Hệ quản trị CSDL (DBMS330284): slide từng chương (PPTX) + giáo trình trích (PDF).

Dùng cho dữ liệu minh họa/kiểm thử: nội dung bám các câu hỏi trong ngân hàng câu hỏi mẫu để xem gợi ý đọc đúng trang/slide.
"""
from __future__ import annotations

import io

# chương → [(tiêu đề slide, nội dung)]
SLIDES: dict[int, list[tuple[str, str]]] = {
    1: [
        ("Kiến trúc hệ quản trị cơ sở dữ liệu", "Hệ quản trị CSDL gồm các thành phần chính: bộ xử lý truy vấn (Query processor) phân tích, tối ưu và sinh kế hoạch thực thi cho câu lệnh SQL; bộ quản lý lưu trữ (Storage manager); bộ quản lý giao tác và bộ quản lý bộ đệm."),
        ("Tính độc lập dữ liệu", "Độc lập dữ liệu logic (logical data independence): thay đổi lược đồ khái niệm không làm ảnh hưởng các chương trình ứng dụng và lược đồ ngoài. Độc lập dữ liệu vật lý: thay đổi cách tổ chức lưu trữ không làm đổi lược đồ khái niệm."),
        ("Cơ sở dữ liệu hệ thống của SQL Server", "master lưu thông tin cấp máy chủ: tài khoản đăng nhập (logins), cấu hình máy chủ, vị trí các CSDL khác. model là khuôn mẫu cho CSDL mới, msdb lưu lịch và công việc của SQL Server Agent, tempdb chứa dữ liệu tạm."),
        ("Tệp dữ liệu và tệp nhật ký giao tác", "Tệp .mdf là tệp dữ liệu chính, .ndf là tệp dữ liệu phụ. Tệp nhật ký giao tác (.ldf) ghi lại mọi thay đổi để phục hồi CSDL về trạng thái nhất quán sau sự cố và để hoàn tác giao tác."),
        ("Các nhóm ngôn ngữ SQL", "DDL (ngôn ngữ định nghĩa dữ liệu): CREATE, ALTER TABLE, DROP. DML (thao tác dữ liệu): SELECT, INSERT, UPDATE, DELETE. DCL (điều khiển truy cập): GRANT, REVOKE, DENY. TCL: BEGIN TRANSACTION, COMMIT, ROLLBACK."),
    ],
    2: [
        ("Biến trong T-SQL", "Khai báo biến bằng DECLARE, tên biến bắt đầu bằng @. Ví dụ: DECLARE @Tong INT = 0; gán giá trị bằng SET @Tong = 5 hoặc SELECT @Tong = COUNT(*) FROM SinhVien."),
        ("Cấu trúc điều khiển và vòng lặp WHILE", "IF...ELSE, CASE và WHILE điều khiển luồng lệnh. WHILE @i < 5 BEGIN ... END lặp tới khi điều kiện sai; BREAK thoát vòng lặp, CONTINUE bỏ qua lần lặp hiện tại. Ví dụ @i bắt đầu 0, mỗi lần tăng 2: các giá trị 0, 2, 4, 6 và dừng khi @i = 6."),
        ("Thủ tục lưu trữ và tham số OUTPUT", "CREATE PROCEDURE khai báo thủ tục lưu trữ. Tham số khai báo với từ khóa OUTPUT dùng để trả giá trị từ thủ tục về cho nơi gọi; khi gọi phải ghi OUTPUT cả ở lệnh EXEC."),
        ("Hàm do người dùng định nghĩa", "Scalar function trả về một giá trị đơn. Table-valued function (inline hoặc multi-statement) trả về một bảng và có thể dùng trong mệnh đề FROM như một bảng thông thường."),
        ("Xử lý lỗi với TRY...CATCH", "BEGIN TRY ... END TRY BEGIN CATCH ... END CATCH. Trong khối CATCH dùng ERROR_MESSAGE() để lấy thông điệp lỗi, ERROR_NUMBER(), ERROR_SEVERITY(), ERROR_LINE() để biết chi tiết; THROW ném lại lỗi."),
    ],
    3: [
        ("Trigger AFTER và hai bảng tạm inserted, deleted", "Trigger AFTER UPDATE chạy sau khi cập nhật. Bảng tạm deleted chứa giá trị CŨ của các dòng bị cập nhật hoặc xóa; bảng inserted chứa giá trị MỚI của các dòng được thêm hoặc cập nhật."),
        ("Khi nào dùng trigger để cài ràng buộc", "Ràng buộc liên quan nhiều bảng, ví dụ tổng số lượng tồn kho không âm sau mỗi lần xuất hàng, không khai báo được bằng CHECK nên đặt trong Trigger. CHECK và khóa ngoại chỉ xét một bảng hoặc một quan hệ khóa."),
        ("Hủy thao tác vi phạm trong trigger", "Trong trigger, khi phát hiện vi phạm ràng buộc dùng ROLLBACK TRANSACTION để hủy bỏ thao tác gây vi phạm, kèm RAISERROR hoặc THROW để báo lỗi cho người dùng."),
        ("Trigger INSTEAD OF", "Trigger INSTEAD OF thay thế thao tác INSERT/UPDATE/DELETE gốc. Thường dùng để cập nhật view nhiều bảng, hoặc thay thao tác xóa thật bằng xóa mềm."),
        ("Khóa ngoại và hành động tham chiếu", "Khóa ngoại khai báo ON DELETE CASCADE sẽ tự động xóa các dòng con tham chiếu khi dòng cha bị xóa; ON DELETE SET NULL đặt khóa ngoại về NULL; NO ACTION từ chối thao tác."),
    ],
    4: [
        ("Tính chất ACID của giao tác", "Atomicity (nguyên tử): hoặc tất cả thao tác của giao tác được thực hiện, hoặc không thao tác nào. Consistency (nhất quán), Isolation (cô lập), Durability (bền vững)."),
        ("Các hiện tượng khi truy cập đồng thời", "Dirty read: một giao tác đọc dữ liệu do giao tác khác ghi nhưng CHƯA COMMIT. Non-repeatable read: đọc lại cùng một dòng nhận giá trị khác. Phantom read: truy vấn lặp lại trả về thêm các dòng mới thỏa điều kiện."),
        ("Các mức cô lập", "READ UNCOMMITTED cho phép dirty read. READ COMMITTED là mức mặc định, ngăn được dirty read nhưng vẫn có thể xảy ra non-repeatable read. REPEATABLE READ ngăn non-repeatable read; SERIALIZABLE ngăn cả phantom read."),
        ("Khóa chết (deadlock)", "Hai giao tác T1 giữ khóa A chờ B, T2 giữ khóa B chờ A tạo thành khóa chết. Hệ quản trị phát hiện chu trình chờ và chọn một giao tác làm nạn nhân (victim) để hủy, giao tác kia tiếp tục."),
        ("Giao thức khóa hai pha (2PL)", "Giao thức khóa hai pha gồm pha tăng trưởng (xin khóa) và pha thu hẹp (nhả khóa), bảo đảm lịch thao tác khả tuần tự (serializable)."),
    ],
    5: [
        ("Chỉ mục clustered và non-clustered", "Mỗi bảng trong SQL Server có tối đa 1 chỉ mục clustered vì dữ liệu chỉ sắp xếp vật lý theo một thứ tự; có thể có nhiều chỉ mục non-clustered."),
        ("Chọn chỉ mục phù hợp", "Truy vấn thường xuyên lọc theo cột MaKH và sắp xếp theo NgayLap nên dùng chỉ mục phức hợp (MaKH, NgayLap): cột lọc bằng đặt trước, cột sắp xếp đặt sau."),
        ("SARGable: điều kiện tận dụng được chỉ mục", "Điều kiện WHERE YEAR(NgayLap) = 2024 thường không tận dụng được chỉ mục trên NgayLap vì áp dụng hàm lên cột làm điều kiện không còn dạng tìm kiếm trực tiếp (non-SARGable). Nên viết NgayLap >= '2024-01-01' AND NgayLap < '2025-01-01'."),
        ("Execution plan", "Execution plan (kế hoạch thực thi) cho biết truy vấn dùng Index Seek hay Index Scan, Table Scan, chi phí từng toán tử. Xem bằng Display Estimated/Actual Execution Plan trong SSMS."),
        ("Chi phí của chỉ mục", "Tạo quá nhiều chỉ mục trên một bảng làm thao tác INSERT/UPDATE/DELETE chậm hơn vì phải cập nhật mọi chỉ mục, đồng thời tốn dung lượng lưu trữ."),
    ],
    6: [
        ("Cấp quyền với GRANT", "Cú pháp: GRANT SELECT ON SinhVien TO u1 cấp quyền SELECT trên bảng SinhVien cho người dùng u1. REVOKE thu hồi quyền, DENY từ chối quyền."),
        ("Vai trò cơ sở dữ liệu", "Vai trò db_datareader cho phép người dùng đọc dữ liệu mọi bảng người dùng; db_datawriter cho phép ghi; db_owner có toàn quyền trên CSDL."),
        ("Các loại sao lưu", "Sao lưu FULL chép toàn bộ CSDL. Bản sao lưu differential chứa các thay đổi kể từ bản sao lưu FULL gần nhất. Sao lưu transaction log ghi lại nhật ký giao tác."),
        ("Mô hình phục hồi (recovery model)", "Để phục hồi CSDL về một thời điểm cụ thể (point-in-time) phải dùng recovery model FULL (hoặc BULK_LOGGED) và có chuỗi sao lưu log. SIMPLE không hỗ trợ phục hồi tới thời điểm."),
    ],
}

# Giáo trình trích (PDF): cùng chủ đề, diễn đạt theo kiểu sách – dùng thử truy xuất trên nhiều loại tài liệu
TEXTBOOK_PAGES = [
    (1, "Chương 1. Tổng quan", "Kiến trúc ba mức của hệ quản trị cơ sở dữ liệu gồm mức ngoài, mức khái niệm và mức trong. Hai ánh xạ giữa các mức tạo nên tính độc lập dữ liệu: độc lập logic và độc lập vật lý."),
    (2, "Chương 1. Tổng quan", "Máy chủ SQL Server có bốn cơ sở dữ liệu hệ thống là master, model, msdb và tempdb. CSDL master giữ thông tin đăng nhập và cấu hình của máy chủ."),
    (3, "Chương 4. Giao tác", "Giao tác là đơn vị công việc logic thỏa bốn tính chất ACID. Khi một giao tác thất bại, hệ thống hoàn tác (rollback) mọi thay đổi dựa vào nhật ký giao tác."),
    (4, "Chương 4. Giao tác", "Cơ chế khóa: khóa chia sẻ (shared) và khóa độc quyền (exclusive). Hiện tượng deadlock được phát hiện bằng đồ thị chờ; hệ quản trị hủy giao tác có chi phí thấp nhất (nạn nhân)."),
    (5, "Chương 5. Chỉ mục", "Cấu trúc chỉ mục B-tree: nút gốc, nút trung gian và nút lá. Chỉ mục clustered quyết định thứ tự vật lý của dòng dữ liệu nên mỗi bảng chỉ có một."),
    (6, "Chương 5. Chỉ mục", "Tối ưu hóa truy vấn: bộ tối ưu chọn kế hoạch có chi phí ước lượng thấp nhất; thống kê (statistics) giúp ước lượng số dòng. Tránh áp dụng hàm lên cột trong mệnh đề WHERE."),
    (7, "Chương 6. Bảo mật và sao lưu", "Mô hình bảo mật SQL Server: đăng nhập (login) ở mức máy chủ, người dùng (user) ở mức CSDL, vai trò (role) gom quyền; cấp quyền bằng GRANT."),
    (8, "Chương 6. Bảo mật và sao lưu", "Chiến lược sao lưu: FULL định kỳ, differential giữa các kỳ FULL và sao lưu log thường xuyên để có thể phục hồi đến thời điểm bất kỳ."),
]


def build_pptx(chapter: int) -> bytes:
    from pptx import Presentation
    prs = Presentation()
    for title, body in SLIDES[chapter]:
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = title
        s.placeholders[1].text = body
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def build_textbook_pdf() -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfgen import canvas
    from reportlab.platypus import Paragraph
    from . import paper
    F = paper.fonts()
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    body = ParagraphStyle("b", fontName=F["serif"], fontSize=12, leading=17)
    for n, head, text in TEXTBOOK_PAGES:
        c.setFont(F["serifB"], 15)
        c.drawString(50, 790, head)
        p = Paragraph(text, body)
        w, h = p.wrap(495, 600)
        p.drawOn(c, 50, 770 - h)
        c.setFont(F["serif"], 10)
        c.drawString(290, 30, str(n))
        c.showPage()
    c.save()
    return buf.getvalue()


def samples(chapter_ids: dict[int, int]) -> list[dict]:
    """chapter_ids: số chương → id course_outlines. Trả danh sách tệp mẫu để thêm vào học phần."""
    out = [{"title": f"Slide bài giảng – Chương {n}", "filename": f"slide_chuong_{n}.pptx", "data": build_pptx(n),
            "material_type": "slide", "outline_id": chapter_ids.get(n)} for n in sorted(SLIDES)]
    out.append({"title": "Giáo trình Hệ quản trị CSDL (trích)", "filename": "giao_trinh_trich.pdf", "data": build_textbook_pdf(),
                "material_type": "textbook", "outline_id": None})
    return out
