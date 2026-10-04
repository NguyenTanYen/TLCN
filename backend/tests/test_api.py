"""Kiểm thử tích hợp API theo các ca sử dụng UC-01 … UC-07 (kịch bản chính + ngoại lệ)."""
import io

from openpyxl import load_workbook
from sqlalchemy import text

from app.config import settings
from app.database import engine

M = settings.moodle_db_name

COURSE = 5  # DBMS330284


def q(sql, **p):
    with engine.connect() as c:
        return c.execute(text(sql), p).fetchall()


# ---------------------------------------------------------------- xác thực & phân quyền
def test_login_wrong_password(client):
    assert client.post("/api/auth/login", json={"username": "gv.son", "password": "sai"}).status_code == 401


def test_requires_token(client):
    assert client.get("/api/exams/2/item-stats").status_code == 401


def test_role_and_ownership(client, gv_other):
    from app.security import create_token
    uid = q("SELECT id FROM users WHERE role='student' LIMIT 1")[0][0]
    sv = {"Authorization": "Bearer " + create_token(uid, "student")}  # kể cả khi có token SV cũng không dùng được API GV
    assert client.get("/api/exams/2/item-stats", headers=gv_other).status_code == 403  # GV không phụ trách lớp
    assert client.get("/api/exams/2/item-stats", headers=sv).status_code == 403        # SV không xem thống kê lớp


# ---------------------------------------------------------------- UC-01 ngân hàng câu hỏi
def _question(**kw):
    body = {"course_id": COURSE, "outline_id": None, "bloom_level_id": 2, "content": "Câu hỏi kiểm thử về khóa chính?",
            "options": [{"content": "A1", "is_correct": True}, {"content": "B1"}, {"content": "C1"}, {"content": "D1"}],
            "clos": [{"clo_id": 1, "weight": 1.0}]}
    body.update(kw)
    return body


def test_question_validation(client, gv):
    bad_two_correct = _question(options=[{"content": "x", "is_correct": True}, {"content": "y", "is_correct": True}])
    assert client.post("/api/questions", json=bad_two_correct, headers=gv).status_code == 422
    bad_weights = _question(clos=[{"clo_id": 1, "weight": 0.5}, {"clo_id": 2, "weight": 0.3}])
    assert client.post("/api/questions", json=bad_weights, headers=gv).status_code == 422
    one_option = _question(options=[{"content": "x", "is_correct": True}])
    assert client.post("/api/questions", json=one_option, headers=gv).status_code == 422


def test_question_crud(client, gv):
    r = client.post("/api/questions", json=_question(clos=[{"clo_id": 1, "weight": 0.6}, {"clo_id": 2, "weight": 0.4}]), headers=gv)
    assert r.status_code == 201, r.text
    qid = r.json()["id"]
    r = client.put(f"/api/questions/{qid}", json=_question(content="Câu hỏi đã sửa nội dung?"), headers=gv)
    assert r.status_code == 200 and r.json()["content"].startswith("Câu hỏi đã sửa")
    assert client.delete(f"/api/questions/{qid}", headers=gv).status_code == 200


def test_question_with_results_is_locked(client, gv):
    # câu 24 đã có kết quả thi -> không được sửa (409), không được xóa (409), chỉ được Hủy
    assert client.put("/api/questions/24", json=_question(), headers=gv).status_code == 409
    assert client.delete("/api/questions/24", headers=gv).status_code == 409


# ---------------------------------------------------------------- UC-02 đề thi
def test_exam_online_xml(client, gv):
    r = client.post("/api/exams", headers=gv, json={"class_section_id": 1, "exam_title": "Quiz kiểm thử", "exam_type": "online",
                                                    "items": [{"question_id": 1, "points": 1}, {"question_id": 2, "points": 1}]})
    assert r.status_code == 201, r.text
    eid = r.json()["id"]
    x = client.get(f"/api/exams/{eid}/moodle-xml", headers=gv)
    assert x.status_code == 200 and b"<idnumber>QB-1</idnumber>" in x.content and x.content.count(b'<question type="multichoice">') == 2
    assert client.get(f"/api/exams/{eid}", headers=gv).json()["status"] == "Published"
    assert client.patch(f"/api/exams/{eid}/link-quiz", json={"moodle_quiz_id": 900}, headers=gv).status_code == 409  # quiz đã gắn bài #2
    assert client.post(f"/api/exams/{eid}/moodle/create-offlinequiz", json={}, headers=gv).status_code == 422  # không phải bài giấy
    assert client.delete(f"/api/exams/{eid}", headers=gv).status_code == 200


def test_exam_rejects_foreign_question(client, gv):
    other = q("SELECT id FROM question_bank WHERE course_id<>:c LIMIT 1", c=COURSE)
    if other:
        r = client.post("/api/exams", headers=gv, json={"class_section_id": 1, "exam_title": "Sai môn", "items": [{"question_id": other[0][0], "points": 1}]})
        assert r.status_code == 422


def test_paper_exam_has_no_system_side_grading(client, gv):
    # Bài giấy: đề in/phiếu/đáp án do Moodle (Offline Quiz) sinh, Moodle nhận diện & chấm – hệ thống không còn nhập phiếu
    paths = client.get("/openapi.json").json()["paths"]
    assert not any("paper-import" in p or "/versions" in p or "answer-key" in p for p in paths)
    assert "/api/exams/{exam_id}/moodle/create-offlinequiz" in paths
    r = client.post("/api/exams/1/moodle/create-offlinequiz", json={"numgroups": 2}, headers=gv)
    assert r.status_code == 409  # đã gắn Offline Quiz 910
    assert client.patch("/api/exams/1/link-quiz", json={"moodle_quiz_id": 12345}, headers=gv).status_code == 422


# ---------------------------------------------------------------- UC-03 đồng bộ / nhập phiếu
def test_moodle_sync_matches_moodle_grades():
    rows = q(f"""SELECT a.total_score, ma.sumgrades FROM exam_attempts a JOIN {M}.mdl_quiz_attempts ma ON ma.id=a.moodle_attempt_id
                WHERE a.exam_id=2""")
    assert len(rows) == 38 and all(round(float(x), 2) == round(float(y), 2) for x, y in rows)


def test_moodle_sync_follows_grade_method_and_skips_preview():
    # quiz mô phỏng dùng grademethod = 1 (điểm cao nhất, hòa -> lượt sớm hơn) như mặc định của Moodle
    graded = {r[0] for r in q(f"""SELECT x.id FROM {M}.mdl_quiz_attempts x WHERE x.quiz=900 AND x.preview=0 AND x.state='finished'
              AND NOT EXISTS (SELECT 1 FROM {M}.mdl_quiz_attempts y WHERE y.quiz=x.quiz AND y.userid=x.userid AND y.preview=0
              AND y.state='finished' AND (y.sumgrades > x.sumgrades OR (y.sumgrades = x.sumgrades AND y.attempt < x.attempt)))""")}
    synced = {r[0] for r in q("SELECT moodle_attempt_id FROM exam_attempts WHERE exam_id=2 AND status='finished'")}
    assert synced == graded and len(synced) == 38


def test_absent_students_recorded():
    assert q("SELECT COUNT(*) FROM exam_attempts WHERE exam_id=2 AND status='absent'")[0][0] == 2
    assert q("""SELECT COUNT(*) FROM item_level_results r JOIN exam_attempts a ON a.id=r.attempt_id
                WHERE a.status='absent' AND r.response_status<>'absent'""")[0][0] == 0


def test_resync_is_idempotent(client, gv):
    before = q("SELECT COUNT(*), SUM(total_score) FROM exam_attempts WHERE exam_id=2")[0]
    r = client.post("/api/exams/2/sync", headers=gv)
    assert r.status_code == 200, r.text
    assert q("SELECT COUNT(*), SUM(total_score) FROM exam_attempts WHERE exam_id=2")[0] == before


def test_sync_warns_average_grade_method(client, gv):
    with engine.begin() as c:
        c.execute(text(f"UPDATE {M}.mdl_quiz SET grademethod=2 WHERE id=900"))
    try:
        r = client.post("/api/exams/2/sync", headers=gv)
        assert r.status_code == 200 and "TRUNG BÌNH" in r.json()["warnings"][0]
    finally:
        with engine.begin() as c:
            c.execute(text(f"UPDATE {M}.mdl_quiz SET grademethod=1 WHERE id=900"))
    assert "warnings" not in client.post("/api/exams/2/sync", headers=gv).json()


def test_sync_stops_when_question_edited_on_moodle(client, gv):
    # GV sửa đáp án đúng của một câu ngay trên Moodle -> cầu nối dừng (409), không giải xáo trộn sai, dữ liệu cũ giữ nguyên
    before = q("SELECT COUNT(*), SUM(total_score) FROM exam_attempts WHERE exam_id=2")[0]
    mq = q("SELECT moodle_question_id FROM question_bank q JOIN exam_questions eq ON eq.question_id=q.id WHERE eq.exam_id=2 LIMIT 1")[0][0]
    ids = [r[0] for r in q(f"SELECT id FROM {M}.mdl_question_answers WHERE question=:m ORDER BY id", m=mq)]
    old = {r[0]: r[1] for r in q(f"SELECT id, fraction FROM {M}.mdl_question_answers WHERE question=:m", m=mq)}
    right = [i for i in ids if float(old[i]) > 0.999][0]
    other = [i for i in ids if i != right][0]
    with engine.begin() as c:
        c.execute(text(f"UPDATE {M}.mdl_question_answers SET fraction = CASE WHEN id=:o THEN 1 ELSE 0 END WHERE question=:m"), {"o": other, "m": mq})
    try:
        r = client.post("/api/exams/2/sync", headers=gv)
        assert r.status_code == 409 and "bị sửa" in r.json()["detail"], r.text
        assert q("SELECT COUNT(*), SUM(total_score) FROM exam_attempts WHERE exam_id=2")[0] == before
    finally:
        with engine.begin() as c:
            for i, fr in old.items():
                c.execute(text(f"UPDATE {M}.mdl_question_answers SET fraction=:f WHERE id=:i"), {"f": fr, "i": i})
    assert client.post("/api/exams/2/sync", headers=gv).status_code == 200


def test_paper_sync_reads_moodle_offlinequiz():
    # tổng điểm = sumgrades Moodle chấm; mỗi SV lấy kết quả complete mới nhất; phiếu partial không tính
    rows = q(f"""SELECT a.total_score, r.sumgrades, r.status FROM exam_attempts a
                 JOIN {M}.mdl_offlinequiz_results r ON r.id=a.moodle_attempt_id AND r.offlinequizid=910 WHERE a.exam_id=1""")
    assert len(rows) == 38 and all(st == "complete" for *_, st in rows)
    assert all(round(float(x), 2) == round(float(y), 2) for x, y, _ in rows)
    latest = {r[0] for r in q(f"""SELECT x.id FROM {M}.mdl_offlinequiz_results x WHERE x.offlinequizid=910 AND x.status='complete'
               AND NOT EXISTS (SELECT 1 FROM {M}.mdl_offlinequiz_results y WHERE y.offlinequizid=910 AND y.userid=x.userid
                               AND y.status='complete' AND y.timemodified > x.timemodified)""")}
    assert {r[0] for r in q("SELECT moodle_attempt_id FROM exam_attempts WHERE exam_id=1 AND status='finished'")} == latest
    assert q("SELECT COUNT(*) FROM exam_attempts WHERE exam_id=1 AND status='absent'")[0][0] == 2


def test_paper_sync_maps_groups_to_versions():
    # mã đề = nhóm A/B của Offline Quiz; mỗi SV ghi đúng mã đề đã làm
    assert {r[0] for r in q("SELECT version_code FROM exam_versions WHERE exam_id=1")} == {"A", "B"}
    bad = q(f"""SELECT COUNT(*) FROM exam_attempts a JOIN exam_versions v ON v.id=a.version_id
                JOIN {M}.mdl_offlinequiz_results r ON r.id=a.moodle_attempt_id AND r.offlinequizid=910
                JOIN {M}.mdl_offlinequiz_groups g ON g.id=r.offlinegroupid WHERE a.exam_id=1 AND v.version_code<>CHAR(64+g.groupnumber)""")
    assert bad[0][0] == 0


def test_paper_reverse_shuffle_per_group():
    # đáp án đúng theo hệ thống = đáp án đúng theo thứ tự đã xáo của từng mã đề trên Moodle
    ok = q(f"""SELECT SUM(r.is_correct), COUNT(*) FROM item_level_results r JOIN exam_attempts a ON a.id=r.attempt_id
               WHERE a.exam_id=1 AND a.status='finished' AND r.response_status='answered'""")[0]
    assert ok[1] > 300 and 0 < int(ok[0]) < ok[1]


def test_paper_sync_requires_offlinequiz(client, gv):
    r = client.post("/api/exams", headers=gv, json={"class_section_id": 1, "exam_title": "Giấy chưa tạo trên Moodle", "exam_type": "paper",
                                                    "items": [{"question_id": 1, "points": 1}]})
    eid = r.json()["id"]
    r = client.post(f"/api/exams/{eid}/sync", headers=gv)
    assert r.status_code == 422 and "Offline Quiz" in r.json()["detail"]
    client.delete(f"/api/exams/{eid}", headers=gv)


def test_paper_resync_is_idempotent(client, gv):
    before = q("SELECT COUNT(*), SUM(total_score) FROM exam_attempts WHERE exam_id=1")[0]
    r = client.post("/api/exams/1/sync", headers=gv)
    assert r.status_code == 200, r.text
    assert q("SELECT COUNT(*), SUM(total_score) FROM exam_attempts WHERE exam_id=1")[0] == before


# ---------------------------------------------------------------- UC-04 thống kê
def test_item_stats_flags_trap_question(client, gv):
    items = {x["question_id"]: x for x in client.get("/api/exams/2/item-stats", headers=gv).json()}
    trap = items[24]
    assert float(trap["di_value"]) < 0 and trap["classification"] == "Cần loại bỏ/kiểm tra đáp án"
    assert sum(d["n"] for d in trap["distribution"]) == 38


def test_clo_results_bm6b(client, gv):
    r = client.get("/api/class-sections/1/clo-results", headers=gv).json()
    assert r["has_data"] and len(r["clos"]) == 4
    for c in r["clos"]:
        assert round(100 * c["n_achieved"] / c["n_evaluated"], 2) == float(c["achieved_pct"])
        assert bool(c["is_achieved"]) == (float(c["achieved_pct"]) >= float(c["target_pct"]))
    clo1 = next(c for c in r["clos"] if c["clo_code"] == "CLO1")
    assert clo1["n_evaluated"] == 76  # evidence_type 'any': cộng dồn 2 bài KT (BM6d)


def test_clo_narrative(client, gv):
    r = client.put("/api/class-sections/1/clo-results/2", headers=gv, json={"analysis": "SV yếu phần giao tác", "improvement": "Bổ sung bài tập"})
    assert r.status_code == 200


# ---------------------------------------------------------------- UC-05: SV xem kết quả trên Moodle
def test_student_cannot_login(client):
    r = client.post("/api/auth/login", json={"username": "22130001", "password": "Sv@123456"})
    assert r.status_code == 403 and "Moodle" in r.json()["detail"]


def test_results_for_moodle(client, gv):
    rows = client.get("/api/exams/2/moodle/results-preview", headers=gv).json()
    done = [r for r in rows if r["status"] == "finished"]
    absent = [r for r in rows if r["status"] == "absent"]
    assert done and absent and all(r["data"] == {} and r["score"] is None for r in absent)  # E2 – chưa đủ bằng chứng
    d = done[0]["data"]
    assert len(d["clos"]) == 4 and all({"code", "pct", "threshold", "class_avg", "achieved"} <= set(c) for c in d["clos"])
    weak = next(r for r in done if r["data"]["items"])
    assert all(i["gap"] > 0 and i["clo"] for i in weak["data"]["items"])
    assert 0 <= done[0]["score"] <= done[0]["max_score"]


def test_push_requires_moodle_config(client, gv, monkeypatch):
    monkeypatch.setattr(settings, "moodle_ws_token", "")
    r = client.post("/api/exams/2/moodle/push-results", headers=gv)
    assert r.status_code == 502 and "Chưa cấu hình" in r.json()["detail"]


def test_publish_requires_analyzed(client, gv):
    r = client.post("/api/exams", headers=gv, json={"class_section_id": 1, "exam_title": "Chưa phân tích", "items": [{"question_id": 3, "points": 1}]})
    eid = r.json()["id"]
    assert client.patch(f"/api/exams/{eid}/publish", json={"publish": True}, headers=gv).status_code == 409
    client.delete(f"/api/exams/{eid}", headers=gv)


# ---------------------------------------------------------------- UC-06/07 CTĐT, PLO/PI, báo cáo
def test_program_structure(client, admin):
    p = client.get("/api/programs/1", headers=admin).json()
    assert len(p["plos"]) == 19  # BM2 KTDL: 19 CĐR (5 nhóm)
    plans = client.get("/api/pi-plans?academic_year=2024-2025", headers=admin).json()
    assert any(x["plo_code"] == "1.3" for x in plans)


def test_plo_summary_is_sum_of_pi(client, admin):
    s = client.get("/api/reports/programs/1/plo-summary?academic_year=2024-2025", headers=admin).json()
    measured = [p for p in s["plos"] if p["n_evaluated"]]
    assert measured
    pr = s["program_result"]
    assert pr["n_evaluated"] == sum(p["n_evaluated"] for p in measured)
    pis = q("""SELECT pi.plo_id, SUM(r.n_evaluated), SUM(r.n_achieved) FROM pi_results r JOIN pi_assessment_plans pl ON pl.id=r.plan_id
               JOIN performance_indicators pi ON pi.id=pl.pi_id JOIN semesters se ON se.id=pl.semester_id
               WHERE se.academic_year='2024-2025' GROUP BY pi.plo_id""")
    by_plo = {p["plo_id"]: p for p in measured}
    for plo_id, n, ok in pis:
        assert by_plo[plo_id]["n_evaluated"] == n and by_plo[plo_id]["n_achieved"] == ok


def test_bm6_workbook(client, gv):
    r = client.get("/api/reports/class-sections/1/bm6.xlsx", headers=gv)
    wb = load_workbook(io.BytesIO(r.content))
    assert r.status_code == 200
    assert any(n.startswith("BM6a") for n in wb.sheetnames) and any(n.startswith("BM6b") for n in wb.sheetnames)
    assert sum(1 for n in wb.sheetnames if n.startswith("BM6c") or n.startswith("BM6d")) == 4
    # bố cục biểu mẫu: mục lục "Biểu mẫu 6", 6a có chữ ký, 6b tiêu đề 2 tầng và số liệu liên kết sang sheet minh chứng
    assert wb.sheetnames[:3] == ["BM6", "BM6a_KH-KQ CĐR môn học", "BM6b_KH-KQ CĐR môn học"]
    assert any("TRƯỞNG ĐƠN VỊ" in str(c.value) for row in wb["BM6a_KH-KQ CĐR môn học"].iter_rows() for c in row if c.value)
    b = wb["BM6b_KH-KQ CĐR môn học"]
    assert "D10:F10" in {str(m) for m in b.merged_cells.ranges} and str(b["D12"].value).startswith("='BM6")
    assert "CÔNG NGHỆ KỸ THUẬT" in str(b["A1"].value)


def test_bm6_evidence_follows_plan_and_matches_bm6b(client, gv):
    """Sheet minh chứng của mỗi CLO chỉ chứa các bài KT đúng loại minh chứng trong BM6a và tổng khớp BM6b."""
    wb = load_workbook(io.BytesIO(client.get("/api/reports/class-sections/1/bm6.xlsx", headers=gv).content))
    plans = dict(q("""SELECT c.clo_code, COALESCE(p.evidence_type,'any') FROM clos c JOIN class_sections cs ON cs.course_id=c.course_id
                      LEFT JOIN clo_assessment_plans p ON p.clo_id=c.id AND p.semester_id=cs.semester_id WHERE cs.id=1"""))
    exams = q("SELECT exam_title, assessment_type FROM exams WHERE class_section_id=1 AND status='Analyzed'")
    for name in wb.sheetnames[3:]:
        clo = name.split("_", 1)[1]
        text_ = " ".join(str(c.value) for row in wb[name].iter_rows() for c in row if c.value)
        for title, typ in exams:
            used = plans[clo] in ("any", typ)
            assert (f": {title}" in text_) == used, (name, title)
        n_used = sum(1 for _, typ in exams if plans[clo] in ("any", typ))
        assert name.startswith("BM6d" if n_used > 1 else "BM6c")


def test_bm6c_per_exam(client, gv, gv_other):
    for eid in (1, 2):
        r = client.get(f"/api/reports/exams/{eid}/bm6c.xlsx", headers=gv)
        assert r.status_code == 200 and "BM6c_" in r.headers["content-disposition"]
        wb = load_workbook(io.BytesIO(r.content))
        clos = [x[0] for x in q("""SELECT DISTINCT c.clo_code FROM attempt_clo_results r JOIN exam_attempts a ON a.id=r.attempt_id
                                   JOIN clos c ON c.id=r.clo_id WHERE a.exam_id=:e AND a.status='finished' ORDER BY 1""", e=eid)]
        assert wb.sheetnames == ["TongHop_BaiKT"] + [f"BM6c_{c}" for c in clos]
        title = q("SELECT exam_title FROM exams WHERE id=:e", e=eid)[0][0]
        for c in clos:   # mỗi sheet chỉ có đúng bài KT này, đủ số SV đã làm bài, đúng số SV đạt
            ws = wb[f"BM6c_{c}"]
            vals = [x.value for row in ws.iter_rows() for x in row if x.value is not None]
            assert any(str(v).startswith("Bài KT 1: " + title) for v in vals)
            assert not any(str(v).startswith("Bài KT 2:") for v in vals)
            n, ok = q("""SELECT COUNT(*), SUM(r.is_achieved) FROM attempt_clo_results r JOIN exam_attempts a ON a.id=r.attempt_id
                         JOIN clos c ON c.id=r.clo_id WHERE a.exam_id=:e AND c.clo_code=:c AND a.status='finished'""", e=eid, c=c)[0]
            summary = {row[1].value.split(":")[0]: (row[2].value, row[3].value) for row in wb["TongHop_BaiKT"].iter_rows()
                       if isinstance(row[1].value, str) and row[1].value.startswith("CLO")}
            assert summary[c] == (int(ok), n)
            mssv = [row[1].value for row in ws.iter_rows() if isinstance(row[0].value, int)]
            assert len(mssv) == n and all(mssv)
    assert client.get("/api/reports/exams/2/bm6c.xlsx", headers=gv_other).status_code == 403
    assert client.get("/api/reports/exams/99999/bm6c.xlsx", headers=gv).status_code == 404
    r = client.post("/api/exams", headers=gv, json={"class_section_id": 1, "exam_title": "Chưa phân tích BM6c", "items": [{"question_id": 3, "points": 1}]})
    assert client.get(f"/api/reports/exams/{r.json()['id']}/bm6c.xlsx", headers=gv).status_code == 409


def test_bm3_workbook(client, admin, gv):
    assert client.get("/api/reports/programs/1/bm3.xlsx?academic_year=2024-2025", headers=gv).status_code == 403
    r = client.get("/api/reports/programs/1/bm3.xlsx?academic_year=2024-2025", headers=admin)
    wb = load_workbook(io.BytesIO(r.content))
    assert r.status_code == 200 and wb.sheetnames[:2] == ["BM2a_KeHoach", "BM2b_TongKet"]
    assert any(n.startswith("BM3b_") for n in wb.sheetnames) and any(n.startswith("BM3c_") for n in wb.sheetnames)


def test_bm2_docx(client, admin, gv):
    from docx import Document
    assert client.get("/api/reports/programs/1/bm2.docx?academic_year=2024-2025", headers=gv).status_code == 403
    r = client.get("/api/reports/programs/1/bm2.docx?academic_year=2024-2025", headers=admin)
    assert r.status_code == 200
    doc = Document(io.BytesIO(r.content))
    text_all = "\n".join(p.text for p in doc.paragraphs)
    assert "KẾ HOẠCH" in text_all and "BÁO CÁO TỔNG KẾT" in text_all and "(năm học trước)" in text_all
    assert "CÔNG NGHỆ KỸ THUẬT" in doc.tables[0].cell(0, 0).text


def test_assignments_xlsx(client, gv):
    sem = q("SELECT semester_id FROM assessment_assignments WHERE all_supervisors=1 LIMIT 1")[0][0]
    r = client.get(f"/api/reports/semesters/{sem}/assignments.xlsx", headers=gv)
    ws = load_workbook(io.BytesIO(r.content)).active
    assert r.status_code == 200 and str(ws["A1"].value).startswith("PHÂN CÔNG ĐÁNH GIÁ PIs HỌC KỲ")
    assert [ws.cell(5, j).value for j in range(1, 6)] == ["STT", "MÃ MH", "TÊN MH", "GV ĐÁNH GIÁ", "GHI CHÚ"]
    gv_col = [ws.cell(i, 4).value for i in range(6, ws.max_row + 1)]
    assert gv_col.count("Tất cả thầy/cô có hướng dẫn") == 2      # PODE434277, POIS431184
    n = q("SELECT COUNT(*) FROM assessment_assignments WHERE semester_id=:s", s=sem)[0][0]
    assert sum(1 for v in gv_col if v) == n
    assert client.get("/api/reports/semesters/999999/assignments.xlsx", headers=gv).status_code == 404


def test_plo_narrative(client, admin):
    plo_id = q("SELECT plo_id FROM plo_results LIMIT 1")[0][0]
    r = client.put(f"/api/reports/plo-results/{plo_id}/2024-2025", headers=admin,
                   json={"analysis": "a", "improvement_actions": "b", "improvement_results": "c", "evidence_tools": "d"})
    assert r.status_code == 200
