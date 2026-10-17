"""Thời gian làm bài trên Moodle Quiz: phát hiện SV chọn đại rồi nộp, chọn xong không nộp để hết giờ (dữ liệu mô phỏng có
2 SV chọn đại 20 câu trong ~1 phút và 2 SV chọn xong trong ~6 phút rồi để Moodle tự nộp khi hết 45 phút)."""
import io
from datetime import datetime, timedelta

from openpyxl import load_workbook
from sqlalchemy import text

from app.database import engine
from app.services import timing


def q(sql, **p):
    with engine.connect() as c:
        return c.execute(text(sql), p).fetchall()


def test_response_time_split_and_page_resolution():
    t0 = datetime(2025, 1, 8, 7, 30)
    at = lambda s: t0 + timedelta(seconds=s)
    assert timing.response_times(t0, [at(30), at(75), at(80)]) == [30.0, 45.0, 5.0]
    # 3 câu cùng một trang được lưu cùng lúc ở giây 90 -> chia đều 90 giây
    assert timing.response_times(t0, [at(90), at(90), at(90), at(120)]) == [30.0, 30.0, 30.0, 30.0]
    assert timing.response_times(None, [at(10), at(15)]) == [0.0, 5.0]
    assert timing.response_times(t0, []) == []


def test_sync_stores_answer_times():
    att = q("""SELECT a.time_limit_sec, a.finished_by_timeout, a.last_answer_at, s.student_code FROM exam_attempts a
               JOIN students s ON s.id=a.student_id WHERE a.exam_id=2 AND a.status='finished' ORDER BY s.student_code""")
    assert len(att) == 38 and all(r[0] == 2700 and r[2] is not None for r in att)
    assert sorted(r[3] for r in att if r[1]) == ["22130006", "22130023"]                  # nộp do hết giờ
    n = q("""SELECT COUNT(*), SUM(r.response_time_sec IS NOT NULL), SUM(r.n_answer_saves > 1) FROM item_level_results r
             JOIN exam_attempts a ON a.id=r.attempt_id WHERE a.exam_id=2 AND r.response_status='answered'""")[0]
    assert n[0] == n[1] and n[2] > 0                                                     # mọi câu đã trả lời có thời gian; có đổi đáp án
    # tổng thời gian trả lời = thời điểm chọn đáp án lần đầu cuối cùng – lúc bắt đầu
    for code in ("22130001", "22130017"):
        tot, start, last = q("""SELECT SUM(r.response_time_sec), a.started_at, MAX(r.first_answer_at) FROM item_level_results r
                                JOIN exam_attempts a ON a.id=r.attempt_id JOIN students s ON s.id=a.student_id
                                WHERE a.exam_id=2 AND s.student_code=:c GROUP BY a.id""", c=code)[0]
        assert abs(float(tot) - (last - start).total_seconds()) < 1
    # bài giấy: thời điểm nhận diện phiếu không phải thời gian làm bài -> không lưu
    assert q("""SELECT COUNT(*) FROM item_level_results r JOIN exam_attempts a ON a.id=r.attempt_id
                WHERE a.exam_id=1 AND r.first_answer_at IS NOT NULL""")[0][0] == 0


def test_timing_report_flags_planted_students(client, gv, gv_other):
    r = client.get("/api/exams/2/timing", headers=gv)
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["available"] and d["n_questions"] == 20 and d["time_limit_min"] == 45 and not d["page_resolution"]
    warn = {a["student_code"]: [f["code"] for f in a["flags"]] for a in d["attempts"] if a["level"] == "canh_bao"}
    assert warn == {"22130006": ["TREO_MAY"], "22130023": ["TREO_MAY"], "22130017": ["LAM_NHANH"], "22130018": ["LAM_NHANH"]}
    assert d["summary"]["n_warning"] == 4 and d["summary"]["n_timeout"] == 2 and d["summary"]["flags"]["TREO_MAY"] == 2
    a17 = next(a for a in d["attempts"] if a["student_code"] == "22130017")
    assert a17["duration_min"] < 2 and a17["sec_per_q"] < 6 and a17["rapid_share"] > 0.9 and a17["rapid_accuracy"] < 0.5
    a06 = next(a for a in d["attempts"] if a["student_code"] == "22130006")
    assert a06["finished_by_timeout"] and a06["active_min"] < 8 and a06["idle_min"] > 35
    normal = next(a for a in d["attempts"] if a["student_code"] == "22130001")
    assert normal["level"] is None and 15 < normal["duration_min"] < 45
    assert all(x["median_sec"] and x["median_sec"] > 20 for x in d["questions"])
    # danh sách bài KT của lớp và trang bài KT hiển thị số cảnh báo
    ex = next(e for e in client.get("/api/class-sections/1/exams", headers=gv).json() if e["id"] == 2)
    assert ex["timing_alerts"] == {"n_warning": 4, "n_notice": 0}
    assert client.get("/api/exams/2", headers=gv).json()["timing_alerts"]["n_warning"] == 4
    assert client.get("/api/exams/2/timing", headers=gv_other).status_code == 403


def test_timing_thresholds_timeline_and_export(client, gv):
    # nới ngưỡng không thao tác (50 phút) -> hết cảnh báo treo máy; siết giây/câu (100) -> nhiều SV bị coi là làm nhanh
    d = client.get("/api/exams/2/timing?idle_min=50", headers=gv).json()
    assert d["summary"]["flags"]["TREO_MAY"] == 0 and d["thresholds"]["idle_min"] == 50
    assert client.get("/api/exams/2/timing?min_sec_per_q=100", headers=gv).json()["summary"]["flags"]["LAM_NHANH"] > 10
    assert client.get("/api/exams/2/timing?rapid_share=2", headers=gv).status_code == 422
    sid = q("SELECT id FROM students WHERE student_code='22130006'")[0][0]
    t = client.get(f"/api/exams/2/timing/students/{sid}", headers=gv).json()
    assert t["finished_by_timeout"] and t["last_answer_min"] < 8 and t["duration_min"] == 45 and len(t["items"]) == 20
    assert all(i["response_time_sec"] is None or i["response_time_sec"] < 30 for i in t["items"])
    absent = q("""SELECT a.student_id FROM exam_attempts a WHERE a.exam_id=2 AND a.status='absent' LIMIT 1""")[0][0]
    assert client.get(f"/api/exams/2/timing/students/{absent}", headers=gv).status_code == 404
    x = client.get("/api/exams/2/timing.xlsx", headers=gv)
    assert x.status_code == 200
    ws = load_workbook(io.BytesIO(x.content))["Canh_bao"]
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=6)]
    assert {r[0] for r in rows if r[2] == "Cảnh báo"} == {"22130006", "22130023", "22130017", "22130018"}
    # bài giấy, bài chưa có dữ liệu thời gian
    p = client.get("/api/exams/1/timing", headers=gv).json()
    assert p["available"] is False and "Moodle Quiz" in p["reason"]
    assert client.get("/api/exams/1/timing.xlsx", headers=gv).status_code == 409
    # chỉ số câu hỏi có trung vị thời gian trả lời; hồ sơ SV có thời gian từng câu
    st = client.get("/api/exams/2/item-stats", headers=gv).json()
    assert all(s["median_time_sec"] for s in st)
    prof = client.get(f"/api/class-sections/1/students/{sid}/profile", headers=gv).json()
    ans = next(e for e in prof["exams"] if e["id"] == 2)["answers"]
    assert any(a["response_time_sec"] is not None for a in ans)
