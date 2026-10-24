"""Chụp màn hình giao diện (Playwright) phục vụ báo cáo + phát hiện lỗi console.
Chạy trên dữ liệu minh họa (sau 4_TAO_DU_LIEU_DEMO_TREN_MOODLE):  python tests/ui_screens.py [http://localhost:8000]
Các màn hình nhập câu hỏi / gán CLO thêm câu tạm vào ngân hàng rồi xóa lại; màn hình tạo đề giấy tạo bài KT tạm rồi xóa.
"""
import io
import sys
from pathlib import Path

import requests
from openpyxl import Workbook, load_workbook
from playwright.sync_api import sync_playwright

BASE = sys.argv[1] if len(sys.argv) > 1 else "http://localhost:8000"
OUT = Path(__file__).resolve().parents[2] / "docs" / "screenshots"
OUT.mkdir(parents=True, exist_ok=True)
TMP = OUT / "_tmp"; TMP.mkdir(exist_ok=True)
errors = []


def login(page, u, p):
    page.goto(BASE + "/login"); page.evaluate("localStorage.clear()"); page.goto(BASE + "/login")
    page.fill("input >> nth=0", u); page.fill("input[type=password]", p); page.click("button.btn.primary")
    page.wait_for_selector(".topbar")


def shot(page, name, url=None, click=None, wait=".card", full=True):
    if url:
        page.goto(BASE + url)
    if click:
        page.click(click)
    try:
        page.wait_for_selector(wait, timeout=15000)
    except Exception:
        print("BỎ QUA (cần dữ liệu Moodle thật):", name); return
    page.wait_for_timeout(900)
    page.screenshot(path=str(OUT / f"{name}.png"), full_page=full)
    print("saved", name)


def api_token(u, p):
    return {"Authorization": "Bearer " + requests.post(BASE + "/api/auth/login", json={"username": u, "password": p}).json()["access_token"]}


def sample_questions() -> Path:
    """Tệp Excel minh họa: 5 câu chưa gán CLO/Bloom, 1 câu trùng, 1 câu thiếu đáp án."""
    wb = Workbook(); ws = wb.active
    ws.append(["STT", "Nội dung câu hỏi", "Phương án A", "Phương án B", "Phương án C", "Phương án D", "Đáp án đúng"])
    rows = [("Hàm nào trả về số dòng bị ảnh hưởng bởi câu lệnh gần nhất?", "@@ROWCOUNT", "@@ERROR", "@@IDENTITY", "SCOPE_IDENTITY()", "A"),
            ("Bảng ảo nào chứa dữ liệu mới trong trigger AFTER UPDATE?", "deleted", "inserted", "updated", "changed", "B"),
            ("Lệnh nào hoàn tác một giao tác đang thực hiện?", "COMMIT", "SAVE TRAN", "ROLLBACK", "END TRAN", "C"),
            ("Chỉ mục clustered quyết định điều gì của bảng?", "Thứ tự lưu vật lý của các dòng", "Quyền truy cập", "Kích thước cột", "Tên bảng", "A"),
            ("Quyền nào cho phép đọc dữ liệu của bảng?", "INSERT", "SELECT", "DELETE", "ALTER", "B"),
            ("Lệnh nào hoàn tác một giao tác đang thực hiện?", "COMMIT", "SAVE TRAN", "ROLLBACK", "END TRAN", "C"),
            ("Câu hỏi thiếu đáp án đúng", "X", "Y", "", "", "")]
    for i, r in enumerate(rows, 1):
        ws.append([i, *r])
    f = TMP / "cau_hoi_moi.xlsx"; wb.save(f)
    return f


with sync_playwright() as pw:
    import os
    b = pw.chromium.launch(executable_path=os.environ.get('PW_CHROME') or None)
    ctx = b.new_context(viewport={"width": 1440, "height": 900}, device_scale_factor=1)
    page = ctx.new_page()
    page.on("console", lambda m: m.type == "error" and errors.append(f"{page.url}: {m.text}"))
    page.on("pageerror", lambda e: errors.append(f"{page.url}: {e}"))
    page.goto(BASE + "/login"); page.wait_for_selector(".login"); page.screenshot(path=str(OUT / "01_login.png")); print("saved login")
    login(page, "gv.son", "Gv@123456")
    gv = api_token("gv.son", "Gv@123456")
    shot(page, "02_lop_hoc_phan", "/")
    shot(page, "03_bai_kiem_tra", "/sections/1")
    shot(page, "04_ngan_hang_cau_hoi", "/questions", wait=".tbl tbody tr")
    page.locator("button:has-text(\"Chi tiết\")").last.click(); page.wait_for_selector(".modal"); page.wait_for_timeout(400)
    page.screenshot(path=str(OUT / "05_cau_hoi_bi_khoa.png")); page.keyboard.press("Escape"); page.goto(BASE + "/questions"); page.wait_for_selector(".tbl tbody tr")
    page.click("text=+ Thêm câu hỏi"); page.wait_for_selector(".modal"); page.wait_for_timeout(300)
    page.screenshot(path=str(OUT / "06_them_cau_hoi.png"))
    page.goto(BASE + "/sections/1"); page.wait_for_selector(".card"); page.click("text=+ Tạo bài kiểm tra"); page.click(".choice:has-text('Trắc nghiệm trực tuyến')")
    page.fill(".modal input >> nth=0", "Kiểm tra giữa kỳ (minh họa)"); page.click("text=Tiếp: chọn câu hỏi"); page.click("text=Tự chọn câu từ ngân hàng"); page.wait_for_selector(".modal tbody tr")
    page.click(".modal button.chip:has-text('CLO2')"); page.click("text=Chọn tất cả đang lọc"); page.wait_for_timeout(300)
    page.screenshot(path=str(OUT / "07_tao_bai_kt.png"))
    shot(page, "08_quy_trinh_online", "/exams/2")
    shot(page, "09_quy_trinh_giay", "/exams/1?step=release", wait="text=Bài đã chấm trên Moodle")
    shot(page, "10_tong_quan_diem", "/exams/2", click="text=Tổng quan điểm", wait="canvas")
    # tạo bài kiểm tra ở trang môn học (đề của môn) – chọn lớp HP khi xuất
    cid = next(x["course_id"] for x in requests.get(BASE + "/api/class-sections", headers=gv).json() if x["id"] == 1)
    page.goto(BASE + f"/courses/{cid}?tab=exams"); page.wait_for_selector("text=Đề dùng chung cho nhiều lớp")
    page.click(".page-h >> text=+ Tạo bài kiểm tra"); page.wait_for_selector(".choice"); page.wait_for_timeout(400)
    page.screenshot(path=str(OUT / "45_chon_hinh_thuc_bai_kt.png")); print("saved chooser"); page.keyboard.press("Escape")
    tpl = requests.post(BASE + "/api/exams", headers=gv, json={"course_id": cid, "exam_title": "Kiểm tra giữa kỳ – đề chung của môn", "exam_type": "online",
                        "assessment_type": "process", "max_score": 10, "duration_minutes": 30,
                        "items": [{"question_id": q, "points": 1} for q in range(1, 11)]}).json()["id"]
    shot(page, "46_de_cua_mon", f"/exams/{tpl}", wait="text=Đề chưa gắn lớp học phần", full=False)
    page.click("text=Giao cho lớp HP / Xuất >> nth=0"); page.wait_for_selector(".pick"); page.click(".pick:has-text('DBMS330284_02') input"); page.wait_for_timeout(300)
    page.screenshot(path=str(OUT / "47_giao_de_cho_lop.png")); print("saved deploy"); page.keyboard.press("Escape")
    requests.delete(BASE + f"/api/exams/{tpl}", headers=gv)
    shot(page, "42_thoi_gian_lam_bai", "/exams/2?tab=timing", wait="text=Cách hệ thống cảnh báo")
    page.click("tr:has-text('22130017') >> text=Diễn biến"); page.wait_for_selector(".modal canvas"); page.wait_for_timeout(800)
    page.screenshot(path=str(OUT / "43_dien_bien_lam_bai.png")); page.keyboard.press("Escape"); print("saved timeline")
    shot(page, "44_canh_bao_lop", "/sections/1", wait="text=cảnh báo", full=False)
    page.goto(BASE + "/exams/2"); page.click("text=Phân tích câu hỏi"); page.wait_for_selector("canvas")
    page.click("tr:has-text('YEAR(NgayLap)') >> text=Phương án"); page.wait_for_timeout(600)
    page.screenshot(path=str(OUT / "11_phan_tich_cau_hoi.png"), full_page=True); print("saved items")
    shot(page, "12_ket_qua_sinh_vien", "/exams/2", click="text=Kết quả sinh viên", wait=".tbl tbody tr")
    shot(page, "13_nhat_ky_dong_bo", "/exams/2", click="button:has-text(\"Nhật ký\")", wait=".tbl tbody tr")
    shot(page, "14_bm6b", "/sections/1/clo", wait="canvas")
    shot(page, "15_bm6a", "/sections/1/clo?tab=bm6a", wait=".tbl tbody tr")
    try:
        page.goto(BASE + "/sections/1/clo?tab=bm6b"); page.wait_for_selector(".tbl tbody tr"); page.locator("button:has-text('Minh chứng'):not([disabled])").first.click(); page.wait_for_selector(".modal"); page.wait_for_timeout(500)
        page.screenshot(path=str(OUT / "16_minh_chung_bm6c.png")); print("saved evidence")
    except Exception:
        print("BỎ QUA 16_minh_chung_bm6c (cần kế hoạch BM6a + dữ liệu Moodle)")
    # thống kê lớp và hồ sơ từng sinh viên
    shot(page, "21_thong_ke_lop", "/sections/1/stats?exam=2", wait="canvas")
    sid = requests.get(BASE + "/api/class-sections/1/students", headers=gv).json()[0]["id"]
    shot(page, "22_ho_so_sinh_vien", f"/sections/1/students/{sid}?exam=2", wait=".tbl")
    # bài thi giấy mới: màn hình tạo đề trên Moodle (bài KT tạm, xóa sau khi chụp)
    tmp_exam = requests.post(BASE + "/api/exams", headers=gv, json={"class_section_id": 1, "exam_title": "Kiểm tra giữa kỳ (giấy)", "exam_type": "paper",
                             "assessment_type": "process", "max_score": 10, "duration_minutes": 30,
                             "items": [{"question_id": q, "points": 1} for q in (2, 3, 6, 7, 9, 11, 12, 13, 16, 17)]}).json()["id"]
    shot(page, "25_tao_de_giay", f"/exams/{tmp_exam}", wait="text=In và chấm ngay trên hệ thống", full=False)
    requests.delete(BASE + f"/api/exams/{tmp_exam}", headers=gv)
    tmp_exam = requests.post(BASE + "/api/exams", headers=gv, json={"class_section_id": 1, "exam_title": "Kiểm tra giữa kỳ", "exam_type": "online",
                             "assessment_type": "process", "max_score": 10, "duration_minutes": 30,
                             "items": [{"question_id": q, "points": 1} for q in range(1, 11)]}).json()["id"]
    shot(page, "08b_tao_quiz_api", f"/exams/{tmp_exam}", wait="text=Tạo bài trắc nghiệm trên Moodle", full=False)
    requests.delete(BASE + f"/api/exams/{tmp_exam}", headers=gv)
    # nhập câu hỏi hàng loạt rồi gán CLO từ file (thêm câu tạm, xóa ở cuối)
    f = sample_questions()
    page.goto(BASE + "/questions"); page.wait_for_selector(".tbl tbody tr")
    page.click("text=Nhập câu hỏi từ file"); page.wait_for_selector(".modal")
    page.set_input_files(".modal input[type=file]", str(f)); page.wait_for_selector(".modal .stats"); page.wait_for_timeout(500)
    page.screenshot(path=str(OUT / "23_nhap_cau_hoi.png")); print("saved import")
    page.click(".modal button.primary:has-text('Nhập')"); page.wait_for_selector(".modal .alert.green"); page.wait_for_timeout(300)
    page.goto(BASE + "/questions"); page.wait_for_selector(".alert.amber"); page.wait_for_timeout(600)
    page.screenshot(path=str(OUT / "04b_ngan_hang_chua_gan.png")); print("saved untagged")
    tpl = requests.get(BASE + "/api/questions/tagging-template?course_id=5&only_untagged=true", headers=gv).content
    wb = load_workbook(io.BytesIO(tpl)); ws = wb["GanCLO"]
    fill = [(2, "3 - Vận dụng", "CLO2"), (3, "2 - Hiểu", "CLO2:0.6; CLO4:0.4"), (4, "Vận dụng", "CLO3"), (5, "4 - Phân tích", "CLO4; CLO9"), (6, "2 - Hiểu", "CLO4")]
    new_ids = []
    for row, (ch, bl, clo) in zip(ws.iter_rows(min_row=2), fill):
        row[3].value, row[4].value, row[5].value = ch, bl, clo; new_ids.append(row[0].value)
    tagged = TMP / "gan_clo_da_dien.xlsx"; wb.save(tagged)
    page.click("button:has-text('Gán CLO, mức độ nhận thức từ file') >> nth=0"); page.wait_for_selector(".modal")
    page.set_input_files(".modal input[type=file]", str(tagged)); page.wait_for_selector(".modal .stats"); page.wait_for_timeout(500)
    page.screenshot(path=str(OUT / "24_gan_clo_tu_file.png")); print("saved tagging")
    for q in new_ids:
        requests.delete(BASE + f"/api/questions/{q}", headers=gv)
    # chức năng bổ sung 15/10
    shot(page, "31_mon_hoc_cua_toi", "/", wait=".secrow")
    shot(page, "32_de_cuong_chi_tiet", "/courses/5?tab=syllabus", wait=".syl")
    page.goto(BASE + "/courses/5?tab=combine"); page.wait_for_selector("text=Phân tích gộp"); page.click("text=Phân tích gộp"); page.wait_for_selector("text=2. Mức đạt CĐR môn học")
    page.wait_for_timeout(500); page.screenshot(path=str(OUT / "33_gop_bai_kt.png"), full_page=True)
    shot(page, "34_phu_dao", "/sections/1/remedial", wait=".tbl tbody tr", full=False)
    page.click(".tabs button:has-text('Theo sinh viên')"); page.wait_for_selector(".tbl tbody tr"); page.wait_for_timeout(300)
    page.click("button:has-text('Câu sai ›') >> nth=0"); page.wait_for_selector(".modal .tbl tbody tr"); page.wait_for_timeout(400)
    page.screenshot(path=str(OUT / "35_phu_dao_cau_sai.png")); page.keyboard.press("Escape")
    man = [e for e in requests.get(BASE + "/api/class-sections/2/exams", headers=gv).json() if e["exam_type"] == "manual"]
    if man:
        shot(page, "36_nhap_diem", f"/exams/{man[0]['id']}", wait=".grade-grid", full=False)
    page.goto(BASE + "/sections/1"); page.wait_for_selector(".card"); page.click("text=+ Tạo bài kiểm tra"); page.click(".choice:has-text('Nhập điểm')"); page.wait_for_timeout(600)
    page.screenshot(path=str(OUT / "37_tao_bai_nhap_diem.png")); page.keyboard.press("Escape")
    qs = [q for q in requests.get(BASE + "/api/questions?course_id=5", headers=gv).json() if q["tagged"]][:12]
    pe = requests.post(BASE + "/api/exams", headers=gv, json={"class_section_id": 1, "exam_title": "Kiểm tra giữa kỳ (giấy, chấm trên hệ thống)", "exam_type": "paper",
                                                              "items": [{"question_id": q["id"], "points": 1} for q in qs]}).json()["id"]
    requests.post(BASE + f"/api/exams/{pe}/paper/versions", headers=gv, json={"n_versions": 2})
    requests.put(BASE + f"/api/exams/{pe}/paper/settings", headers=gv, json={"paper": {"columns": 2, "duration": "30 phút"}, "sheet": {}})
    page.goto(BASE + f"/exams/{pe}"); page.wait_for_selector(".paper-preview"); page.wait_for_timeout(700)
    page.screenshot(path=str(OUT / "38_de_giay_xem_truoc.png"), full_page=True)
    page.click("button:has-text('3. Phiếu trả lời')"); page.wait_for_selector(".sheet-img"); page.wait_for_timeout(700)
    page.screenshot(path=str(OUT / "39_phieu_tra_loi.png"), full_page=False)
    # chấm phiếu: tô phiếu (in PDF -> ảnh), một phiếu chụp nghiêng, một phiếu tô 2 ô, một phiếu gạch chéo đổi đáp án
    import random as _r, sys as _s
    _s.path.insert(0, str(Path(__file__).resolve().parents[1]))
    import numpy as np, cv2, fitz
    from app.services import paper as _P
    stp = requests.get(BASE + f"/api/exams/{pe}/paper", headers=gv).json()
    vs = stp["versions"]; nq = len(vs[0]["questions"])
    per = [max(len(v["questions"][i]["options"]) for v in vs) for i in range(nq)]
    so = _P.SheetOptions(n_questions=nq, options_per_question=per, version_codes=[v["version_code"] for v in vs])
    roster = [x["student_code"] for x in requests.get(BASE + "/api/class-sections/1/students", headers=gv).json()][:4]
    rg = _r.Random(3); fills = []
    for i, code in enumerate(roster):
        v = vs[i % len(vs)]
        fills.append({"student": code, "version": v["version_code"],
                      "answers": {q["position"]: (q["answer"] if rg.random() < .75 else rg.choice([o["label"] for o in q["options"]])) for q in v["questions"]}})
    L = _P.sheet_layout(so); S_ = 150 / 25.4
    pos = {(b["q"], b["label"]): (int(b["x"] * S_), int(b["y"] * S_)) for b in L["answers"]}; rr = int(L["r"] * S_)
    imgs = []
    for k, pg in enumerate(fitz.open(stream=_P.sheet_pdf(so, prefills=fills), filetype="pdf")):
        px = pg.get_pixmap(matrix=fitz.Matrix(150 / 72, 150 / 72), colorspace=fitz.csGRAY)
        g = np.frombuffer(px.samples, np.uint8).reshape(px.height, px.width).copy()
        if k == 1:      # tô thêm một ô ở câu 3 -> cảnh báo tô 2 ô
            lab = [l for (q, l) in pos if q == 3 and l != fills[1]["answers"][3]][0]; cv2.circle(g, pos[(3, lab)], int(rr * .9), 40, -1)
        if k == 2:      # gạch chéo một ô khác ở câu 5 (đổi đáp án đúng cách)
            lab = [l for (q, l) in pos if q == 5 and l != fills[2]["answers"][5]][0]; x, y = pos[(5, lab)]
            cv2.line(g, (x - rr, y - rr), (x + rr, y + rr), 30, 2); cv2.line(g, (x - rr, y + rr), (x + rr, y - rr), 30, 2)
        if k == 3:      # ảnh chụp nghiêng
            h, w = g.shape; g = cv2.warpAffine(g, cv2.getRotationMatrix2D((w / 2, h / 2), 6, .9), (w, h), borderValue=110)
        f = TMP / f"phieu_{k + 1}.jpg"; cv2.imwrite(str(f), g); imgs.append(str(f))
    page.click("button:has-text('4. Chấm phiếu')"); page.wait_for_timeout(300)
    page.set_input_files("input[type=file][multiple]", imgs); page.wait_for_selector(".scan-grid img", timeout=60000); page.wait_for_timeout(600)
    page.click(".tbl tbody tr >> nth=1"); page.wait_for_timeout(400)
    page.screenshot(path=str(OUT / "41_cham_phieu.png"), full_page=True); print("saved scan")
    requests.delete(BASE + f"/api/exams/{pe}/paper/versions", headers=gv); requests.delete(BASE + f"/api/exams/{pe}", headers=gv)
    page.goto(BASE + "/questions"); page.wait_for_selector(".tbl tbody tr"); page.click("text=Đồng bộ với Moodle"); page.wait_for_timeout(1500)
    page.screenshot(path=str(OUT / "40_dong_bo_cau_hoi.png")); page.keyboard.press("Escape")
    login(page, "admin", "Admin@123")
    shot(page, "17_ctdt_plo_pi", "/programs/1")
    shot(page, "26_ds_ctdt", "/programs", wait=".card")
    shot(page, "27_ctdt_2026_plo", "/programs/3", wait=".tbl tbody tr")
    shot(page, "28_ma_tran_hp_pi", "/programs/3?tab=matrix", wait=".tbl tbody tr", full=False)
    shot(page, "29_hoc_phan_theo_hk", "/programs/3?tab=courses", wait=".tbl tbody tr")
    shot(page, "30_khai_bao_hoc_phan", "/courses/5?tab=setup", wait=".tbl tbody tr")
    page.evaluate("localStorage.setItem('tlcn_program', '1')")      # CTĐT KTDL (BM2 minh họa)
    shot(page, "18_ke_hoach_pi_bm3b", "/pi-plans", wait=".tbl tbody tr")
    shot(page, "19_tong_hop_plo", "/plo-summary", wait="canvas")
    shot(page, "20_ket_noi_moodle", "/system", wait=".tbl")
    b.close()

print("\nConsole errors:" if errors else "\nKhông có lỗi console.")
for e in errors:
    print(" -", e)
