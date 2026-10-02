"""Nhập hàng loạt câu hỏi (Excel / Aiken / Moodle XML) và gán CLO / Bloom / chương từ file Excel."""
import io

from openpyxl import Workbook, load_workbook

from app.services import question_import as qi


def _course(client, gv):
    return client.get("/api/class-sections", headers=gv).json()[0]["course_id"]


def _xlsx(rows):
    wb = Workbook(); ws = wb.active
    for r in rows:
        ws.append(r)
    b = io.BytesIO(); wb.save(b)
    return b.getvalue()


def _up(client, gv, url, name, data):
    return client.post(url, headers=gv, files={"file": (name, data)})


def test_parsers():
    clos = {"clo1": 1, "clo2": 2, "clo3": 3}
    assert qi.parse_clos("CLO2", clos) == ([(2, 1.0)], None)
    assert qi.parse_clos("CLO1; CLO3", clos) == ([(1, 0.5), (3, 0.5)], None)
    assert qi.parse_clos("clo1, clo2, clo3", clos)[0] == [(1, 0.33), (2, 0.33), (3, 0.34)]
    assert qi.parse_clos("CLO1:0.6; CLO3:0.4", clos)[0] == [(1, 0.6), (3, 0.4)]
    assert qi.parse_clos("2", clos)[0] == [(2, 1.0)]
    assert qi.parse_clos("CLO1:0.6; CLO3:0.6", clos)[1]           # tổng ≠ 1
    assert qi.parse_clos("CLO1:0.6; CLO3", clos)[1]               # thiếu trọng số
    assert qi.parse_clos("CLO9", clos)[1]
    levels = [{"id": 3, "code": "Apply", "name_vi": "Vận dụng"}]
    assert qi.parse_bloom("3 - Vận dụng", levels)[0] == 3
    assert qi.parse_bloom("van dung", levels)[0] == 3 and qi.parse_bloom("APPLY", levels)[0] == 3
    assert qi.parse_chapter("Chương 2", {2: 20})[0] == 20 and qi.parse_chapter("9", {2: 20})[1]
    aiken = qi._from_aiken(["Câu 1: Hỏi gì?", "A. x", "B) y", "ANSWER: B", "CLO: CLO2", "", "Câu hai", "nhiều dòng", "A. p", "B. q"])
    assert aiken[0]["content"] == "Hỏi gì?" and aiken[0]["options"] == ["x", "y"] and aiken[0]["clo"] == "CLO2"
    assert aiken[1]["content"] == "Câu hai\nnhiều dòng" and aiken[1]["errors"]


def test_bulk_import_then_tag(client, gv):
    cid = _course(client, gv)
    before = len(client.get(f"/api/questions?course_id={cid}", headers=gv).json())
    data = _xlsx([["Nội dung", "A", "B", "C", "Đáp án"],
                  ["Câu nhập hàng loạt số một?", "x", "y", "z", "b"],
                  ["Câu nhập hàng loạt số hai?", "x", "y", "z", "3"],
                  ["Câu nhập hàng loạt số hai?", "x", "y", "z", "C"],   # trùng trong file
                  ["Câu lỗi", "x", "", "", ""]])
    url = f"/api/questions/import?course_id={cid}"
    prev = _up(client, gv, url + "&dry_run=true", "q.xlsx", data).json()
    assert prev["summary"] == {"format": "Excel/CSV", "total": 4, "ok": 2, "duplicate": 1, "error": 1, "will_add": 2, "untagged": 2}
    assert len(client.get(f"/api/questions?course_id={cid}", headers=gv).json()) == before     # xem trước không ghi
    res = _up(client, gv, url + "&dry_run=false", "q.xlsx", data).json()
    ids = res["created_ids"]
    assert len(ids) == 2
    q = client.get(f"/api/questions/{ids[1]}", headers=gv).json()
    assert q["tagged"] is False and q["bloom_level_id"] is None and [o["is_correct"] for o in q["options"]] == [False, False, True]
    untagged = {x["id"] for x in client.get(f"/api/questions?course_id={cid}&untagged=true", headers=gv).json()}
    assert set(ids) <= untagged
    # nhập lại → trùng với câu đã có
    assert _up(client, gv, url + "&dry_run=true", "q.xlsx", data).json()["summary"]["duplicate"] == 3

    # câu chưa gán không được đưa vào đề
    cs = client.get("/api/class-sections", headers=gv).json()[0]["id"]
    r = client.post("/api/exams", headers=gv, json={"class_section_id": cs, "exam_title": "Đề thử", "exam_type": "online",
                                                    "assessment_type": "process", "max_score": 10, "items": [{"question_id": ids[0], "points": 1}]})
    assert r.status_code == 422 and "chưa gán" in r.json()["detail"]

    # tải file gán CLO, điền, tải lên
    t = client.get(f"/api/questions/tagging-template?course_id={cid}&ids={','.join(map(str, ids))}", headers=gv)
    wb = load_workbook(io.BytesIO(t.content)); ws = wb["GanCLO"]
    assert [c.value for c in ws[1]][:6] == ["ID", "Nội dung câu hỏi", "Đáp án", "Chương", "Bloom", "CLO"]
    assert {ws.cell(row=i, column=1).value for i in (2, 3)} == set(ids)
    for i in (2, 3):
        ws.cell(row=i, column=4, value=2); ws.cell(row=i, column=5, value="3 - Vận dụng")
        ws.cell(row=i, column=6, value="CLO2" if i == 2 else "CLO1:0.7; CLO2:0.3")
    ws.append([999999, "không tồn tại"])
    b = io.BytesIO(); wb.save(b)
    turl = f"/api/questions/tagging-import?course_id={cid}"
    p = _up(client, gv, turl + "&dry_run=true", "t.xlsx", b.getvalue()).json()
    assert p["summary"]["update"] == 2 and p["summary"]["error"] == 1
    a = _up(client, gv, turl + "&dry_run=false", "t.xlsx", b.getvalue()).json()
    assert a["summary"]["update"] == 2
    q = client.get(f"/api/questions/{ids[1]}", headers=gv).json()
    assert q["tagged"] and q["bloom_level_id"] == 3 and q["chapter"] == "Chương 2"
    assert sorted((c["clo_code"], c["weight"]) for c in q["clos"]) == [("CLO1", 0.7), ("CLO2", 0.3)]
    # tải lại cùng file → không đổi
    assert _up(client, gv, turl + "&dry_run=true", "t.xlsx", b.getvalue()).json()["summary"]["unchanged"] == 2
    # giờ đưa vào đề được
    r = client.post("/api/exams", headers=gv, json={"class_section_id": cs, "exam_title": "Đề thử", "exam_type": "online",
                                                    "assessment_type": "process", "max_score": 10, "items": [{"question_id": ids[0], "points": 1}]})
    assert r.status_code == 201
    client.delete(f"/api/exams/{r.json()['id']}", headers=gv)
    for i in ids:
        client.delete(f"/api/questions/{i}", headers=gv)


def test_tagging_locks_used_questions(client, gv):
    cid = _course(client, gv)
    used = next(q for q in client.get(f"/api/questions?course_id={cid}", headers=gv).json() if q["usage"]["n_results"])
    other = "CLO1" if used["clos"][0]["clo_code"] != "CLO1" else "CLO2"
    data = _xlsx([["ID", "CLO"], [used["id"], other]])
    p = _up(client, gv, f"/api/questions/tagging-import?course_id={cid}&dry_run=false", "t.xlsx", data).json()
    assert p["rows"][0]["status"] == "locked"
    assert client.get(f"/api/questions/{used['id']}", headers=gv).json()["clos"] == used["clos"]


def test_aiken_and_moodle_xml(client, gv):
    cid = _course(client, gv)
    aiken = "Câu A1 kiểm thử?\nA. một\nB. hai\nANSWER: A\nBLOOM: Hiểu\nCLO: CLO1\n\nCâu A2 kiểm thử?\nA. một\nB. hai\nANSWER: B\n"
    p = _up(client, gv, f"/api/questions/import?course_id={cid}&dry_run=true", "a.txt", aiken.encode()).json()
    assert p["summary"]["ok"] == 2 and p["summary"]["untagged"] == 1
    xml = ('<quiz><question type="multichoice"><questiontext><text>&lt;p&gt;Câu XML kiểm thử?&lt;/p&gt;</text></questiontext>'
           '<answer fraction="0"><text>sai</text></answer><answer fraction="100"><text>đúng</text></answer>'
           '<tags><tag><text>CLO2</text></tag><tag><text>Bloom:2</text></tag></tags></question>'
           '<question type="essay"><questiontext><text>Tự luận</text></questiontext></question></quiz>')
    p = _up(client, gv, f"/api/questions/import?course_id={cid}&dry_run=true", "m.xml", xml.encode()).json()
    ok = [r for r in p["rows"] if r["status"] == "ok"]
    assert len(ok) == 1 and ok[0]["content"] == "Câu XML kiểm thử?" and ok[0]["answer"] == "B" and ok[0]["tagged"]
    assert p["summary"]["error"] == 1
    assert _up(client, gv, f"/api/questions/import?course_id={cid}", "x.pdf", b"x").status_code == 422
