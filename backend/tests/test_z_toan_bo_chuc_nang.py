"""Kiểm thử phủ TOÀN BỘ chức năng còn lại của API (chạy sau cùng – tên tệp bắt đầu bằng test_z).

Cùng với các tệp kiểm thử khác, mọi endpoint của hệ thống đều được gọi ít nhất một lần với kết quả thành công
(xem tools/kiem_thu/bao_phu_api.py). Các thao tác ghi sang Moodle qua Web Service được thay bằng một "Moodle giả"
(FakeMoodle) ghi lại lời gọi và trả về đúng cấu trúc của plugin local_clo; phần ĐỌC Moodle dùng CSDL Moodle mô phỏng
(moodle_sim) như các kiểm thử khác. Kiểm thử với Moodle thật: tools/moodle_e2e/.
"""
import io
import zipfile

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from app.config import settings
from app.database import engine
from app.services import lms

M = settings.moodle_db_name


def q(sql, **p):
    with engine.connect() as c:
        return c.execute(text(sql), p).fetchall()


def ex(sql, **p):
    with engine.begin() as c:
        return c.execute(text(sql), p)


# ------------------------------------------------------------------ Moodle giả cho kênh Web Service
class FakeMoodle:
    def __init__(self):
        self.calls = []
        self.next_q = 7000

    def __call__(self, fn, params=None, timeout=60):
        params = params or {}
        self.calls.append((fn, params))
        base = settings.moodle_url.rstrip("/")
        if fn == "core_webservice_get_site_info":
            return {"sitename": "UTE LMS (giả lập)", "username": "clo_service", "release": "4.5.14",
                    "functions": [{"name": n} for n in ("core_webservice_get_site_info", "local_clo_create_quiz",
                                                         "local_clo_create_offlinequiz", "local_clo_get_offlinequiz",
                                                         "local_clo_process_scans", "local_clo_push_results",
                                                         "local_clo_sync_questions")]}
        if fn == "local_clo_create_quiz":
            assert "<quiz>" in params["questionsxml"] and params["courseid"]
            return {"quizid": 9800 + params["examref"], "cmid": 8800 + params["examref"], "courseid": params["courseid"],
                    "created": True, "nquestions": params["questionsxml"].count("<question type=\"multichoice\"")}
        if fn in ("local_clo_create_offlinequiz", "local_clo_get_offlinequiz", "local_clo_process_scans"):
            ref = params["examref"]
            files = lambda g: [{"filename": f"{k}_{g}.pdf", "kind": k, "url": f"{base}/webservice/pluginfile.php/1/mod_offlinequiz/{k}_{g}.pdf"}
                               for k in ("question", "answer", "correction")]
            return {"offlinequizid": 9900 + ref, "cmid": 8900 + ref, "docscreated": True, "status": "ready",
                    "npending": 0, "nerror": 0, "ngraded": 38 if fn == "local_clo_process_scans" else 0,
                    "groups": [{"letter": g, "number": i + 1, "files": files(g)} for i, g in enumerate("AB")],
                    "uploadurl": f"{base}/mod/offlinequiz/upload.php", "correcturl": f"{base}/mod/offlinequiz/correct.php"}
        if fn == "local_clo_push_results":
            return {"examid": 1, "saved": len(params["results"]), "notfound": 0, "notified": 0}
        if fn == "local_clo_sync_questions":
            out = []
            for item in params["questions"]:
                mid = item.get("moodlequestionid") or self._new_id()
                out.append({"idnumber": item["idnumber"], "questionid": mid, "version": 1, "action": "tagged" if item.get("moodlequestionid") else "created"})
            return {"results": out}
        raise AssertionError(f"Hàm Web Service không mong đợi: {fn}")

    def _new_id(self):
        self.next_q += 1
        return self.next_q


@pytest.fixture()
def moodle_ws(monkeypatch):
    fake = FakeMoodle()
    monkeypatch.setattr(settings, "moodle_url", "http://moodle.test")
    monkeypatch.setattr(settings, "moodle_ws_token", "token-gia-lap")
    monkeypatch.setattr(lms, "call", fake)
    monkeypatch.setattr(lms, "download", lambda url: (b"%PDF-1.4 tep gia lap " + url.encode(), "application/pdf"))
    return fake


@pytest.fixture()
def gv_teacher_of_course_50():
    """gv.son là giảng viên (editingteacher) của khóa học Moodle 50 trong CSDL Moodle mô phỏng."""
    ex(f"INSERT INTO `{M}`.mdl_role_assignments (id, roleid, contextid, userid) VALUES (9101, 3, 70, 99)")
    yield
    ex(f"DELETE FROM `{M}`.mdl_role_assignments WHERE id=9101")


# ================================================================== 1. Xác thực, danh mục
def test_auth_me_and_catalogs(client, gv, admin):
    me = client.get("/api/auth/me", headers=gv).json()
    assert me["username"] == "gv.son" and me["role"] == "lecturer"
    assert client.get("/api/auth/me").status_code == 401
    assert client.get("/api/auth/me", headers=admin).json()["role"] == "admin"
    sem = client.get("/api/semesters", headers=gv).json()
    assert sem and {"id", "name", "academic_year"} <= set(sem[0])
    bloom = client.get("/api/bloom-levels", headers=gv).json()
    assert [b["id"] for b in bloom] == [1, 2, 3, 4, 5, 6]
    lec = client.get("/api/lecturers", headers=gv).json()
    assert any(x["full_name"] == "Nguyễn Thành Sơn" for x in lec)
    courses = client.get("/api/courses", headers=gv).json()
    assert any(c["course_code"] == "DBMS330284" for c in courses)
    assert client.get("/api/courses?program_id=2", headers=admin).status_code == 200
    assert client.get("/api/health").json()["status"] == "ok"
    assert "moodle_url" in client.get("/api/public-config").json()


# ================================================================== 2. CTĐT, PLO, kế hoạch đo PI (Bộ môn)
def test_programs_plo_and_pi_plans(client, gv, admin):
    body = {"code": "TEST-CTDT", "name": "CTĐT kiểm thử", "faculty": "Khoa CNTT", "cohort_year": 2026, "total_credits": 150}
    assert client.post("/api/programs", headers=gv, json=body).status_code == 403
    r = client.post("/api/programs", headers=admin, json=body)
    assert r.status_code == 201, r.text
    pid = r.json()["id"]
    assert client.post("/api/programs", headers=admin, json=body).status_code == 409             # trùng mã
    assert client.put(f"/api/programs/{pid}/courses", headers=admin, json={"course_id": 5, "semester_no": 3}).json()["ok"]
    assert any(c["id"] == 5 and c["semester_no"] == 3 for c in client.get(f"/api/programs/{pid}", headers=admin).json()["courses"])
    assert client.delete(f"/api/programs/{pid}/courses/5", headers=gv).status_code == 403
    assert client.delete(f"/api/programs/{pid}/courses/5", headers=admin).json() == {"deleted": 5}
    assert client.delete(f"/api/programs/{pid}/courses/5", headers=admin).status_code == 404
    # PLO: sửa chỉ tiêu -> đánh giá lại, rồi trả lại như cũ
    plo = q("SELECT id, target_pct, description FROM plos WHERE program_id=1 ORDER BY id LIMIT 1")[0]
    assert client.put(f"/api/plos/{plo[0]}", headers=gv, json={"target_pct": 50}).status_code == 403
    assert client.put(f"/api/plos/{plo[0]}", headers=admin, json={"target_pct": 50}).json()["ok"]
    assert float(q("SELECT target_pct FROM plos WHERE id=:i", i=plo[0])[0][0]) == 50
    client.put(f"/api/plos/{plo[0]}", headers=admin, json={"target_pct": float(plo[1])})
    assert client.put("/api/plos/999999", headers=admin, json={"target_pct": 50}).status_code == 404
    # Kế hoạch đo PI: thêm (CLO phải thuộc môn), sửa, xóa
    pi = q("SELECT pi.id FROM performance_indicators pi JOIN plos p ON p.id=pi.plo_id WHERE p.program_id=1 ORDER BY pi.id LIMIT 1")[0][0]
    sem = q("SELECT id FROM semesters ORDER BY id LIMIT 1")[0][0]
    plan = {"pi_id": pi, "course_id": 5, "semester_id": sem, "method": "Bài kiểm tra trắc nghiệm", "target_pct": 70, "clo_ids": [1, 2]}
    assert client.post("/api/pi-plans", headers=admin, json={**plan, "clo_ids": [999]}).status_code == 422
    r = client.post("/api/pi-plans", headers=admin, json=plan)
    assert r.status_code == 200, r.text
    plan_id = r.json()["id"]
    assert client.put(f"/api/pi-plans/{plan_id}", headers=admin, json={**plan, "method": "Tự luận", "clo_ids": [1]}).json()["ok"]
    assert q("SELECT method FROM pi_assessment_plans WHERE id=:i", i=plan_id)[0][0] == "Tự luận"
    assert any(p["id"] == plan_id for p in client.get("/api/pi-plans?program_id=1", headers=gv).json())
    assert client.delete(f"/api/pi-plans/{plan_id}", headers=admin).json()["ok"]
    assert client.delete(f"/api/pi-plans/{plan_id}", headers=admin).status_code == 404
    assert client.post("/api/reports/programs/1/recompute", headers=gv).status_code == 403
    assert client.post("/api/reports/programs/1/recompute", headers=admin).status_code == 200


# ================================================================== 3. CLO (Bộ môn), câu hỏi, tệp mẫu
def test_clo_lifecycle_cancel_question_and_template(client, gv, admin):
    body = {"course_id": 5, "clo_code": "CLO9", "description": "CLO tạm cho kiểm thử", "bloom_level_id": 2, "plos": []}
    assert client.post("/api/clos", headers=gv, json=body).status_code == 403                    # GV không tạo CLO
    r = client.post("/api/clos", headers=admin, json=body)
    assert r.status_code in (200, 201), r.text
    cid = r.json()["id"]
    assert client.delete(f"/api/clos/{cid}", headers=gv).status_code == 403
    assert client.delete(f"/api/clos/{cid}", headers=admin).json() == {"deleted": cid}
    assert client.delete("/api/clos/1", headers=admin).status_code == 409                       # CLO đã dùng
    # Hủy câu hỏi (có lý do) – câu đã Hủy không vào đề mới
    qb = {"course_id": 5, "outline_id": None, "bloom_level_id": 1, "content": "Câu hỏi tạm để Hủy?",
          "options": [{"content": "Đúng", "is_correct": True}, {"content": "Sai", "is_correct": False}],
          "clos": [{"clo_id": 1, "weight": 1.0}]}
    r = client.post("/api/questions", headers=gv, json=qb)
    assert r.status_code == 201, r.text
    qid = r.json()["id"]
    assert client.patch(f"/api/questions/{qid}/cancel", headers=gv, json={"reason": "ab"}).status_code == 422
    assert client.patch(f"/api/questions/{qid}/cancel", headers=gv, json={"reason": "Nội dung trùng câu khác"}).json()["is_cancelled"]
    assert client.get(f"/api/questions/{qid}", headers=gv).json()["is_cancelled"]
    assert client.delete(f"/api/questions/{qid}", headers=gv).status_code == 200
    # Tệp Excel mẫu nhập câu hỏi
    t = client.get("/api/questions/import-template?course_id=5", headers=gv)
    assert t.status_code == 200
    ws = load_workbook(io.BytesIO(t.content)).active
    assert [c.value for c in ws[1]][:7] == ["Nội dung", "A", "B", "C", "D", "E", "Đáp án"]


# ================================================================== 4. Đồng bộ & phân tích, cầu nối
def test_sync_analyze_runs_and_bridge(client, gv, admin):
    r = client.post("/api/exams/2/analyze", headers=gv)
    assert r.status_code == 200, r.text
    runs = client.get("/api/exams/2/sync-runs", headers=gv).json()
    assert runs and runs[0]["status"] == "success"
    st = client.get("/api/system/moodle-status", headers=gv).json()
    assert st["moodle_available"] and st["bridge_installed"]
    assert client.post("/api/system/install-bridge", headers=gv).status_code == 403
    r = client.post("/api/system/install-bridge", headers=admin)
    assert r.status_code == 200 and r.json()["statements"] > 0, r.text
    ev = client.get("/api/class-sections/1/clo-results/1/evidence", headers=gv).json()
    assert {r["exam_id"] for r in ev["evidence"]} == {1, 2}            # CLO1: minh chứng từ cả hai bài (BM6d)


# ================================================================== 5. Liên kết lớp HP ⇄ khóa học Moodle
def test_link_section_with_moodle_both_ways(client, gv, admin):
    # khóa học Moodle 51 có shortname trùng mã lớp DBMS330284_02, GV là giảng viên của khóa
    ex(f"INSERT INTO `{M}`.mdl_course (id, category, fullname, shortname, idnumber, visible) VALUES (51, 1, 'HQT CSDL lớp 02', 'DBMS330284_02', '', 1)")
    ex(f"INSERT INTO `{M}`.mdl_context (id, contextlevel, instanceid, path, depth) VALUES (71, 50, 51, '/1/71', 2)")
    try:
        # chưa là giảng viên của khóa 51 -> không tự nối
        r = client.post("/api/class-sections/2/moodle-autolink", headers=gv)
        assert r.status_code == 404 and "không phải giảng viên" in r.json()["detail"]
        ex(f"INSERT INTO `{M}`.mdl_role_assignments (id, roleid, contextid, userid) VALUES (9102, 3, 71, 99)")
        r = client.post("/api/class-sections/2/moodle-autolink", headers=gv)
        assert r.status_code == 200 and r.json()["moodle_course_id"] == 51 and r.json()["status"] == "vừa gắn theo mã lớp"
        # cả môn: lớp 01 đã gắn, lớp 02 vừa gắn
        secs = client.post("/api/courses/5/moodle-sync-sections", headers=gv).json()["sections"]
        assert {s["section_code"]: s["moodle_course_id"] for s in secs} == {"DBMS330284_01": 50, "DBMS330284_02": 51}
        # gắn thủ công: khóa đã gắn lớp khác -> 409; gắn lại đúng khóa -> 200
        assert client.patch("/api/class-sections/2/moodle-course", headers=gv, json={"moodle_course_id": 50}).status_code == 409
        r = client.patch("/api/class-sections/2/moodle-course", headers=gv, json={"moodle_course_id": 51})
        assert r.status_code == 200 and r.json()["moodle_course_id"] == 51
        assert client.patch("/api/class-sections/2/moodle-course", headers=gv, json={"moodle_course_id": 4040}).status_code == 404
    finally:
        ex("UPDATE class_sections SET moodle_course_id=NULL WHERE id=2")
        ex(f"DELETE FROM `{M}`.mdl_role_assignments WHERE id=9102")
        ex(f"DELETE FROM `{M}`.mdl_context WHERE id=71")
        ex(f"DELETE FROM `{M}`.mdl_course WHERE id=51")


# ================================================================== 6. Web Service: Quiz, Offline Quiz, công bố, trạng thái
def test_moodle_web_service_flows(client, gv, admin, moodle_ws, gv_teacher_of_course_50):
    st = client.get("/api/system/lms-status", headers=gv).json()
    assert st["ws_configured"] and st["ws_ok"] and "local_clo_sync_questions" in st["site"]["functions"]
    qs = [x["id"] for x in client.get("/api/questions?course_id=5", headers=gv).json() if x["tagged"]][:5]
    items = [{"question_id": i, "points": 2} for i in qs]
    # Tạo Quiz trên Moodle
    on = client.post("/api/exams", headers=gv, json={"class_section_id": 1, "exam_title": "Quiz tạo qua Web Service", "exam_type": "online",
                                                      "assessment_type": "process", "max_score": 10, "items": items}).json()["id"]
    r = client.post(f"/api/exams/{on}/moodle/create-quiz", headers=gv, json={})
    assert r.status_code == 200, r.text
    assert r.json()["quizid"] == 9800 + on and "/mod/quiz/view.php?id=" in r.json()["quiz_url"]
    e = client.get(f"/api/exams/{on}", headers=gv).json()
    assert e["moodle_quiz_id"] == 9800 + on and e["status"] == "Published"
    assert client.post(f"/api/exams/{on}/moodle/create-quiz", headers=gv, json={}).status_code == 409          # đã có Quiz
    # Tạo Offline Quiz (đề giấy trên Moodle), xem trạng thái, xử lý phiếu, tải tệp
    pp = client.post("/api/exams", headers=gv, json={"class_section_id": 1, "exam_title": "Đề giấy Moodle Offline Quiz", "exam_type": "paper",
                                                      "assessment_type": "process", "max_score": 10, "items": items}).json()["id"]
    assert client.post(f"/api/exams/{pp}/moodle/create-offlinequiz", headers=gv, json={"numgroups": 9}).status_code == 422
    r = client.post(f"/api/exams/{pp}/moodle/create-offlinequiz", headers=gv, json={"numgroups": 2, "intro": "Không dùng tài liệu"})
    assert r.status_code == 200, r.text
    oq = r.json()
    assert oq["offlinequizid"] == 9900 + pp and all(f["download"].startswith(f"/api/exams/{pp}/") and "url" not in f
                                                  for g in oq["groups"] for f in g["files"])          # không lộ token/URL Moodle
    assert client.get(f"/api/exams/{pp}/moodle/offlinequiz", headers=gv).json()["groups"][1]["letter"] == "B"
    assert client.post(f"/api/exams/{pp}/moodle/process-scans", headers=gv).json()["ngraded"] == 38
    f = client.get(f"/api/exams/{pp}/moodle/offlinequiz/files/answer_A.pdf", headers=gv)
    assert f.status_code == 200 and f.content.startswith(b"%PDF") and "phieu_tra_loi_ma_de_A" in f.headers["content-disposition"]
    assert client.get(f"/api/exams/{pp}/moodle/offlinequiz/files/khong_co.pdf", headers=gv).status_code == 404
    z = zipfile.ZipFile(io.BytesIO(client.get(f"/api/exams/{pp}/moodle/offlinequiz/files.zip", headers=gv).content))
    assert {"ma_de_A/de_thi_A.pdf", "ma_de_B/dap_an_B.pdf"} <= set(z.namelist())
    assert client.get(f"/api/exams/{on}/moodle/offlinequiz", headers=gv).status_code == 404                  # bài trực tuyến
    # Công bố / gửi kết quả sang Moodle (bài cuối kỳ đã phân tích)
    r = client.patch("/api/exams/2/publish", headers=gv, json={"publish": True})
    assert r.status_code == 200 and r.json()["publish_flag"] and r.json()["moodle"]["saved"] == 40, r.text
    pushed = [p for fn, p in moodle_ws.calls if fn == "local_clo_push_results"][-1]
    assert pushed["published"] is True and len(pushed["results"]) == 40
    r = client.post("/api/exams/2/moodle/push-results", headers=gv)
    assert r.status_code == 200 and r.json()["sent"] == 40
    assert client.post(f"/api/exams/{on}/moodle/push-results", headers=gv).status_code == 409                # chưa phân tích
    # dọn
    for eid in (on, pp):
        ex("DELETE FROM exams WHERE id=:e", e=eid)


# ================================================================== 7. Dùng đề cho lớp khác, sửa câu/ý bài nhập điểm, xóa mã đề
def test_copy_exam_manual_items_and_paper_versions(client, gv, gv_other):
    r = client.post("/api/exams/2/copy", headers=gv, json={"class_section_ids": [2]})
    assert r.status_code == 200, r.text
    cp = r.json()["copies"][0]
    assert cp["section_code"] == "DBMS330284_02" and "chưa gắn khóa học Moodle" in cp["moodle"]
    n_src = q("SELECT COUNT(*) FROM exam_questions WHERE exam_id=2")[0][0]
    assert q("SELECT COUNT(*) FROM exam_questions WHERE exam_id=:e", e=cp["exam_id"])[0][0] == n_src
    assert client.post("/api/exams/2/copy", headers=gv_other, json={"class_section_ids": [2]}).status_code == 403
    ex("DELETE FROM exams WHERE id=:e", e=cp["exam_id"])
    # bài nhập điểm: đổi cấu trúc câu/ý khi chưa có điểm
    body = {"exam_title": "Thực hành tuần 3", "assessment_type": "process", "max_score": 10,
            "items": [{"title": "Ý 1", "points": 5, "clos": [{"clo_id": 1, "weight": 1}]}]}
    me = client.post("/api/class-sections/1/manual-exams", headers=gv, json=body).json()["id"]
    new_items = {"items": [{"title": "Ý 1", "points": 4, "clos": [{"clo_id": 1, "weight": 1}]},
                           {"title": "Ý 2", "points": 6, "clos": [{"clo_id": 2, "weight": 1}]}]}
    r = client.put(f"/api/exams/{me}/manual-items", headers=gv, json=new_items)
    assert r.status_code == 200 and r.json()["n_items"] == 2, r.text
    assert [i["points"] for i in client.get(f"/api/exams/{me}/grades", headers=gv).json()["items"]] == [4, 6]
    ex("DELETE FROM exams WHERE id=:e", e=me)
    # đề giấy chấm trên hệ thống: xóa mã đề khi chưa chấm (để sinh lại)
    qs = [x["id"] for x in client.get("/api/questions?course_id=5", headers=gv).json() if x["tagged"]][:4]
    pe = client.post("/api/exams", headers=gv, json={"class_section_id": 1, "exam_title": "Giấy – sinh lại mã đề", "exam_type": "paper",
                                                      "items": [{"question_id": i, "points": 1} for i in qs]}).json()["id"]
    assert client.post(f"/api/exams/{pe}/paper/versions", headers=gv, json={"n_versions": 2}).status_code == 200
    r = client.delete(f"/api/exams/{pe}/paper/versions", headers=gv)
    assert r.status_code == 200, r.text
    assert client.get(f"/api/exams/{pe}/paper", headers=gv).json()["mode"] is None
    ex("DELETE FROM exams WHERE id=:e", e=pe)


# ================================================================== 8. Đồng bộ ngân hàng câu hỏi hai chiều
def test_question_bank_two_way_sync(client, gv, gv_other, moodle_ws):
    t = client.get("/api/courses/5/moodle-sync/targets", headers=gv).json()
    assert [x["moodle_course_id"] for x in t["targets"]] == [50] and t["ws_configured"]
    assert client.get("/api/courses/5/moodle-sync?moodle_course_id=4040", headers=gv).status_code == 403
    before = client.get("/api/courses/5/moodle-sync?moodle_course_id=50", headers=gv).json()
    assert before["n_moodle"] == 30
    # (1) câu soạn trực tiếp trên Moodle, chưa có idnumber -> "chỉ có trên Moodle"
    ex(f"""INSERT INTO `{M}`.mdl_question (id, parent, name, questiontext, generalfeedback, defaultmark, qtype, timecreated, timemodified)
           VALUES (1100, 0, 'Soạn trên Moodle', '<p>Lệnh nào xóa toàn bộ dữ liệu của bảng nhưng giữ cấu trúc?</p>', '', 1, 'multichoice', 1, 1)""")
    ex(f"INSERT INTO `{M}`.mdl_question_bank_entries (id, questioncategoryid, idnumber) VALUES (400, 80, NULL)")
    ex(f"INSERT INTO `{M}`.mdl_question_versions (id, questionbankentryid, version, questionid, status) VALUES (100, 400, 1, 1100, 'ready')")
    for i, (a, f) in enumerate([("TRUNCATE TABLE", 1), ("DROP TABLE", 0), ("DELETE DATABASE", 0)]):
        ex(f"INSERT INTO `{M}`.mdl_question_answers (id, question, answer, fraction, feedback) VALUES (:i, 1100, :a, :f, '')", i=6001 + i, a=a, f=f)
    ex(f"""INSERT INTO `{M}`.mdl_qtype_multichoice_options (id, questionid, single, correctfeedback, partiallycorrectfeedback, incorrectfeedback)
           VALUES (900, 1100, 1, '', '', '')""")
    # (2) câu mới soạn ở hệ thống -> "chỉ có ở hệ thống"
    new_q = client.post("/api/questions", headers=gv, json={
        "course_id": 5, "bloom_level_id": 1, "content": "Câu soạn ở hệ thống để đẩy lên Moodle?",
        "options": [{"content": "Có", "is_correct": True}, {"content": "Không", "is_correct": False}],
        "clos": [{"clo_id": 1, "weight": 1.0}]}).json()["id"]
    created = None
    try:
        cmp = client.get("/api/courses/5/moodle-sync?moodle_course_id=50", headers=gv).json()
        rows = {(r["state"], r.get("question_id"), r.get("moodle_questionid")) for r in cmp["rows"]}
        assert ("only_moodle", None, 1100) in rows and ("only_system", new_q, None) in rows
        assert cmp["summary"]["synced"] + cmp["summary"].get("moodle_newer", 0) >= 1
        assert client.post("/api/courses/5/moodle-sync", headers=gv_other,
                           json={"moodle_course_id": 50, "push": [new_q]}).status_code == 403
        r = client.post("/api/courses/5/moodle-sync", headers=gv, json={"moodle_course_id": 50, "push": [new_q], "import_new": [1100]})
        assert r.status_code == 200, r.text
        res = r.json()
        assert [p["idnumber"] for p in res["pushed"]] == [f"QB-{new_q}"] and len(res["created"]) == 1 and res["tagged"] == 1
        created = res["created"][0]
        sent = [p for fn, p in moodle_ws.calls if fn == "local_clo_sync_questions"]
        assert sent[0]["questions"][0]["answers"][0] == {"text": "Có", "fraction": 1}
        assert sent[1]["questions"] == [{"idnumber": f"QB-{created}", "moodlequestionid": 1100}]
        imp = client.get(f"/api/questions/{created}", headers=gv).json()
        assert "TRUNCATE" in [o["content"] for o in imp["options"] if o["is_correct"]][0] and imp["tagged"] is False   # chờ gán CLO
        # (3) Moodle gắn QB-n cho câu vừa nhập -> đã khớp; GV sửa câu trên Moodle (phiên bản 2) -> "Moodle mới hơn" -> kéo về
        ex(f"UPDATE `{M}`.mdl_question_bank_entries SET idnumber=:n WHERE id=400", n=f"QB-{created}")
        cmp = client.get("/api/courses/5/moodle-sync?moodle_course_id=50", headers=gv).json()
        assert next(r for r in cmp["rows"] if r.get("question_id") == created)["state"] == "synced"
        ex(f"""INSERT INTO `{M}`.mdl_question (id, parent, name, questiontext, generalfeedback, defaultmark, qtype, timecreated, timemodified)
               VALUES (1101, 0, 'Soạn trên Moodle', 'Lệnh nào xóa NHANH toàn bộ dữ liệu của bảng nhưng giữ cấu trúc?', '', 1, 'multichoice', 2, 2)""")
        ex(f"INSERT INTO `{M}`.mdl_question_versions (id, questionbankentryid, version, questionid, status) VALUES (101, 400, 2, 1101, 'ready')")
        for i, (a, f) in enumerate([("TRUNCATE TABLE", 1), ("DROP TABLE", 0), ("DELETE FROM … (không WHERE)", 0)]):
            ex(f"INSERT INTO `{M}`.mdl_question_answers (id, question, answer, fraction, feedback) VALUES (:i, 1101, :a, :f, '')", i=6011 + i, a=a, f=f)
        ex(f"""INSERT INTO `{M}`.mdl_qtype_multichoice_options (id, questionid, single, correctfeedback, partiallycorrectfeedback, incorrectfeedback)
               VALUES (901, 1101, 1, '', '', '')""")
        cmp = client.get("/api/courses/5/moodle-sync?moodle_course_id=50", headers=gv).json()
        row = next(r for r in cmp["rows"] if r.get("question_id") == created)
        assert row["state"] == "moodle_newer" and row["action"] == "pull" and row["moodle_version"] == 2
        r = client.post("/api/courses/5/moodle-sync", headers=gv, json={"moodle_course_id": 50, "pull": [created]})
        assert r.status_code == 200 and r.json()["updated"] == [created], r.text
        got = client.get(f"/api/questions/{created}", headers=gv).json()
        assert "NHANH" in got["content"] and got["options"][2]["content"].startswith("DELETE FROM")
        cmp = client.get("/api/courses/5/moodle-sync?moodle_course_id=50", headers=gv).json()
        assert next(r for r in cmp["rows"] if r.get("question_id") == created)["state"] == "synced"
        # (4) góp ý F3: sửa ở CẢ hai bên sau lần đồng bộ -> "Hai bên khác nhau", không tự chọn chiều (trước đây báo
        #     "Moodle mới hơn" và kéo về sẽ ghi đè mất bản sửa ở hệ thống)
        ex("UPDATE question_bank SET content='Lệnh nào xóa nhanh dữ liệu bảng (sửa ở hệ thống)?' WHERE id=:q", q=created)
        ex(f"""INSERT INTO `{M}`.mdl_question (id, parent, name, questiontext, generalfeedback, defaultmark, qtype, timecreated, timemodified)
               VALUES (1102, 0, 'Soạn trên Moodle', 'Lệnh nào xóa dữ liệu bảng (sửa trên Moodle)?', '', 1, 'multichoice', 3, 3)""")
        ex(f"INSERT INTO `{M}`.mdl_question_versions (id, questionbankentryid, version, questionid, status) VALUES (102, 400, 3, 1102, 'ready')")
        for i, (a, f) in enumerate([("TRUNCATE TABLE", 1), ("DROP TABLE", 0), ("DELETE FROM … (không WHERE)", 0)]):
            ex(f"INSERT INTO `{M}`.mdl_question_answers (id, question, answer, fraction, feedback) VALUES (:i, 1102, :a, :f, '')", i=6021 + i, a=a, f=f)
        ex(f"""INSERT INTO `{M}`.mdl_qtype_multichoice_options (id, questionid, single, correctfeedback, partiallycorrectfeedback, incorrectfeedback)
               VALUES (902, 1102, 1, '', '', '')""")
        cmp = client.get("/api/courses/5/moodle-sync?moodle_course_id=50", headers=gv).json()
        row = next(r for r in cmp["rows"] if r.get("question_id") == created)
        assert row["state"] == "conflict" and row["action"] is None and "Cả hệ thống và Moodle" in row["note"]
        # GV chọn lấy bản Moodle -> hai bên khớp lại
        r = client.post("/api/courses/5/moodle-sync", headers=gv, json={"moodle_course_id": 50, "pull": [created]})
        assert r.status_code == 200, r.text
        cmp = client.get("/api/courses/5/moodle-sync?moodle_course_id=50", headers=gv).json()
        assert next(r for r in cmp["rows"] if r.get("question_id") == created)["state"] == "synced"
        # chỉ hệ thống sửa -> "Hệ thống mới hơn" (đề xuất đẩy lên)
        ex("UPDATE question_bank SET content='Chỉ sửa ở hệ thống?' WHERE id=:q", q=created)
        cmp = client.get("/api/courses/5/moodle-sync?moodle_course_id=50", headers=gv).json()
        row = next(r for r in cmp["rows"] if r.get("question_id") == created)
        assert row["state"] == "system_newer" and row["action"] == "push"
        # câu đã có kết quả thi không bị ghi đè từ Moodle
        locked = [r for r in cmp["rows"] if r["state"] == "moodle_newer" and r["action"] is None]
        assert all("đã có kết quả thi" in r["note"] for r in locked)
    finally:
        ex(f"DELETE FROM `{M}`.mdl_qtype_multichoice_options WHERE id IN (900, 901, 902)")
        ex(f"DELETE FROM `{M}`.mdl_question_answers WHERE question IN (1100, 1101, 1102)")
        ex(f"DELETE FROM `{M}`.mdl_question_versions WHERE id IN (100, 101, 102)")
        ex(f"DELETE FROM `{M}`.mdl_question_bank_entries WHERE id=400")
        ex(f"DELETE FROM `{M}`.mdl_question WHERE id IN (1100, 1101, 1102)")
        for i in (new_q, created):
            if i:
                ex("DELETE FROM question_bank WHERE id=:i", i=i)
