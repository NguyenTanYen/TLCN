"""Nhập điểm thủ công / Excel (bài làm ngoài hệ thống) và gộp nhiều bài KT để đánh giá CLO, PLO (góp ý 15/10)."""
import io

from openpyxl import Workbook, load_workbook
from sqlalchemy import text

import pytest

from app.database import SessionLocal, engine
from app.services import results


@pytest.fixture(scope="module", autouse=True)
def _cleanup():
    """Các bài nhập điểm tạo trong tệp này bị xóa sau cùng để không ảnh hưởng kiểm thử khác (thống kê lớp, BM6…)."""
    yield
    with engine.begin() as c:
        c.execute(text("DELETE a FROM exam_attempts a JOIN exams e ON e.id=a.exam_id WHERE e.exam_type='manual' AND e.class_section_id=1"))
        c.execute(text("DELETE FROM exams WHERE exam_type='manual' AND class_section_id=1"))
        c.execute(text("DELETE e FROM enrollments e JOIN students s ON s.id=e.student_id WHERE s.student_code='99999901'"))
    db = SessionLocal()
    results.recompute_class_section(db, 1)
    db.close()


def _q(sql, **p):
    with engine.connect() as c:
        return c.execute(text(sql), p).fetchall()


def test_manual_exam_grid_excel_and_analysis(client, gv, gv_other):
    items = [{"title": "Câu 1 – Thiết kế lược đồ", "points": 4, "clos": [{"clo_id": 1, "weight": 1}]},
             {"title": "Câu 2 – Viết truy vấn", "points": 6, "clos": [{"clo_id": 2, "weight": 0.5}, {"clo_id": 3, "weight": 0.5}]}]
    body = {"exam_title": "Bài tập lớn (tự luận)", "assessment_type": "process", "max_score": 10, "items": items}
    assert client.post("/api/class-sections/1/manual-exams", headers=gv_other, json=body).status_code == 403
    bad = {**body, "items": [{**items[0], "clos": [{"clo_id": 1, "weight": 0.5}]}]}
    assert client.post("/api/class-sections/1/manual-exams", headers=gv, json=bad).status_code == 422      # tổng trọng số ≠ 1
    r = client.post("/api/class-sections/1/manual-exams", headers=gv, json=body)
    assert r.status_code == 201, r.text
    eid = r.json()["id"]
    gb = client.get(f"/api/exams/{eid}/grades", headers=gv).json()
    assert len(gb["items"]) == 2 and gb["raw_max"] == 10 and len(gb["students"]) == 40
    q1, q2 = (str(i["question_id"]) for i in gb["items"])
    studs = gb["students"]
    # câu/ý chấm tay không xuất hiện trong ngân hàng câu hỏi trắc nghiệm
    assert not any(q["id"] == int(q1) for q in client.get("/api/questions?course_id=5", headers=gv).json())
    # lưới: điểm vượt tối đa bị chặn
    assert client.put(f"/api/exams/{eid}/grades", headers=gv, json={"rows": [{"student_id": studs[0]["student_id"], "scores": {q1: 5}}]}).status_code == 422
    rows = [{"student_id": s["student_id"], "absent": i == 0, "scores": {q1: (i % 5), q2: (i % 7)}} for i, s in enumerate(studs[:20])]
    r = client.put(f"/api/exams/{eid}/grades", headers=gv, json={"rows": rows})
    assert r.status_code == 200, r.text
    assert r.json()["saved"] == 19 and r.json()["absent"] == 1 and r.json()["analysis"]["n_students"] == 19
    e = client.get(f"/api/exams/{eid}", headers=gv).json()
    assert e["status"] == "Analyzed" and e["exam_type"] == "manual"
    # CLO tính theo lớp (không theo từng SV): danh sách SV không còn trả CLO riêng
    st = {s["student_code"]: s for s in client.get(f"/api/exams/{eid}/students", headers=gv).json()}
    assert "clos" not in st[studs[1]["student_code"]]
    cs = client.get(f"/api/exams/{eid}/clo-summary", headers=gv).json()
    assert cs and all({"clo_code", "n_evaluated", "n_achieved"} <= set(x) for x in cs)
    ov = client.get(f"/api/exams/{eid}/overview", headers=gv).json()
    assert ov["reliability"] == "Cronbach alpha" and ov["n_students"] == 19 and ov["n_absent"] == 1
    items_st = client.get(f"/api/exams/{eid}/item-stats", headers=gv).json()
    assert all(0 <= float(x["p_value"]) <= 1 for x in items_st)
    # Excel: tải bảng điểm mẫu, sửa và tải lên
    x = client.get(f"/api/exams/{eid}/grades.xlsx", headers=gv)
    wb = load_workbook(io.BytesIO(x.content))
    ws = wb["NhapDiem"]
    assert ws["B3"].value == "MSSV" and ws.cell(6, 5).value == f"QB-{q1}"
    ws.cell(7 + 25, 5).value = 3.5; ws.cell(7 + 25, 6).value = "4,5"   # SV thứ 26 (chưa có điểm)
    ws.cell(7 + 26, 4).value = "x"                                       # SV thứ 27 vắng
    buf = io.BytesIO(); wb.save(buf)
    files = {"file": ("bd.xlsx", buf.getvalue(), "application/octet-stream")}
    pre = client.post(f"/api/exams/{eid}/grades/import?dry_run=true", headers=gv, files=files).json()
    assert pre["n_matched"] == 40 and not pre["missing_students"] and not pre["errors"]
    r = client.post(f"/api/exams/{eid}/grades/import?dry_run=false", headers=gv, files=files)
    assert r.status_code == 200, r.text
    gb = client.get(f"/api/exams/{eid}/grades", headers=gv).json()
    s26 = gb["students"][25]
    assert s26["status"] == "finished" and s26["total"] == 8.0 and gb["students"][26]["status"] == "absent"
    # đổi cấu trúc câu/ý khi đã có điểm: bị chặn
    assert client.put(f"/api/exams/{eid}/manual-items", headers=gv, json={"items": items}).status_code == 409


def test_create_exam_from_any_excel(client, gv):
    """Bài làm ngoài hệ thống: file Excel tự soạn – tiêu đề ghi điểm tối đa và CLO; SV lạ được ghi danh khi đồng ý."""
    codes = [r[0] for r in _q("""SELECT s.student_code FROM enrollments e JOIN students s ON s.id=e.student_id
                                 WHERE e.class_section_id=1 ORDER BY s.student_code LIMIT 6""")]
    wb = Workbook(); ws = wb.active
    ws.append(["DANH SÁCH ĐIỂM THỰC HÀNH"]); ws.append([])
    ws.append(["STT", "Mã SV", "Họ tên", "Bài 1 (5đ) [CLO3]", "Bài 2 (5đ) [CLO4]", "Tổng"])
    for i, c in enumerate(codes):
        ws.append([i + 1, c, f"SV {i}", i % 6, 5 - (i % 6), 5])
    ws.append([7, "99999901", "Sinh viên ngoài danh sách", 4, 4, 8])
    buf = io.BytesIO(); wb.save(buf)
    f = lambda: {"file": ("th.xlsx", buf.getvalue(), "application/octet-stream")}   # noqa: E731
    form = {"exam_title": "Thực hành tuần 1–5", "assessment_type": "process", "max_score": "10"}
    pre = client.post("/api/class-sections/1/manual-exams/import?dry_run=true", headers=gv, files=f(), data=form).json()
    assert [i["points"] for i in pre["items"]] == [5, 5] and not pre["problems"] and pre["n_rows"] == 7
    assert [m["student_code"] for m in pre["missing_students"]] == ["99999901"]
    r = client.post("/api/class-sections/1/manual-exams/import?dry_run=false&enroll_missing=true", headers=gv, files=f(), data=form)
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["saved"] == 7 and d["analysis"]["n_students"] == 7
    assert _q("SELECT COUNT(*) FROM enrollments e JOIN students s ON s.id=e.student_id WHERE s.student_code='99999901' AND e.class_section_id=1")[0][0] == 1
    # thiếu CLO trong tiêu đề -> xem trước báo cột cần chọn CLO; chọn CLO qua clo_map thì tạo được
    wb2 = Workbook(); w2 = wb2.active
    w2.append(["MSSV", "Họ tên", "Điểm thuyết trình"]); w2.append(["", "Điểm tối đa", 10])
    for c in codes[:3]:
        w2.append([c, "x", 8])
    b2 = io.BytesIO(); wb2.save(b2)
    pre = client.post("/api/class-sections/1/manual-exams/import?dry_run=false", headers=gv,
                      files={"file": ("tt.xlsx", b2.getvalue(), "application/octet-stream")}, data=form).json()
    assert pre["dry_run"] and pre["problems"][0]["need_clos"] and pre["items"][0]["points"] == 10
    r = client.post("/api/class-sections/1/manual-exams/import?dry_run=false", headers=gv,
                    files={"file": ("tt.xlsx", b2.getvalue(), "application/octet-stream")},
                    data={**form, "clo_map": '{"%d": [1, 2]}' % pre["items"][0]["col"]})
    assert r.status_code == 200 and r.json()["saved"] == 3, r.text
    w = _q("SELECT m.weight FROM question_clo_mapping m JOIN exam_questions eq ON eq.question_id=m.question_id WHERE eq.exam_id=:e",
           e=r.json()["exam_id"])
    assert sorted(float(x[0]) for x in w) == [0.5, 0.5]


def test_combine_exams_and_evidence_flag(client, gv, gv_other):
    ex = client.get("/api/courses/5/analyzed-exams", headers=gv).json()
    ids = [e["id"] for e in ex if e["exam_type"] in ("paper", "online")]
    assert len(ids) == 2
    assert client.post("/api/combined-analysis", headers=gv_other, json={"exam_ids": ids}).status_code == 403
    m = client.post("/api/combined-analysis", headers=gv, json={"exam_ids": ids, "mode": "merge"}).json()
    cu = client.post("/api/combined-analysis", headers=gv, json={"exam_ids": ids, "mode": "cumulative"}).json()
    c1 = next(c for c in m["clos"] if c["clo_code"] == "CLO1")
    # merge: mỗi SV đếm 1 lần; cumulative: mỗi bài một lượt
    assert c1["n_evaluated"] <= 40 and next(c for c in cu["clos"] if c["clo_code"] == "CLO1")["n_evaluated"] == sum(b["n"] for b in c1["by_exam"])
    # kiểm tra lại công thức gộp cho một SV
    s = m["students"][0]
    rows = _q("""SELECT SUM(r.score_earned), SUM(r.score_max) FROM attempt_clo_results r JOIN exam_attempts a ON a.id=r.attempt_id
                 JOIN students s ON s.id=a.student_id JOIN clos c ON c.id=r.clo_id
                 WHERE s.student_code=:s AND c.clo_code='CLO1' AND a.exam_id IN (1, 2) AND a.status='finished'""", s=s["student_code"])
    assert abs(s["clos"]["CLO1"]["pct"] - round(100 * float(rows[0][0]) / float(rows[0][1]), 2)) < 0.01
    assert m["programs"] and all(p["plos"] for p in m["programs"])
    x = client.get(f"/api/combined-analysis.xlsx?ids={','.join(map(str, ids))}", headers=gv)
    assert x.status_code == 200 and load_workbook(io.BytesIO(x.content)).sheetnames == ["TongHop_CLO", "TungSinhVien"]
    # bỏ bài 1 khỏi minh chứng -> BM6b của lớp chỉ còn bài 2
    before = {r[0]: r[1] for r in _q("SELECT clo_id, n_evaluated FROM clo_results WHERE class_section_id=1")}
    assert client.patch(f"/api/exams/{ids[0]}/evidence", headers=gv, json={"counts_for_clo": False}).status_code == 200
    after = {r[0]: r[1] for r in _q("SELECT clo_id, n_evaluated FROM clo_results WHERE class_section_id=1")}
    assert sum(after.values()) < sum(before.values())
    assert client.patch(f"/api/exams/{ids[0]}/evidence", headers=gv, json={"counts_for_clo": True}).status_code == 200
    assert {r[0]: r[1] for r in _q("SELECT clo_id, n_evaluated FROM clo_results WHERE class_section_id=1")} == before


def test_remedial_flow(client, gv, gv_other):
    ov = client.get("/api/class-sections/1/remedial", headers=gv).json()
    assert client.get("/api/class-sections/1/remedial", headers=gv_other).status_code == 403
    assert ov["students"] and ov["questions"]
    worst = ov["students"][0]
    assert worst["n_wrong"] == max(s["n_wrong"] for s in ov["students"]) and worst["n_wrong"] > 0
    # số câu sai = số dòng chi tiết; mỗi dòng đúng là câu chưa đủ điểm
    d = client.get(f"/api/class-sections/1/remedial/students/{worst['student_id']}", headers=gv).json()
    assert len(d["wrong"]) == worst["n_wrong"]
    assert all(w["score_earned"] < w["points"] for w in d["wrong"])
    assert all(w["chosen_label"] != w["correct_label"] for w in d["wrong"] if w["question_type"] == "multichoice")
    q0 = ov["questions"][0]
    n = _q("""SELECT COUNT(*) FROM item_level_results r JOIN exam_attempts a ON a.id=r.attempt_id WHERE a.exam_id=:e
              AND r.question_id=:q AND a.status='finished' AND NOT r.is_correct""", e=q0["exam_id"], q=q0["question_id"])[0][0]
    assert q0["n_wrong"] == n
    # tạo bài phụ đạo cho 2 SV yếu nhất: câu đã sai + 1 câu tương tự mỗi câu
    two = [s["student_id"] for s in ov["students"][:2]]
    r = client.post("/api/class-sections/1/remedial-exams", headers=gv, json={"student_ids": two, "similar_per_question": 1, "seed": 3})
    assert r.status_code == 201, r.text
    res = r.json()
    e = client.get(f"/api/exams/{res['id']}", headers=gv).json()
    assert e["counts_for_clo"] is False and e["parent_exam_id"] and len(e["target_students"].split(",")) == 2
    assert len(e["questions"]) == res["n_questions"] and abs(sum(q["points"] for q in e["questions"]) - 10) < 0.01
    wrong_ids = {w["question_id"] for sid in two for w in client.get(f"/api/class-sections/1/remedial/students/{sid}", headers=gv).json()["wrong"]}
    assert {q["question_id"] for q in e["questions"]} >= set(list(wrong_ids)[:1])
    assert client.post("/api/class-sections/1/remedial-exams", headers=gv, json={"student_ids": [999999]}).status_code == 422
    assert client.delete(f"/api/exams/{res['id']}", headers=gv).status_code == 200
