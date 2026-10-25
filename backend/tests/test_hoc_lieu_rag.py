"""Học liệu theo học phần + gợi ý đọc đúng trang/slide (RAG BM25 cấp trang) cho câu SV làm sai."""
import io

import pytest
from sqlalchemy import text

from app.database import SessionLocal, engine
from app.services import hoc_lieu_mau, materials


def _q(sql, **kw):
    with engine.connect() as c:
        return c.execute(text(sql), kw).all()


@pytest.fixture(scope="module")
def course():
    return _q("SELECT id FROM courses WHERE course_code='DBMS330284'")[0][0]


def test_seed_co_hoc_lieu(client, gv, course):
    r = client.get(f"/api/courses/{course}/materials", headers=gv)
    assert r.status_code == 200
    ms = r.json()
    assert len(ms) == 7 and sum(m["n_pages"] for m in ms) == sum(len(v) for v in hoc_lieu_mau.SLIDES.values()) + len(hoc_lieu_mau.TEXTBOOK_PAGES)
    assert {m["file_ext"] for m in ms} == {"pptx", "pdf"}


@pytest.mark.parametrize("query,title,page", [
    ("Tệp nhật ký giao tác (.ldf) trong SQL Server được dùng chủ yếu để phục hồi", "Slide bài giảng – Chương 1", 4),
    ("Trong trigger AFTER UPDATE bảng tạm nào chứa giá trị CŨ của các dòng bị cập nhật", "Slide bài giảng – Chương 3", 1),
    ("Hai giao tác giữ khóa chờ nhau deadlock chọn nạn nhân victim", "Slide bài giảng – Chương 4", 4),
    ("Mỗi bảng trong SQL Server có tối đa bao nhiêu chỉ mục clustered", "Slide bài giảng – Chương 5", 1),
    ("Bản sao lưu differential chứa gì", "Slide bài giảng – Chương 6", 3),
])
def test_tim_dung_slide(client, gv, course, query, title, page):
    hits = client.get(f"/api/courses/{course}/materials/search", params={"q": query, "k": 3}, headers=gv).json()["hits"]
    assert hits, query
    assert (hits[0]["material"], hits[0]["page"]) == (title, page), hits[:2]
    assert hits[0]["unit"] == "slide" and hits[0]["snippet"]


def test_cau_ngoai_pham_vi_khong_goi_y(client, gv, course):
    assert client.get(f"/api/courses/{course}/materials/search", params={"q": "công thức nấu phở bò Hà Nội"}, headers=gv).json()["hits"] == []


def test_chi_tim_trong_hoc_lieu_cua_dung_khoa(client, gv, course):
    """Học liệu của học phần khác không bao giờ xuất hiện trong gợi ý của học phần này."""
    other = _q("SELECT id FROM courses WHERE id<>:c LIMIT 1", c=course)[0][0]
    db = SessionLocal()
    try:
        from pptx import Presentation
        prs = Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[1]); s.shapes.title.text = "Giao tác"
        s.placeholders[1].text = "Tính chất Atomicity của giao tác deadlock victim READ COMMITTED dirty read"
        b = io.BytesIO(); prs.save(b)
        mid = materials.add_material(db, course_id=other, title="Tài liệu môn khác", filename="khac.pptx", data=b.getvalue())["id"]
    finally:
        db.close()
    hits = client.get(f"/api/courses/{course}/materials/search", params={"q": "Atomicity giao tác deadlock victim", "k": 10}, headers=gv).json()["hits"]
    assert hits and all(h["material"] != "Tài liệu môn khác" for h in hits)
    assert client.get(f"/api/courses/{other}/materials/search", params={"q": "Atomicity giao tác deadlock victim"}, headers=gv).json()["hits"][0]["material"] == "Tài liệu môn khác"
    assert client.delete(f"/api/materials/{mid}", headers=gv).status_code == 200


def test_tai_len_docx_txt_va_loi(client, gv, course):
    from docx import Document
    d = Document(); d.add_heading("Chuẩn hóa dữ liệu", 1); d.add_paragraph("Dạng chuẩn 3NF loại bỏ phụ thuộc bắc cầu giữa các thuộc tính không khóa.")
    b = io.BytesIO(); d.save(b)
    r = client.post(f"/api/courses/{course}/materials", headers=gv, files={"file": ("chuan_hoa.docx", b.getvalue())},
                    data={"title": "Chuẩn hóa", "material_type": "reference"})
    assert r.status_code == 201 and r.json()["n_pages"] >= 1
    hit = client.get(f"/api/courses/{course}/materials/search", params={"q": "dạng chuẩn 3NF phụ thuộc bắc cầu"}, headers=gv).json()["hits"][0]
    assert hit["material"] == "Chuẩn hóa" and hit["unit"] == "mục" and "3NF" in hit["snippet"]
    assert client.post(f"/api/courses/{course}/materials", headers=gv, files={"file": ("x.exe", b"MZ")}).status_code == 422
    assert client.post(f"/api/courses/{course}/materials", headers=gv, files={"file": ("rong.txt", b"   \n")}).status_code == 422
    assert client.post(f"/api/courses/{course}/materials", headers=gv, files={"file": ("a.txt", b"x")}, data={"material_type": "zzz"}).status_code == 422
    mid = client.get(f"/api/courses/{course}/materials", headers=gv).json()[-1]["id"]
    assert client.delete(f"/api/materials/{mid}", headers=gv).status_code == 200


def test_goi_y_doc_cho_sv_lam_sai(client, gv):
    """Phân tích bài KT → mỗi SV làm sai câu nào thì có gợi ý đọc trang/slide tương ứng của học phần."""
    ex = _q("SELECT id FROM exams WHERE status='Analyzed' AND exam_type='online' ORDER BY id LIMIT 1")[0][0]
    rows = client.get(f"/api/exams/{ex}/moodle/results-preview", headers=gv)
    assert rows.status_code == 200, rows.text
    with_reads = [r for r in rows.json() if r["data"].get("readings")]
    assert len(with_reads) >= 20
    r0 = with_reads[0]["data"]["readings"]
    assert 1 <= len(r0) <= 8 and all({"material", "page", "unit", "snippet", "score"} <= set(x) for x in r0)
    assert "question_ids" not in r0[0] and r0[0]["n_wrong"] >= 1


def test_tai_tao_goi_y_sau_khi_them_hoc_lieu(client, gv):
    ex = _q("SELECT id FROM exams WHERE status='Analyzed' AND exam_type='online' ORDER BY id LIMIT 1")[0][0]
    r = client.post(f"/api/exams/{ex}/readings/rebuild", headers=gv)
    assert r.status_code == 200 and r.json()["n_with_readings"] >= 20
    assert client.post("/api/exams/999999/readings/rebuild", headers=gv).status_code == 404


# ------------------------------------------------------------------ đồng bộ hai chiều với Moodle (Moodle giả lập bằng monkeypatch)
def _pptx(text_):
    from pptx import Presentation
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1]); s.shapes.title.text = "Chuẩn hóa"
    s.placeholders[1].text = text_
    b = io.BytesIO(); prs.save(b)
    return b.getvalue()


def test_lay_hoc_lieu_tu_moodle_ve(client, gv, course, monkeypatch):
    from app.services import lms
    state = {"mod": 100, "body": "Dạng chuẩn BCNF chặt hơn 3NF: mọi định thức đều là siêu khóa."}
    base = "http://moodle.test/webservice/pluginfile.php/55/mod_resource/content/1/"
    monkeypatch.setattr(lms, "course_files", lambda mcid: [
        {"cmid": 7, "name": "Slide chuẩn hóa (Moodle)", "section": "Tuần 1", "filename": "chuan_hoa.pptx", "url": base + "chuan_hoa.pptx", "size": 1, "modified": state["mod"]},
        {"cmid": 8, "name": "Ảnh", "section": "Tuần 1", "filename": "logo.png", "url": base + "logo.png", "size": 1, "modified": 1}])
    monkeypatch.setattr(lms, "download", lambda url: (_pptx(state["body"]), "application/octet-stream"))
    cs_mcid = _q("SELECT moodle_course_id FROM class_sections WHERE course_id=:c AND moodle_course_id IS NOT NULL LIMIT 1", c=course)[0][0]

    r = client.post(f"/api/courses/{course}/materials/sync-moodle", headers=gv, json={})
    assert r.status_code == 200, r.text
    j = r.json()
    assert (j["moodle_course_id"], j["added"], j["updated"], j["skipped"], j["errors"]) == (cs_mcid, 1, 0, 1, [])   # .png bị bỏ qua
    hit = client.get(f"/api/courses/{course}/materials/search", params={"q": "BCNF định thức siêu khóa"}, headers=gv).json()["hits"][0]
    assert hit["material"] == "Slide chuẩn hóa (Moodle)"

    j = client.post(f"/api/courses/{course}/materials/sync-moodle", headers=gv, json={}).json()
    assert (j["added"], j["updated"], j["skipped"]) == (0, 0, 2)                                 # chưa đổi → không tải lại

    state["mod"], state["body"] = 200, "Chuẩn hóa mới: dạng chuẩn 4NF xử lý phụ thuộc đa trị."
    j = client.post(f"/api/courses/{course}/materials/sync-moodle", headers=gv, json={}).json()
    assert (j["added"], j["updated"]) == (0, 1)                                                  # sửa trên Moodle → lấy lại bản mới
    assert client.get(f"/api/courses/{course}/materials/search", params={"q": "4NF phụ thuộc đa trị"}, headers=gv).json()["hits"][0]["material"] == "Slide chuẩn hóa (Moodle)"
    ms = [m for m in client.get(f"/api/courses/{course}/materials", headers=gv).json() if m["source"] == "moodle"]
    assert len(ms) == 1 and ms[0]["moodle_cmid"] == 7
    assert client.post(f"/api/materials/{ms[0]['id']}/push-moodle", headers=gv, json={}).status_code == 409   # tệp từ Moodle không đẩy ngược
    assert client.delete(f"/api/materials/{ms[0]['id']}", headers=gv).status_code == 200


def test_day_hoc_lieu_len_moodle(client, gv, course, monkeypatch):
    from app.services import lms
    calls = []
    monkeypatch.setattr(lms, "add_resource", lambda mcid, name, filename, data, section=0: calls.append((mcid, name, filename, len(data))) or {"cmid": 321, "created": True, "url": "http://moodle.test/mod/resource/view.php?id=321"})
    r = client.post(f"/api/courses/{course}/materials", headers=gv, files={"file": ("de_day.pptx", _pptx("Nội dung đẩy lên Moodle để sinh viên xem"))}, data={"title": "Slide đẩy lên"})
    mid = r.json()["id"]
    j = client.post(f"/api/materials/{mid}/push-moodle", headers=gv, json={})
    assert j.status_code == 200 and j.json()["cmid"] == 321 and calls and calls[0][2] == "de_day.pptx"
    again = client.post(f"/api/materials/{mid}/push-moodle", headers=gv, json={}).json()
    assert again == {"cmid": 321, "created": False} and len(calls) == 1                            # gọi lại không tạo trùng
    assert any(m["id"] == mid and m["moodle_cmid"] == 321 for m in client.get(f"/api/courses/{course}/materials", headers=gv).json())
    assert client.delete(f"/api/materials/{mid}", headers=gv).status_code == 200


def test_moodle_loi_tra_502(client, gv, course, monkeypatch):
    from app.services import lms
    def boom(_):
        raise lms.MoodleWSError("Lỗi kết nối LMS: giả lập")
    monkeypatch.setattr(lms, "course_files", boom)
    r = client.post(f"/api/courses/{course}/materials/sync-moodle", headers=gv, json={})
    assert r.status_code == 502 and "giả lập" in r.json()["detail"]


def test_on_tap_co_goi_y_doc(client, gv):
    """Trang Ôn tập của GV: mỗi SV yếu và mỗi chương đều có 'nên đọc lại' trỏ đúng tài liệu của môn."""
    ov = client.get("/api/class-sections/1/remedial", headers=gv).json()
    mat_titles = {r[0] for r in _q("SELECT title FROM course_materials")}
    with_reads = [s for s in ov["students"] if s.get("readings")]
    assert with_reads, "không SV nào có gợi ý đọc"
    for s in with_reads[:5]:
        for r in s["readings"]:
            assert r["material"] in mat_titles and r["page"] >= 1
    chs = [c for c in ov.get("chapters", []) if c.get("readings")]
    assert chs and all(r["material"] in mat_titles for c in chs for r in c["readings"])


def test_xem_tep_goc_trang_va_lap_lai_chi_muc(client, gv, course):
    """Xem lại đúng trang đã được gợi ý, tải tệp gốc và lập lại chỉ mục không làm mất trang."""
    ms = client.get(f"/api/courses/{course}/materials", headers=gv).json()
    m = next(x for x in ms if x["title"] == "Slide bài giảng – Chương 4")
    p = client.get(f"/api/materials/{m['id']}/pages/2", headers=gv)
    assert p.status_code == 200 and "dirty read" in p.text.lower()
    assert client.get(f"/api/materials/{m['id']}/pages/999", headers=gv).status_code == 404
    f = client.get(f"/api/materials/{m['id']}/file", headers=gv)
    assert f.status_code == 200 and f.content[:2] == b"PK"          # PPTX là tệp zip
    r = client.post(f"/api/materials/{m['id']}/reindex", headers=gv)
    assert r.status_code == 200
    again = client.get(f"/api/courses/{course}/materials", headers=gv).json()
    assert next(x for x in again if x["id"] == m["id"])["n_pages"] == m["n_pages"]


def test_lay_ve_khong_trung_voi_tep_da_day_len(client, gv, course, monkeypatch):
    """Lỗi gặp trên Moodle thật: tệp hệ thống đã đẩy lên (cmid 555) bị lấy về lần nữa thành bản trùng."""
    from app.services import lms
    monkeypatch.setattr(lms, "add_resource", lambda mcid, name, filename, data, section=0: {"cmid": 555, "created": True, "url": "http://moodle.test/mod/resource/view.php?id=555"})
    mid = client.post(f"/api/courses/{course}/materials", headers=gv, files={"file": ("trung.pptx", _pptx("Tệp đẩy lên rồi lấy về không được nhân đôi"))}, data={"title": "Slide trùng"}).json()["id"]
    assert client.post(f"/api/materials/{mid}/push-moodle", headers=gv, json={}).status_code == 200
    base = "http://moodle.test/webservice/pluginfile.php/55/mod_resource/content/1/trung.pptx"
    monkeypatch.setattr(lms, "course_files", lambda mcid: [{"cmid": 555, "name": "Slide trùng", "section": "Tuần 1", "filename": "trung.pptx", "url": base, "size": 1, "modified": 999}])
    monkeypatch.setattr(lms, "download", lambda url: (_pptx("x"), "application/octet-stream"))
    j = client.post(f"/api/courses/{course}/materials/sync-moodle", headers=gv, json={}).json()
    assert (j["added"], j["updated"], j["skipped"]) == (0, 0, 1)
    assert len([m for m in client.get(f"/api/courses/{course}/materials", headers=gv).json() if m["file_name"] == "trung.pptx"]) == 1
    assert client.delete(f"/api/materials/{mid}", headers=gv).status_code == 200
